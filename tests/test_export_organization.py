"""Export scope, provenance and deterministic arithmetic survive the layout cleanup."""
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import liquidity_simulation as ls
import scenario_export

ROOT = Path(__file__).resolve().parents[1]


class ExportOrganization(unittest.TestCase):
    def test_optional_summary_has_no_detailed_company_sheets(self):
        w = load_workbook(ROOT / "samples/liquidity_pilots.xlsx", data_only=True)
        self.assertEqual(w.sheetnames, ["Debt and liquidity", "Liquidity inputs", "About"])
        self.assertIn("not a peer ranking", w["About"]["B7"].value)
        self.assertEqual(w["Debt and liquidity"]["C2"].value, 8_950_000_000)
        self.assertIsNotNone(w["Liquidity inputs"]["K2"].hyperlink)
        w.close()

    def test_scenario_sources_and_formulas_are_traceable(self):
        i = ls.oracle_inputs(); decisions = ls.oracle_decisions()
        rs = [(ls.simulate(i, d, 1000, 7), d) for d in decisions]
        raw = scenario_export.workbook(i, rs, 1000, 7, uses=ls.oracle_uses_outside_model())
        formulas = load_workbook(BytesIO(raw)); values = load_workbook(BytesIO(raw), data_only=True)
        self.assertIn("Data", formulas.sheetnames); self.assertIn("About", formulas.sheetnames)
        self.assertEqual(formulas["Scenario"]["C6"].value, "='Data'!C2")
        self.assertEqual(values["Scenario"]["C7"].value, i.principal_due)
        self.assertEqual(values["Uses not modelled"]["C3"].value, 3_937_000_000)
        self.assertIn("+'Data'!C7", formulas["Uses not modelled"]["C3"].value)
        self.assertIn("not an outstanding balance", values["Uses not modelled"]["D3"].value)
        for row in values["Data"].iter_rows(min_row=2):
            self.assertIsNotNone(row[8].hyperlink)
            self.assertEqual(row[8].hyperlink.target, row[8].value)
        self.assertIn("does not rerun Monte Carlo", values["About"]["B4"].value)
        formulas.close(); values.close()

    def test_template_fingerprint_refuses_unreviewed_layout(self):
        i = ls.oracle_inputs(); d = ls.oracle_decisions()[0]
        with patch.object(scenario_export, "TEMPLATE", ROOT / "samples/liquidity_pilots.xlsx"):
            with self.assertRaises((ValueError, KeyError)):
                scenario_export.workbook(i, [(ls.simulate(i, d, 10), d)], 10, 2026)

    def test_light_vertical_separators_and_original_links_remain(self):
        for stem in ("lly", "oracle", "vistra"):
            w = load_workbook(ROOT / f"samples/{stem}_liquidity.xlsx")
            self.assertEqual(w["Data"]["B2"].border.left.style, "thin")
            self.assertEqual(len(w["Core charts"]._charts), 2)
            self.assertIsNotNone(next(c for row in w["Data"] for c in row if c.hyperlink).hyperlink)
            w.close()
