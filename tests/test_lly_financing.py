"""Contract dates must not turn initial face principal or old facilities into current funding."""
import json
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import lly_financing as lf

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT.parent / "data/raw/sec/LLY"


class LillyFinancing(unittest.TestCase):
    def test_eight_initial_series_reconcile_without_becoming_current_debt(self):
        r = lf.load()
        self.assertEqual(r["status"], "Checked")
        self.assertEqual([x["initial_principal"] for x in r["rows"]],
                         [750e6, 500e6, 750e6, 1500e6, 1250e6, 1500e6, 1750e6, 1000e6])
        self.assertEqual(sum(x["initial_principal"] for x in r["rows"]), r["reported_issue_total"])
        self.assertEqual(r["initial_principal_total"], 9e9)
        self.assertTrue(all(x["current_outstanding"] is None and x["tag"] is None for x in r["rows"]))
        self.assertIsNone(r["drawable"])
        self.assertIsNone(r["coverage"])

    def test_stated_maturities_rates_and_contract_sources_are_preserved(self):
        r = lf.load()
        self.assertEqual([x["stated_maturity"] for x in r["rows"]],
                         [f"{y}-05-20" for y in [2028, 2029, 2029, 2031, 2033, 2036, 2056, 2066]])
        self.assertEqual(r["rows"][2]["coupon"], 0.0415)
        self.assertEqual([x["sofr_margin"] for x in r["rows"][:2]], [.0035, .0046])
        self.assertTrue(all(r["sources"][x["source"]]["accession"] == "0001193125-26-232707" for x in r["rows"]))
        self.assertIn("business-day", r["date_detail"])
        self.assertIn("4.650%", r["coupon_conflict"])
        self.assertIn("4.150%", r["coupon_conflict"])

    def test_historical_renewal_and_resolved_redemption_are_not_new_funding_or_bill(self):
        r = lf.load()
        self.assertIn("not an August 2026 renewal", r["credit_detail"])
        self.assertIn("June 24, 2026", r["redemption_detail"])
        self.assertIn("not presented as a future 101%", r["redemption_detail"])
        self.assertIn("not used to allocate H1 repayments", r["delisting_detail"])

    def test_inventory_covers_cutoff_without_carrying_old_counts_forward(self):
        inv = lf.load()["inventory"]
        self.assertEqual(inv["cutoff"], "2026-10-15")
        self.assertLessEqual(inv["oldest"], "2026-01-01")
        after = [r for r in inv["eligible"] if r["filingDate"] > "2026-06-30"]
        self.assertEqual([(r["form"], r["accessionNumber"]) for r in after],
                         [("10-Q", "0000059478-26-000081"), ("8-K", "0000059478-26-000077")])
        self.assertTrue(all(r["filingDate"] <= inv["cutoff"] for r in inv["eligible"]))

    def test_changed_or_missing_saved_review_is_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / "review.json"
            with patch.object(lf, "SAMPLE", p):
                self.assertEqual(lf.load()["status"], "Unavailable")
                r = lf._review();r["drawable"] = 10e9
                p.write_text(json.dumps(r))
                self.assertEqual(lf.load()["status"], "Unavailable")
                r = lf._review();r["rows"][0]["current_outstanding"] = 750e6
                p.write_text(json.dumps(r))
                self.assertEqual(lf.load()["status"], "Unavailable")

    def test_snapshot_date_and_company_cannot_inherit_terms(self):
        s = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"][0]
        self.assertEqual(lf.for_snapshot(s)["status"], "Checked")
        other = deepcopy(s);other["metrics"]["cash"]["period_end"] = "2026-09-30"
        self.assertEqual(lf.for_snapshot(other)["status"], "Unavailable")
        other["ticker"] = "ORCL"
        self.assertIsNone(lf.for_snapshot(other))

    @unittest.skipUnless((RAW / "0001193125-26-232707-d148580dex43.htm").exists(), "Full saved certificates unavailable")
    def test_saved_source_hashes_reproduce_audit_and_changed_source_is_refused(self):
        self.assertEqual(lf.build(RAW), lf.load())
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder);name = lf.SOURCES["offering"]["saved"]
            self.assertIn("Missing", lf.build(p)["detail"])
            (p / name).write_bytes((RAW / name).read_bytes() + b"changed")
            self.assertIn("Changed", lf.build(p)["detail"])

    def test_workbook_initial_amounts_and_dates_match_shared_audit(self):
        from openpyxl import load_workbook
        import company_review
        snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]
        b = company_review.saved_workbook(snapshots, ticker="LLY")
        self.assertIsNotNone(b)
        v, f = load_workbook(BytesIO(b), data_only=True), load_workbook(BytesIO(b), data_only=False)
        r = lf.load()
        for i, x in enumerate(r["rows"], 77):
            self.assertEqual(v["Repayment review"][f"B{i}"].value, x["initial_principal"])
            self.assertEqual(v["Repayment review"][f"C{i}"].value.strftime("%Y-%m-%d"), x["stated_maturity"])
            self.assertTrue(f["Repayment review"][f"B{i}"].value.startswith("="))
        self.assertEqual(v["Repayment review"]["B87"].value, 0)
        self.assertEqual(v["Repayment review"]["D79"].value, .0415)
        self.assertEqual(v["Repayment review"]["B88"].value, "Unavailable")
        self.assertEqual(v["Repayment review"]["B89"].value, "Unavailable")

    def test_download_refuses_stale_financing_source(self):
        import company_review
        snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for name in ["liquidity_pilots.json", "lly_liquidity.json", "lly_financing.json", "lly_interest.json", "lly_liquidity.xlsx", "lly_liquidity.manifest.json"]:
                (root / name).write_bytes((ROOT / "samples" / name).read_bytes())
            self.assertIsNotNone(company_review.saved_workbook(snapshots, root, ticker="LLY"))
            (root / "lly_financing.json").write_text("{}")
            self.assertIsNone(company_review.saved_workbook(snapshots, root, ticker="LLY"))


if __name__ == "__main__":
    unittest.main()
