"""Sign reconciliation and a checked original comparative gate Vistra's coverage."""
import html
import json
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

import vistra_interest as vi
import vistra_interest_view as view
import vistra_liquidity as vl
from company_review import money

ROOT = Path(__file__).resolve().parents[1]
FILES = [ROOT / "tests/fixtures" / n for n in ["VST_2025_interest.htm", "VST_2026H1_interest.htm", "VST_2025H1_interest.htm"]]
M = 1_000_000


class VistraInterest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.review = vi.build(*FILES)
        cls.snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]

    def test_checked_periods_and_ttm_coverage(self):
        c = self.review["calculation"]
        self.assertEqual([c["periods"][p]["interest"] for p in ["fy", "current", "prior", "prior_original"]], [1107*M,557*M,547*M,547*M])
        self.assertEqual([c["periods"][p]["operating_income"] for p in ["fy", "current", "prior"]], [1906*M,2052*M,395*M])
        self.assertEqual((c["ttm"]["operating_income"], c["ttm"]["interest"]), (3563*M,1117*M))
        self.assertAlmostEqual(c["value"], 3.189794091316025)
        self.assertTrue(c["prior_agrees"])
        self.assertEqual((self.review["period_start"],self.review["period_end"],self.review["tag"]), ("2025-07-01","2026-06-30",None))

    def test_raw_gain_signs_are_preserved_and_expense_signs_reconcile(self):
        c = self.review["calculation"]
        for role, swaps, extinguishment in [("fy",-67,0),("current",7,-35),("prior",-74,0)]:
            p = c["periods"][role]
            self.assertEqual(p["swaps"],swaps*M);self.assertEqual(p["extinguishment"],extinguishment*M)
            self.assertEqual(p["effects"]["swaps"],-swaps*M)
            self.assertEqual(p["effects"]["extinguishment"],-extinguishment*M)
            self.assertEqual(p["difference"],0)
        t = c["ttm"]
        self.assertEqual({k:v/M for k,v in t.items() if k!='operating_income'},
                         dict(interest=1117,swaps=14,amortization=45,extinguishment=-35,capitalized=151,other=100,total=1132))
        self.assertEqual(sum(t[k]*s for k,s in vi.EFFECTS.items()),1132*M)

    def test_capitalized_interest_is_deducted_only_in_mixed_total(self):
        c = self.review["calculation"]
        self.assertEqual([c["periods"][p]["capitalized"] for p in ["fy","current","prior"]],[125*M,82*M,56*M])
        self.assertEqual(c["ttm"]["capitalized"],151*M)
        self.assertNotEqual(c["ttm"]["interest"],(1117-151)*M)
        self.assertNotEqual(c["ttm"]["interest"],(1117+151)*M)
        self.assertIn("not added to or subtracted",self.review["basis"])

    def test_component_and_unknown_allocation_are_distinct(self):
        self.assertNotEqual(self.review["calculation"]["ttm"]["interest"],self.review["calculation"]["ttm"]["total"])
        for gap in self.review["unavailable"].values():
            self.assertIsNone(gap["value"]);self.assertEqual(gap["status"],"Unavailable")
        self.assertIn("Receivables-facility and Nuveen",self.review["composition"])
        self.assertIn("commodity-hedge",self.review["limitations"])

    def test_exact_contexts_and_parent_facts_are_retained_but_not_selected(self):
        root = ET.parse(FILES[0]).getroot()
        self.assertTrue(any(c.find('.//xbrldi:explicitMember',vi.fx.NS) is not None for c in root.iter(f"{{{vi.fx.XBRLI}}}context")))
        for role, context in [("fy","c-1"),("current","c-1"),("prior","c-7"),("prior_original","c-1")]:
            for key in vi.TAGS:
                f = self.review["inputs"][role+'_'+key]
                self.assertEqual(f["context_id"],context);self.assertEqual(f["dimensions"],[])
                self.assertEqual(f["unit"],vi.fx.USD)

    def test_wrong_tag_period_unit_dimension_entity_accession_and_source_are_refused(self):
        for key,val in [("tag","us-gaap:InterestExpenseNonoperating"),("period_start","2026-04-01"),
                        ("period_end","2026-06-29"),("unit","iso4217:EUR"),("dimensions",[["a","b"]]),
                        ("entity_identifier","1"),("accession","other"),("source_sha256","unchecked")]:
            with self.subTest(key=key):
                i = deepcopy(self.review["inputs"]);i["current_interest"][key]=val
                c=vi.calculate(i);self.assertIsNone(c["value"]);self.assertIsNone(c["ttm"]["interest"])
                self.assertEqual(c["ttm"]["operating_income"],3563*M)

    def test_missing_bridge_input_blocks_ratio_but_retains_supported_component(self):
        for key in ["current_swaps","fy_amortization","prior_extinguishment","current_capitalized","fy_other","current_total"]:
            with self.subTest(key=key):
                i=deepcopy(self.review["inputs"]);i[key]["value"]=None
                c=vi.calculate(i);self.assertIsNone(c["value"]);self.assertEqual(c["ttm"]["interest"],1117*M)

    def test_wrong_swap_extinguishment_and_capitalization_signs_block_coverage(self):
        for key in ["fy_swaps","current_extinguishment","current_capitalized"]:
            i=deepcopy(self.review["inputs"]);i[key]["value"] *= -1
            self.assertIsNone(vi.calculate(i)["value"])

    def test_original_comparative_mismatch_is_caught_even_when_each_table_still_ties(self):
        i=deepcopy(self.review["inputs"])
        i["prior_original_other"]["value"] += M;i["prior_original_total"]["value"] += M
        c=vi.calculate(i)
        self.assertEqual(c["periods"]["prior_original"]["difference"],0)
        self.assertFalse(c["prior_agrees"]);self.assertIsNone(c["value"])

    def test_zero_and_negative_ttm_denominators_do_not_divide(self):
        for fy,current in [(0,547*M),(0,0)]:
            i=deepcopy(self.review["inputs"])
            for role,v in [("fy",fy),("current",current)]:
                delta=v-i[role+'_interest']["value"];i[role+'_interest']["value"]=v;i[role+'_total']["value"]+=delta
            c=vi.calculate(i)
            self.assertTrue(all(p["difference"]==0 for p in c["periods"].values()))
            self.assertIsNone(c["value"]);self.assertIn("zero or negative",c["detail"])

    def test_operating_loss_is_retained_without_addbacks(self):
        i=deepcopy(self.review["inputs"]);i["fy_operating_income"]["value"]=-1906*M
        c=vi.calculate(i);self.assertEqual(c["ttm"]["operating_income"],-249*M);self.assertLess(c["value"],0)

    def changed(self,mutate):
        root=ET.parse(FILES[1]).getroot();mutate(root)
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'changed.htm';ET.ElementTree(root).write(p)
            return vi.build(FILES[0],p,FILES[2])

    def test_conflicting_inline_interest_duplicate_is_refused(self):
        def mutate(root):
            f=next(f for f in root.iter(f"{{{vi.fx.IX}}}nonFraction") if f.get('name')==vi.TAGS['interest'] and f.get('contextRef')=='c-1')
            extra=deepcopy(f);extra.text='558';root.append(extra)
        c=self.changed(mutate)["calculation"];self.assertIsNone(c["value"])

    def test_missing_consolidated_operating_income_is_not_replaced_with_segment(self):
        def mutate(root):
            for f in root.iter(f"{{{vi.fx.IX}}}nonFraction"):
                if f.get('name')==vi.TAGS['operating_income'] and f.get('contextRef')=='c-1':f.set('name','x:removed')
        c=self.changed(mutate)["calculation"];self.assertIsNone(c["value"]);self.assertIsNone(c["ttm"]["operating_income"])
        self.assertEqual(c["ttm"]["interest"],1117*M)

    def test_fixed_zero_extinguishment_is_reported_zero(self):
        for role in ['fy','prior','prior_original']:
            self.assertEqual(self.review['inputs'][role+'_extinguishment']['value'],0)
        self.assertTrue(self.review['calculation']['prior_agrees'])

    def test_saved_review_matches_trimmed_and_all_full_sources(self):
        self.assertEqual(vi.load(),self.review)
        raw=ROOT.parent/'data/raw/sec/VST'
        if (raw/vi.SOURCES['annual']['document']).exists():
            self.assertEqual(vi.build(*(raw/vi.SOURCES[s]['document'] for s in ['annual','quarter','prior_original'])),self.review)

    def test_stale_missing_review_and_wrong_snapshot_are_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'review.json'
            with patch.object(vi,'SAMPLE',p):
                self.assertIsNone(vi.load())
                r=deepcopy(self.review);r['inputs']['fy_interest']['value']+=M;r['calculation']=vi.calculate(r['inputs'])
                p.write_text(json.dumps(r));self.assertIsNone(vi.load())
        self.assertIsNotNone(vi.for_snapshot(self.snapshots[2]))
        self.assertIsNone(vi.for_snapshot(self.snapshots[0]));self.assertIsNone(vi.for_snapshot(self.snapshots[1]))
        s=deepcopy(self.snapshots[2]);s['metrics']['cash']['accession']='other';self.assertIsNone(vi.for_snapshot(s))

    def test_view_shows_period_build_signs_limits_and_sources(self):
        shown=html.unescape(view.render(self.snapshots[2],money))
        for text in ['3.19x','$3.563B','$1.117B','before capitalized interest','1,132','151','Receivables-facility and Nuveen','gain-positive','original H1 2025','Unavailable']:
            self.assertIn(text,shown)
        self.assertEqual(view.render(self.snapshots[0],money),'');self.assertNotIn('$',view.render(self.snapshots[2],money))

    def test_workbook_calculations_and_stale_source_guard(self):
        from openpyxl import load_workbook
        b=vl.workbook_bytes(json.loads(vl.SAMPLE.read_text()));self.assertIsNotNone(b)
        v,f=load_workbook(BytesIO(b),data_only=True),load_workbook(BytesIO(b),data_only=False)
        self.assertEqual(v.sheetnames,['Funding conclusion','Borrowing and leases','Debt and liquidity','History','Core charts','Cash review','Payment dates','Financing capacity','Data','About'])
        for cell,value in {'B68':3563*M,'B69':1117*M,'B70':-14*M,'B71':45*M,'B72':35*M,'B73':-151*M,'B74':100*M,'B75':1132*M,'B76':1132*M,'B77':0,'B79':3563/1117,'D91':0}.items():
            self.assertAlmostEqual(v['Cash review'][cell].value,value)
            self.assertTrue(f['Cash review'][cell].value.startswith('='))
        self.assertEqual(v['Cash review']['B47'].value,435*M)
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'interest.json';p.write_bytes(vi.SAMPLE.read_bytes()+b'\n')
            with patch.object(vi,'SAMPLE',p):self.assertIsNone(vl.workbook_bytes(json.loads(vl.SAMPLE.read_text())))


if __name__=='__main__':
    unittest.main()
