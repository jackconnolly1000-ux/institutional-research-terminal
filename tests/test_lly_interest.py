"""Exact period/dimension selection and a reconciled, derived numerator are required."""
import json
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

import filing_xbrl as fx
import lly_interest as li

ROOT = Path(__file__).resolve().parents[1]
ANNUAL = ROOT / "tests/fixtures/LLY_2025_interest.htm"
QUARTER = ROOT / "tests/fixtures/LLY_2026H1_interest.htm"


class LillyInterest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.review = li.build(ANNUAL, QUARTER)

    def test_hand_checked_ttm_and_independent_controls(self):
        r = self.review;c = r["calculation"]
        self.assertEqual([c["periods"][k]["interest"] for k in ["fy", "current", "prior"]], [895e6, 677e6, 493e6])
        self.assertEqual([c["periods"][k]["operating_income"] for k in ["fy", "current", "prior"]], [26302e6, 17893e6, 10562e6])
        self.assertTrue(all(p["operating_income"] == p["control"] for p in c["periods"].values()))
        self.assertEqual((c["operating_income"], c["interest"]), (33633e6, 1079e6))
        self.assertAlmostEqual(c["value"], 31.17052826691381)
        self.assertEqual((r["period_start"], r["period_end"], r["tag"]), ("2025-07-01", "2026-06-30", None))

    def test_checked_segment_rule_and_six_month_periods(self):
        i = self.review["inputs"]
        self.assertEqual(i["current_interest"]["context_id"], "c-397")
        self.assertEqual(i["prior_interest"]["context_id"], "c-398")
        self.assertEqual(i["current_interest"]["dimensions"], [list(d) for d in li.SEGMENT])
        self.assertEqual(i["fy_interest"]["dimensions"], [])
        rejected = fx.read_duration(QUARTER, i["current_interest"]["accession"], li.INTEREST, "2026-01-01", "2026-06-30")
        self.assertIsNone(rejected["value"])
        self.assertNotEqual(i["current_interest"]["value"], 345e6)

    def test_statement_selection_excludes_conflicting_rounded_iprd_without_weakening_reader(self):
        f = self.review["inputs"]["current_iprd"]
        self.assertEqual(f["value"], 3360e6)
        self.assertEqual(self.review["inputs"]["fy_iprd"]["value"], 2910e6)
        whole = fx.read_duration(QUARTER, f["accession"], f["tag"], f["period_start"], f["period_end"])
        self.assertIsNone(whole["value"])
        self.assertIn("conflicting", whole["detail"])

    def test_capitalization_basis_stays_qualified_and_source_signs_preserved(self):
        self.assertIn("FY2025", self.review["basis"])
        self.assertIn("without explicitly stating", self.review["basis"])
        self.assertIn("not added or subtracted", self.review["basis"])
        self.assertEqual([self.review["inputs"][k + "_other_income"]["value"] for k in ["fy", "current", "prior"]], [-571e6, 204e6, -329e6])

    def test_incompatible_interest_metadata_is_refused_without_losing_numerator(self):
        for key, value in [("period_end", "2026-06-29"), ("period_start", "2026-04-01"), ("unit", "EUR"),
                           ("tag", "us-gaap:InterestAndDebtExpense"), ("accession", "other"), ("dimensions", []),
                           ("entity_identifier", "1"), ("source_sha256", "unchecked")]:
            with self.subTest(field=key):
                inputs = deepcopy(self.review["inputs"]);inputs["current_interest"][key] = value
                c = li.calculate(inputs)
                self.assertIsNone(c["value"]);self.assertIsNone(c["interest"])
                self.assertEqual(c["operating_income"], 33633e6)

    def test_missing_numerator_and_denominator_are_independent(self):
        for key in ["fy_revenue", "current_interest"]:
            inputs = deepcopy(self.review["inputs"]);inputs[key]["value"] = None
            c = li.calculate(inputs)
            self.assertIsNone(c["value"])
            self.assertEqual(c["interest"] if key == "fy_revenue" else c["operating_income"], 1079e6 if key == "fy_revenue" else 33633e6)

    def test_zero_and_negative_ttm_interest_do_not_divide(self):
        for prior, fy, current in [(0, 0, 0), (2000e6, 895e6, 677e6)]:
            inputs = deepcopy(self.review["inputs"])
            for role, value in [("fy", fy), ("current", current), ("prior", prior)]:inputs[role + "_interest"]["value"] = value
            c = li.calculate(inputs)
            self.assertIsNone(c["value"])
            self.assertIn("zero or negative", c["detail"])

    def test_negative_operating_income_is_retained(self):
        inputs = deepcopy(self.review["inputs"])
        inputs["fy_rd"]["value"] += 100e9;inputs["fy_pretax"]["value"] -= 100e9
        c = li.calculate(inputs)
        self.assertEqual(c["operating_income"], -66367e6)
        self.assertLess(c["value"], 0)

    def test_unreconciled_costs_block_numerator_and_ratio(self):
        inputs = deepcopy(self.review["inputs"]);inputs["fy_revenue"]["value"] += 1e6
        c = li.calculate(inputs)
        self.assertIsNone(c["value"]);self.assertIsNone(c["operating_income"])
        self.assertEqual(c["interest"], 1079e6)
        self.assertIn("do not reconcile", c["detail"])

    def _changed_quarter(self, mutate):
        root = ET.parse(QUARTER).getroot();mutate(root)
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / "changed.htm";ET.ElementTree(root).write(p)
            return li.build(ANNUAL, p)

    def test_unexpected_extra_dimension_in_filing_is_refused(self):
        def mutate(root):
            ctx = next(c for c in root.iter(f"{{{fx.XBRLI}}}context") if c.get("id") == "c-397")
            segment = ctx.find("xbrli:entity/xbrli:segment", fx.NS)
            ET.SubElement(segment, f"{{{fx.XBRLDI}}}explicitMember", {"dimension": "us-gaap:ProductOrServiceAxis"}).text = "lly:UnexpectedMember"
        r = self._changed_quarter(mutate)
        self.assertIsNone(r["calculation"]["value"])
        self.assertIsNone(r["inputs"]["current_interest"]["value"])

    def test_duplicate_conflicting_interest_in_filing_is_refused(self):
        def mutate(root):
            f = next(f for f in root.iter(f"{{{fx.IX}}}nonFraction") if f.get("name") == li.INTEREST and f.get("contextRef") == "c-397")
            extra = deepcopy(f);extra.text = "678";root.append(extra)
        r = self._changed_quarter(mutate)
        self.assertIsNone(r["calculation"]["value"])
        self.assertIn("conflicting", r["inputs"]["current_interest"]["detail"])

    def test_missing_and_ambiguous_statement_are_not_substituted(self):
        for mutation in [lambda root:root.append(deepcopy(li.statement_table(root))),
                         lambda root:li.statement_table(root).clear()]:
            r = self._changed_quarter(mutation)
            self.assertIsNone(r["calculation"]["operating_income"])
            self.assertEqual(r["calculation"]["interest"], 1079e6)

    def test_missing_document_and_stale_saved_review_are_unavailable(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / "missing.htm"
            self.assertIsNone(li.build(ANNUAL, p)["calculation"]["value"])
            with patch.object(li, "SAMPLE", p):
                self.assertIsNone(li.load())
                r = deepcopy(self.review);r["inputs"]["fy_interest"]["value"] += 1;r["calculation"] = li.calculate(r["inputs"])
                p.write_text(json.dumps(r));self.assertIsNone(li.load())

    def test_frozen_review_matches_real_fixture_and_full_sources_when_present(self):
        self.assertEqual(li.load(), self.review)
        raw = ROOT.parent / "data/raw/sec/LLY"
        if (raw / "lly-20251231.htm").exists():
            self.assertEqual(li.build(raw / "lly-20251231.htm", raw / "lly-20260630.htm"), self.review)
        s = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"][0]
        self.assertIsNotNone(li.for_snapshot(s))
        s["metrics"]["cash"]["period_end"] = "2026-09-30";self.assertIsNone(li.for_snapshot(s))
        s["ticker"] = "VST";self.assertIsNone(li.for_snapshot(s))

    def test_workbook_reconciles_and_is_bound_to_the_coverage_source(self):
        from openpyxl import load_workbook
        import company_review
        snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]
        b = company_review.saved_workbook(snapshots, ticker="LLY");self.assertIsNotNone(b)
        v, f = load_workbook(BytesIO(b), data_only=True), load_workbook(BytesIO(b), data_only=False)
        for cell, value in {"B40":26302e6,"C40":17893e6,"D40":10562e6,"B44":0,"C44":0,"D44":0,"B48":33633e6,"B49":1079e6,"B50":33633/1079,"B15":33633/1079}.items():
            self.assertAlmostEqual(v["Debt and liquidity"][cell].value, value)
            self.assertTrue(f["Debt and liquidity"][cell].value.startswith("="))
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for name in ["liquidity_pilots.json","lly_liquidity.json","lly_financing.json","lly_interest.json","lly_liquidity.xlsx","lly_liquidity.manifest.json"]:
                (root / name).write_bytes((ROOT / "samples" / name).read_bytes())
            self.assertIsNotNone(company_review.saved_workbook(snapshots, root, ticker="LLY"))
            (root / "lly_interest.json").write_text("{}")
            self.assertIsNone(company_review.saved_workbook(snapshots, root, ticker="LLY"))


if __name__ == "__main__":
    unittest.main()
