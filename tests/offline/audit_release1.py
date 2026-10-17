"""Read-only Release 1 reconciliation; optional actual downloads and saved raw documents."""
import argparse
import hashlib
from io import BytesIO
import importlib
import json
import math
from pathlib import Path
import sys
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
import borrowing_leases
import cash_generation
import company_history
import company_review
import core_charts
import export_guide
import funding_conclusions
import oracle_liquidity
import ratio_review
import vistra_liquidity
from openpyxl import load_workbook

STEMS = {"LLY": "lly", "ORCL": "oracle", "VST": "vistra"}
CIKS = {"LLY": "59478", "ORCL": "1341439", "VST": "1692819"}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def source_url(url, accession, ticker):
    parsed = urlparse(url)
    prefix = f"/Archives/edgar/data/{CIKS[ticker]}/{accession.replace('-', '')}/"
    require(parsed.scheme == "https" and parsed.netloc == "www.sec.gov" and
            parsed.path.startswith(prefix) and len(parsed.path) > len(prefix),
            f"{ticker}: source URL does not match filing accession/entity")


def reported_facts(value):
    if isinstance(value, dict):
        if value.get("tag") and value.get("accession") and value.get("value") is not None:
            yield value
        for child in value.values():
            yield from reported_facts(child)
    elif isinstance(value, (list, tuple)):
        for child in value:
            yield from reported_facts(child)


def audit(downloads=None, raw=None):
    snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]
    report = []
    for snapshot in snapshots:
        t = snapshot["ticker"]
        ratios, history = ratio_review.build(snapshot), company_history.build(snapshot)
        leases, charts = borrowing_leases.build(snapshot), core_charts.build(snapshot)
        conclusion = funding_conclusions.load(t)
        require(conclusion is not None, t + ": checked conclusion unavailable")
        interest = importlib.import_module(ratio_review.INTEREST_MODULES[t]).load()
        require(interest is not None, t + ": checked interest unavailable")
        content = (company_review.saved_workbook(snapshots, ticker=t) if t == "LLY" else
                   oracle_liquidity.workbook_bytes(oracle_liquidity.for_snapshot(snapshot)) if t == "ORCL" else
                   vistra_liquidity.workbook_bytes(vistra_liquidity.for_snapshot(snapshot)))
        require(content is not None, t + ": workbook failed its source/interpretation/layout gates")
        if downloads is not None:
            require(downloads.get(t) == content, t + ": actual download differs from reviewed bytes")
        values = load_workbook(BytesIO(content), data_only=True)
        formulas = load_workbook(BytesIO(content), data_only=False)
        try:
            require(values.sheetnames == [name for name, _ in export_guide.contents(t)], t + ": worksheet scope differs")
            require(values["Funding conclusion"]["B6"].value == conclusion["conclusion"], t + ": conclusion differs")
            require("Unavailable" in values["Funding conclusion"]["B17"].value, t + ": unsupported future coverage became a number")
            for i, row in enumerate(ratios["rows"], 6):
                actual = values["Debt and liquidity"].cell(i, 2).value
                require(actual == row["value"], t + ": ratio/metric mismatch " + row["key"])
                require(formulas["Debt and liquidity"].cell(i, 2).data_type == "f", t + ": result lost its formula")
            for i, row in enumerate(history["rows"], 6):
                for col, expected in [(2, row["annual_value"]), (3, row["value"])]:
                    actual = values["History"].cell(i, col).value
                    require(actual == ("Unavailable" if expected is None else expected), t + ": history mismatch " + row["key"])
            require(len(formulas["Core charts"]._charts) == 2, t + ": missing core chart")
            for i, (_, expected) in enumerate(charts["flow"], 6):
                require(values["Core charts"].cell(i, 2).value == expected, t + ": chart flow mismatch")
            require(sum(f["value"] for f in leases["borrowing"]) == snapshot["metrics"]["gross_debt"]["value"], t + ": borrowing does not reconcile")
            if t == "VST":
                require(charts["series"] == ["Cash and equivalents"] and
                        values["History"]["B7"].value == "Unavailable", "VST: unsupported older debt plotted or calculated")
            facts = {}
            for f in reported_facts([snapshot["metrics"], history["annual_facts"], leases, interest["inputs"]]):
                require(f.get("unit") in ("USD", "iso4217:USD") and math.isfinite(f["value"]), t + ": invalid fact unit/value")
                require(all(f.get(k) for k in ("period_end", "filed", "form", "tag", "selection_rule")), t + ": incomplete fact provenance")
                url = f.get("source_url")
                if not url and f["accession"] == snapshot["metrics"]["cash"]["accession"]:
                    url = snapshot["source_url"]
                source_url(url or "", f["accession"], t)
                facts[(f["tag"], f.get("period_start"), f["period_end"], f["accession"], str(f.get("dimensions", [])))] = f
            for _, _, sources in [*conclusion["rows"], ("Historical coverage", None, conclusion["coverage_sources"])]:
                for source in sources:
                    source_url(source["url"], source["accession"], t)
            raw_checked = 0
            if raw is not None:
                # Re-read exact annual instant contexts and their source hashes, without rewriting evidence.
                import filing_xbrl
                for f in history["annual_facts"]:
                    document = Path(f["source_url"]).name
                    path = raw / t / document
                    require(path.is_file(), t + ": required raw filing missing: " + document)
                    require(hashlib.sha256(path.read_bytes()).hexdigest() == f["source_sha256"], t + ": raw filing changed")
                    actual = filing_xbrl.read_instant(path, f["accession"], f["tag"], f["period_end"])
                    require(actual["value"] == f["value"] and actual["context_id"] == f["context_id"], t + ": raw annual fact differs")
                    raw_checked += 1
                spec = cash_generation.SOURCES[t]
                flows = cash_generation.build(t, {spec[r][0]: raw / t / spec[r][1] for r in ("annual", "quarter")})
                for key in ("operating_cash_flow", "capital_spending", "free_cash_flow"):
                    require(flows[key]["value"] == snapshot["metrics"][key]["value"], t + ": full raw cash flows differ")
            report.append({"ticker": t, "balance_sheet": ratios["period_end"], "annual_date": history["annual_end"],
                           "metrics": len(ratios["rows"]), "reported_facts": len(facts), "sheets": len(values.sheetnames),
                           "core_charts": 2, "raw_annual_facts": raw_checked,
                           "workbook_sha256": hashlib.sha256(content).hexdigest(), "future_funding_coverage": "Unavailable"})
        finally:
            values.close()
            formulas.close()
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--download-dir", type=Path, help="Directory containing actual LLY.xlsx, ORCL.xlsx, VST.xlsx downloads")
    parser.add_argument("--raw", type=Path, help="Saved data/raw/sec directory; require and re-read the checked annual/quarter filings")
    args = parser.parse_args()
    downloads = {t: (args.download_dir / (t + ".xlsx")).read_bytes() for t in STEMS} if args.download_dir else None
    print(json.dumps(audit(downloads, args.raw), indent=2))
