"""Read-only presentation records for Artifact Tool; never writes a workbook."""
import json
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook


def records(path):
    values = load_workbook(path, data_only=True)
    result = {}
    for s in values:
        cells = []
        for row in s:
            for c in row:
                v = c.value
                if isinstance(v, (date, datetime)):
                    v = v.strftime("%Y-%m-%d")
                color = c.font.color.rgb if c.font.color and c.font.color.type == "rgb" else "000000"
                if v is not None:
                    cells.append([c.coordinate, v, color[-6:], c.number_format, c.font.bold, c.font.italic])
        result[s.title] = dict(rows=s.max_row, cols=s.max_column, cells=cells,
                               heights={str(i): r.height for i, r in s.row_dimensions.items() if r.height},
                               widths={k: v.width for k, v in s.column_dimensions.items()})
    values.close()
    return result


if __name__ == "__main__":
    import sys
    print(json.dumps(records(Path(sys.argv[1]))))
