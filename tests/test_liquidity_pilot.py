"""Pilot figures reconcile to the hand-read evidence, and exports trace calculations to inputs."""
import html
import json
import sys
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import liquidity_pilot as pilot  # noqa: E402
import cash_generation  # noqa: E402


class LiquidityPilot(unittest.TestCase):
    def setUp(self):
        self.snapshots = json.loads(pilot.SAMPLE.read_text())["snapshots"]

    def test_three_saved_pilots_reconcile_and_rebuild(self):
        expected = {"LLY": (8950, 54908, 45958, 1.3546427598919168),
                    "ORCL": (36369, 125337, 88968, 1.1708128130656228),
                    "VST": (435, 19895, 19460, 0.9733825897508762)}
        for saved in self.snapshots:
            ticker = saved["ticker"]
            m = saved["metrics"]
            self.assertEqual(tuple(m[k]["value"] / 1_000_000 for k in ("cash", "gross_debt", "net_debt")),
                             expected[ticker][:3])
            self.assertAlmostEqual(m["current_ratio"]["value"], expected[ticker][3])
            facts = json.loads((ROOT / "tests/fixtures" / f"{ticker}_companyfacts.json").read_text())
            docs = {pilot.PILOTS[ticker]["accession"]: ROOT / "tests/fixtures/VST_2026Q2_ixbrl.htm"}
            flows = {cash_generation.SOURCES[ticker][role][0]: ROOT / f"tests/fixtures/{ticker}_{role}_cashflow.htm"
                     for role in ("annual", "quarter")}
            self.assertEqual(pilot.build_snapshot(facts, ticker, docs, flows), saved)
            self.assertIsNone(m["net_debt"]["tag"])

    def test_every_reported_input_keeps_source_metadata(self):
        for snapshot in self.snapshots:
            for metric in snapshot["metrics"].values():
                for fact in pilot.leaf_inputs(metric):
                    for key in ("period_end", "form", "filed", "accession", "tag", "unit", "selection_rule"):
                        self.assertTrue(fact[key], (snapshot["ticker"], key))
            if snapshot["ticker"] == "VST":
                extension = snapshot["metrics"]["gross_debt"]["inputs"][-1]
                self.assertEqual((extension["context_id"], extension["scale"]), ("c-8", 6))

    def test_net_debt_refuses_missing_or_incompatible_inputs(self):
        m = self.snapshots[0]["metrics"]
        cash, debt = m["cash"], m["gross_debt"]
        for key, bad in (("value", None), ("period_end", "2026-06-29"),
                         ("accession", "another-version"), ("unit", "EUR")):
            changed = {**cash, key: bad}
            self.assertIsNone(pilot.net_debt(changed, debt)["value"], key)

    def test_negative_net_debt_is_kept(self):
        m = self.snapshots[0]["metrics"]
        cash = {**m["cash"], "value": m["gross_debt"]["value"] + 100}
        self.assertEqual(pilot.net_debt(cash, m["gross_debt"])["value"], -100)

    def test_wrong_company_facts_are_refused(self):
        with self.assertRaisesRegex(ValueError, "CIK does not match"):
            pilot.build_snapshot({"cik": 1}, "LLY")

    def test_display_has_dates_source_links_and_caveats(self):
        rendered = html.unescape(pilot.render_html(self.snapshots[1]))
        self.assertIn("$125.34B", rendered)
        self.assertIn("2026-08-31", rendered)
        self.assertIn("0001193125-26-389274", rendered)
        self.assertIn("commercial paper", rendered)
        self.assertIn("Source 10-Q", rendered)
        self.assertIn("excludes leases", rendered)
        self.assertIn("$288B", rendered)
        self.assertIn("$11.4B of customer prepayments", rendered)
        self.assertIn("−$28.72B", rendered)

    def test_excel_stores_dollars_and_calculates_from_input_cells(self):
        wb = load_workbook(BytesIO(pilot.workbook(self.snapshots)))
        ws, raw = wb["Debt and liquidity"], wb["Liquidity inputs"]
        self.assertEqual(raw["D2"].value, 8_950_000_000)
        self.assertEqual(ws["C2"].value, "='Liquidity inputs'!D2")
        self.assertEqual(ws["C3"].value, "='Liquidity inputs'!D3+'Liquidity inputs'!D4")
        self.assertEqual(ws["C4"].value, "=C3-C2")
        self.assertEqual(ws["C5"].value, "='Liquidity inputs'!D5/'Liquidity inputs'!D6")
        self.assertEqual(raw["H2"].value, "0000059478-26-000081")
        self.assertEqual(raw["I2"].value, "us-gaap:CashAndCashEquivalentsAtCarryingValue")
        self.assertEqual(ws["C6"].value, "='Liquidity inputs'!D7+'Liquidity inputs'!D8-'Liquidity inputs'!D9")
        self.assertEqual(ws["C8"].value, "=C6-C7")
        self.assertEqual(ws["C9"].value, "=C6/C3")
        self.assertEqual(ws["C10"].value, "=C8/C3")
        self.assertEqual(raw["H7"].value, "0000059478-26-000013")
        self.assertIn("lly-20251231.htm", raw["K7"].value)
        self.assertIn("Other obligations", wb.sheetnames)
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertEqual(ws["C2"].font.name, "Arial")
        self.assertIn("About", wb.sheetnames)

    def test_unavailable_export_never_writes_zero_or_formula(self):
        snapshot = deepcopy(self.snapshots[0])
        snapshot["metrics"]["cash"].update(value=None, detail="Source not checked")
        snapshot["metrics"]["net_debt"] = pilot.net_debt(snapshot["metrics"]["cash"],
                                                       snapshot["metrics"]["gross_debt"])
        wb = load_workbook(BytesIO(pilot.workbook([snapshot])))
        self.assertIsNone(wb["Debt and liquidity"]["C2"].value)
        self.assertIsNone(wb["Debt and liquidity"]["C4"].value)
        self.assertEqual(wb["Debt and liquidity"]["G4"].value, "Unavailable")

    def test_styled_saved_export_matches_all_three_snapshots(self):
        import company_review
        content = company_review.saved_workbook(self.snapshots)
        self.assertIsNotNone(content)
        wb = load_workbook(BytesIO(content), data_only=True)
        row = 2
        for snapshot in self.snapshots:
            for metric in snapshot["metrics"].values():
                self.assertAlmostEqual(wb["Debt and liquidity"].cell(row, 3).value, metric["value"])
                row += 1
        changed = deepcopy(self.snapshots)
        changed[0]["metrics"]["cash"]["value"] = None
        self.assertIsNone(company_review.saved_workbook(changed))

    def test_lilly_download_contains_only_lilly_metrics_with_source_formulas(self):
        import company_review
        content=company_review.saved_workbook(self.snapshots,ticker="LLY")
        self.assertIsNotNone(content)
        values=load_workbook(BytesIO(content),data_only=True)
        formulas=load_workbook(BytesIO(content),data_only=False)
        self.assertEqual(values.sheetnames,["Funding conclusion","Borrowing and leases","Debt and liquidity","History","Repayment review","Data","About"])
        for row,metric in enumerate(self.snapshots[0]["metrics"].values(),6):
            self.assertAlmostEqual(values["Debt and liquidity"].cell(row,2).value,metric["value"])
            self.assertTrue(formulas["Debt and liquidity"].cell(row,2).value.startswith("="))
        inputs=values["Data"]
        for row in range(2,inputs.max_row+1):
            if inputs.cell(row,1).value is None:
                break
            self.assertTrue(str(inputs.cell(row,7).value).startswith("0000059478-"))
        changed=deepcopy(self.snapshots);changed[0]["metrics"]["cash"]["value"]=None
        self.assertIsNone(company_review.saved_workbook(changed,ticker="LLY"))
        self.assertIsNone(company_review.saved_workbook(self.snapshots,ticker="ORCL"))

    def test_styled_export_refuses_changed_source_or_file(self):
        import company_review
        source = ROOT / "samples"
        names = ("liquidity_pilots.json", "oracle_financing.json", "liquidity_pilots.xlsx",
                 "liquidity_pilots.manifest.json")
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            for name in names:
                (folder / name).write_bytes((source / name).read_bytes())
            self.assertIsNotNone(company_review.saved_workbook(self.snapshots, folder))
            financing = folder / "oracle_financing.json"
            original = financing.read_bytes()
            financing.write_bytes(original + b"\n")
            self.assertIsNone(company_review.saved_workbook(self.snapshots, folder))
            financing.write_bytes(original)
            (folder / "liquidity_pilots.xlsx").write_bytes(b"not the checked workbook")
            self.assertIsNone(company_review.saved_workbook(self.snapshots, folder))

    def test_flow_ratios_refuse_incompatible_dates_missing_or_nonpositive_debt(self):
        m = self.snapshots[1]["metrics"]
        flow, debt = m["free_cash_flow"], m["gross_debt"]
        self.assertLess(pilot.flow_to_debt(flow, debt)["value"], 0)
        for changes in ({"period_end": "2026-06-30"}, {"value": None}, {"value": 0},
                        {"value": -1}, {"unit": "EUR"}):
            self.assertIsNone(pilot.flow_to_debt(flow, {**debt, **changes})["value"])


if __name__ == "__main__":
    unittest.main()
