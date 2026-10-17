"""Source freshness, snapshot isolation and app/export agreement for funding prose."""
from copy import deepcopy
from io import BytesIO
import html
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

import funding_conclusions as fc
import company_review
import oracle_liquidity as ol
import vistra_liquidity as vl


class FundingConclusions(unittest.TestCase):
    def setUp(self):
        self.snapshots = json.loads((fc.FOLDER / 'liquidity_pilots.json').read_text())['snapshots']

    def test_checked_models_and_coverage(self):
        for ticker,expected in [('LLY',33633/1079),('ORCL',23057/5104),('VST',3563/1117)]:
            r=fc.load(ticker)
            self.assertEqual(r['status'],'Checked')
            self.assertAlmostEqual(r['coverage'],expected)
            self.assertEqual(len(r['rows']),4)
            self.assertIn('Future funding coverage: Unavailable',r['rows'][3][1])
            self.assertIn('do not rank funding strength',r['comparison'])

    def test_every_bound_source_change_or_missing_file_refuses_its_conclusion(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            for name in fc.CHECKED:shutil.copy(fc.FOLDER/name,folder/name)
            for ticker,names in fc.FILES.items():
                for name in names:
                    original=(folder/name).read_bytes()
                    (folder/name).write_bytes(original+b'\n')
                    self.assertIsNone(fc.load(ticker,folder),(ticker,name))
                    (folder/name).unlink()
                    self.assertIsNone(fc.load(ticker,folder),(ticker,name))
                    (folder/name).write_bytes(original)
                self.assertIsNotNone(fc.load(ticker,folder))

    def test_other_company_evidence_does_not_invalidate_checked_company(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            for name in fc.CHECKED:shutil.copy(fc.FOLDER/name,folder/name)
            (folder/'oracle_revolver.json').write_text('{}')
            self.assertIsNotNone(fc.load('LLY',folder))
            self.assertIsNotNone(fc.load('VST',folder))
            self.assertIsNone(fc.load('ORCL',folder))

    def test_changed_snapshot_and_unknown_ticker_refused(self):
        self.assertIsNone(fc.load('AAPL'))
        for s in self.snapshots:
            self.assertIsNotNone(fc.for_snapshot(s))
            for key in ['accession','period_end','value']:
                changed=deepcopy(s);changed['metrics']['cash'][key]='changed'
                self.assertIsNone(fc.for_snapshot(changed))
                self.assertIn('Unavailable',fc.render(changed))

    def test_returned_model_cannot_mutate_following_review(self):
        r=fc.load('LLY');r['rows'][0][2][0]['label']='changed'
        self.assertNotEqual(fc.load('LLY')['rows'][0][2][0]['label'],'changed')

    def test_app_shows_same_claims_and_clickable_sources(self):
        for s in self.snapshots:
            r=fc.load(s['ticker']);shown=html.unescape(company_review.render(s))
            self.assertIn(r['conclusion'],shown)
            self.assertLess(shown.index('Funding conclusion'),shown.index('What remained after capital spending?'))
            for _,text,sources in r['rows']:
                self.assertIn(text,shown)
                for source in sources:self.assertIn(source['url'],shown)
            self.assertIn(r['coverage_basis'],shown)
            self.assertIn(r['comparison'],shown)
            self.assertNotIn('$',fc.render(s))

    def test_workbooks_use_same_prose_and_link_existing_calculation(self):
        from openpyxl import load_workbook
        for s in self.snapshots:
            ticker=s['ticker'];r=fc.load(ticker)
            stem={'LLY':'lly','ORCL':'oracle','VST':'vistra'}[ticker]
            content=(fc.FOLDER/f'{stem}_liquidity.xlsx').read_bytes()
            v=load_workbook(BytesIO(content),data_only=True)
            f=load_workbook(BytesIO(content),data_only=False)
            self.assertEqual(v.sheetnames[0],'Funding conclusion')
            sheet=v['Funding conclusion']
            self.assertEqual(sheet['B6'].value,r['conclusion'])
            for row,(_,text,_) in enumerate(r['rows'],7):self.assertEqual(sheet[f'B{row}'].value,text)
            self.assertAlmostEqual(sheet['B13'].value,r['coverage'])
            self.assertIn(r['coverage_formula'][1:],f['Funding conclusion']['B13'].value)
            self.assertIn('ISNUMBER',f['Funding conclusion']['B13'].value)
            self.assertEqual(sheet['B14'].value,r['coverage_basis'])
            self.assertEqual(sheet['B15'].value,r['comparison'])
            provenance=' '.join(str(c.value) for row in v['Data'] for c in row if c.value)
            for _,_,sources in r['rows']:
                for source in sources:self.assertIn(source['accession'],provenance)

    def test_oracle_conclusion_states_the_checked_principal_calendar(self):
        """The re-reviewed Oracle text quotes the calendar's figures and keeps coverage Unavailable."""
        import oracle_principal
        r, cal = fc.load('ORCL'), oracle_principal.load()
        self.assertEqual(r['reviewed'], '2026-10-17')
        self.assertIn('oracle_principal.json', r['evidence_files'])
        near = cal['next_twelve_months']
        tl = cal['term_loan']
        installments = sum(s['amount'] for s in tl['installments'] if '2026-08-31' < s['date'] <= near['end'])
        for amount in (near['principal'], tl['final']['amount']):
            self.assertIn(f"${amount / 1e9:.3f}B", r['conclusion'] + r['rows'][0][1])
        self.assertIn(f"${installments / 1e6:,.0f}M of term-loan installments", r['rows'][0][1])
        self.assertIn('Coverage remains Unavailable', r['conclusion'])
        self.assertIn('Future funding coverage: Unavailable', r['rows'][3][1])
        self.assertNotIn('remaining debt principal dates', r['rows'][3][1])
        self.assertEqual({fc.REVIEWED['LLY'], fc.REVIEWED['VST']}, {'2026-10-16'})

    def test_a_changed_principal_calendar_refuses_the_oracle_conclusion(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            for name in fc.FILES['ORCL']:
                shutil.copy(fc.FOLDER / name, folder / name)
            self.assertIsNotNone(fc.load('ORCL', folder))
            (folder / 'oracle_principal.json').write_bytes((fc.FOLDER / 'oracle_principal.json').read_bytes() + b' ')
            self.assertIsNone(fc.load('ORCL', folder))

    def test_changed_interpretation_refuses_all_individual_downloads(self):
        with patch.object(fc,'fingerprint',return_value='changed'):
            self.assertIsNone(company_review.saved_workbook(self.snapshots,ticker='LLY'))
            self.assertIsNone(ol.workbook_bytes(json.loads(ol.SAMPLE.read_text())))
            self.assertIsNone(vl.workbook_bytes(json.loads(vl.SAMPLE.read_text())))

    def test_unavailable_evidence_refuses_all_individual_downloads(self):
        with patch.object(fc,'load',return_value=None):
            self.assertIsNone(company_review.saved_workbook(self.snapshots,ticker='LLY'))
            self.assertIsNone(ol.workbook_bytes(json.loads(ol.SAMPLE.read_text())))
            self.assertIsNone(vl.workbook_bytes(json.loads(vl.SAMPLE.read_text())))


if __name__=='__main__':unittest.main()
