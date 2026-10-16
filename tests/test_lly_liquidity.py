"""Checks against real saved facts, plus failures that must not become funding forecasts."""
import hashlib
import json
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET

import filing_xbrl as fx
import lly_liquidity as ll

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests/fixtures"
ANNUAL, QUARTER = FIX / "LLY_2025_liquidity.htm", FIX / "LLY_2026H1_liquidity.htm"


class LillyReviewTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.review = ll.build(ANNUAL, QUARTER, text_checked=True)

    def test_h1_cash_reconciles_without_reusing_ttm_or_debt_proceeds(self):
        r, v = self.review, {k: f["value"] for k, f in self.review["inputs"].items()}
        self.assertEqual(r["cash"], {"ocf_less_capex": 10_764_000_000, "acquisitions_iprd": 13_291_000_000,
                                    "payouts": 7_051_000_000, "net_borrowing": 12_602_000_000, "other": -1_342_000_000})
        c = r["cash"]
        self.assertEqual(v["opening"] + c["ocf_less_capex"] - c["acquisitions_iprd"] - c["payouts"] + c["net_borrowing"] + c["other"], v["closing"])
        self.assertEqual(v["closing"], 8_950_000_000)
        self.assertEqual(v["long_repayments"], 1_623_000_000)
        for key in ll.FLOWS:
            f = r["inputs"][key]
            self.assertEqual((f["period_start"], f["period_end"], f["accession"], f["dimensions"]),
                             (ll.START, ll.END, ll.SOURCES["quarter"]["accession"], []))

    def test_due_year_schedule_preserves_carrying_basis_and_printed_adjustment(self):
        r = self.review
        self.assertEqual(r["notes_schedule"]["total_notes"], 42_842_000_000)
        self.assertEqual(r["notes_schedule"]["total_debt"], 42_503_000_000)
        self.assertEqual(r["inputs"]["adjustment"]["value"], 339_000_000)
        self.assertEqual(r["notes_schedule"]["adjustment_coefficient"], -1)
        for label, member in ll.NOTE_MEMBERS:
            f = r["inputs"]["notes_" + member]
            self.assertEqual(f["period_end"], ll.ANNUAL)
            self.assertEqual(f["dimensions"], [["us-gaap:DebtInstrumentAxis", "lly:NotesDue" + member + "Member"]])
        self.assertIsNone(r["coverage"])

    def test_text_amounts_keep_basis_and_do_not_become_inline_facts(self):
        d = {x["key"]: x for x in self.review["disclosures"]}
        self.assertEqual(d["unpaid_kelonia"]["value"] + d["unpaid_centessa"]["value"], 296_000_000)
        self.assertEqual(d["manufacturing"]["period_end"], ll.ANNUAL)
        self.assertEqual(d["indiana"]["accession"], "0000059478-26-000077")
        self.assertIn("Includes upfront", d["july_ceiling"]["detail"])
        self.assertEqual(d["atai"]["period_end"], "Disclosed August 5, 2026")
        self.assertTrue(all(f["tag"] is None and f["source_kind"] == "Reviewed filing text" and f["sha256"] for f in d.values()))
        self.assertEqual(ll.build(ANNUAL, QUARTER)["disclosures"], [])

    def test_contingent_fair_value_is_separate_from_payments_and_maxima(self):
        fs = [self.review["inputs"][k] for k in ["contingent_current", "contingent_noncurrent"]]
        self.assertEqual([f["value"] for f in fs], [814_000_000, 1_704_000_000])
        self.assertTrue(all(f["dimensions"] == [["us-gaap:FairValueByFairValueHierarchyLevelAxis", "us-gaap:FairValueInputsLevel3Member"]] for f in fs))

    def _changed(self, mutate):
        root = ET.parse(QUARTER).getroot()
        mutate(root)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "changed.htm"
            ET.ElementTree(root).write(path)
            return ll.build(ANNUAL, path, text_checked=False)

    def test_missing_cash_input_disables_bridge_not_historical_notes(self):
        def mutate(root):
            for parent in root.iter():
                for fact in list(parent):
                    if fact.get("name") == "us-gaap:OtherPaymentsToAcquireBusinesses":
                        parent.remove(fact)
        r = self._changed(mutate)
        self.assertIsNone(r["cash"])
        self.assertIn("acquisitions", r["cash_detail"])
        self.assertIsNotNone(r["notes_schedule"])
        self.assertIsNone(r["coverage"])

    def test_wrong_period_cannot_substitute_prior_year_ocf(self):
        def mutate(root):
            for c in root.iter(f"{{{fx.XBRLI}}}context"):
                if c.get("id") == "c-1":
                    c.find("xbrli:period/xbrli:endDate", fx.NS).text = "2026-06-29"
        r = self._changed(mutate)
        self.assertIsNone(r["inputs"]["ocf"]["value"])
        self.assertIsNone(r["cash"])

    def test_cash_control_mismatch_blocks_derived_results(self):
        def mutate(root):
            for f in root.iter(f"{{{fx.IX}}}nonFraction"):
                if f.get("name") == "us-gaap:NetCashProvidedByUsedInInvestingActivities" and f.get("contextRef") == "c-1":
                    f.text = "19,311"
        r = self._changed(mutate)
        self.assertIsNone(r["cash"])
        self.assertIn("do not reconcile", r["cash_detail"])

    def test_wrong_entity_cannot_supply_funding(self):
        def mutate(root):
            for ident in root.iter(f"{{{fx.XBRLI}}}identifier"):
                ident.text = "0000000001"
        r = self._changed(mutate)
        self.assertIsNone(r["cash"])
        self.assertIn("not Lilly", r["inputs"]["closing"]["detail"])

    def test_frozen_model_matches_fixture_values_and_full_sources_when_present(self):
        saved = ll.load()
        self.assertEqual(saved["cash"], self.review["cash"])
        self.assertEqual(saved["notes_schedule"], self.review["notes_schedule"])
        self.assertEqual(saved["disclosures"], self.review["disclosures"])
        raw = ROOT.parent / "data/raw/sec/LLY"
        if ll.audit_sources(raw):
            full = ll.build(raw / ll.SOURCES["annual"]["document"], raw / ll.SOURCES["quarter"]["document"], text_checked=True)
            self.assertEqual(full, saved)

    def test_changed_source_needs_new_prose_review(self):
        with tempfile.TemporaryDirectory() as folder:
            self.assertFalse(ll.audit_sources(Path(folder)))

    def test_review_refuses_incompatible_snapshot_and_other_companies(self):
        snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]
        self.assertIsNotNone(ll.for_snapshot(snapshots[0]))
        self.assertIsNone(ll.for_snapshot(snapshots[1]))
        s = deepcopy(snapshots[0]);s["metrics"]["cash"]["value"] += 1
        self.assertIsNone(ll.for_snapshot(s))

    def test_workbook_shared_review_values_formulas_and_manifest(self):
        from openpyxl import load_workbook
        path = ROOT / "samples/lly_liquidity.xlsx"
        values, formulas = load_workbook(path, data_only=True), load_workbook(path, data_only=False)
        self.assertEqual(values.sheetnames, ["Debt and liquidity", "Repayment review", "Data", "About"])
        for cell, value in {"B22": 0, "B23": 10_764_000_000, "B24": 12_602_000_000, "B25": 13_291_000_000,
                            "B26": 7_051_000_000, "B48": 2_518_000_000, "B63": 42_842_000_000, "B68": 0}.items():
            self.assertEqual(values["Repayment review"][cell].value, value)
            self.assertTrue(formulas["Repayment review"][cell].value.startswith("="))
        self.assertEqual(values["Repayment review"]["B49"].value, "Unavailable")
        self.assertEqual(values["Repayment review"]["B34"].value, "Unavailable")
        self.assertEqual(len(formulas["Repayment review"]._charts), 1)
        manifest = json.loads(path.with_suffix(".manifest.json").read_text())
        self.assertEqual(set(manifest["sources"]), {"liquidity_pilots.json", "lly_liquidity.json"})
        self.assertEqual(manifest["workbook_sha256"], hashlib.sha256(path.read_bytes()).hexdigest())


if __name__ == "__main__":
    unittest.main()
