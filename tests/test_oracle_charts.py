"""Charts must show the checked calendar and the selected run, never missing values as zero."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import unittest
import numpy as np
from openpyxl import load_workbook
import liquidity_simulation as ls
import oracle_principal as op
import oracle_principal_view as ov
import company_review
import scenario_export as se
import simulation_view as sv

ROOT=Path(__file__).resolve().parents[1]


class OracleCharts(unittest.TestCase):
    def test_ladder_is_the_complete_checked_calendar(self):
        review=op.load();text=ov.ladder(review,company_review.money)
        for b in review['calendar']:
            self.assertIn(f'data-principal="{b["principal"]}"',text)
        self.assertIn('FY2032 onward (combined)',text)
        self.assertIn('September 2026–May 2027',text)
        self.assertIn('commercial paper (Unavailable)',text)
        self.assertAlmostEqual(sum(b['principal'] for b in review['calendar']),review['reconciliation']['known_principal_excluding_cp'])
        bad=deepcopy(review);bad['calendar'][0]['principal']=None
        with self.assertRaises(ValueError):ov.ladder(bad,company_review.money)

    def test_monthly_band_is_independently_computed_from_the_trials(self):
        i=ls.oracle_inputs();d=ls.oracle_decisions()[1];p=ls.monthly_path(i,d,1000,42)
        for q in (10,50,90):np.testing.assert_array_equal(p[f'month_p{q}'],np.percentile(p['path'],q,axis=0))
        f=sv.monthly_figure(p)
        np.testing.assert_array_equal(f.data[2].y,p['month_p50']/ls.B)
        self.assertEqual(list(f.data[0].x),list(ls.MONTH_ENDS))
        self.assertEqual(f.data[1].fill,'tonexty')
        self.assertEqual(f.layout.shapes[0].y0,0)

    def test_scenario_chart_follows_the_selected_decision_and_seed(self):
        i=ls.oracle_inputs();ds=ls.oracle_decisions();rs=[(ls.simulate(i,d,1000,42),d) for d in ds]
        monthly=[(d,r,ls.monthly_path(i,d,1000,42)) for r,d in rs]
        raw=se.workbook(i,rs,1000,42,ls.compare_correlations(i,ds[2],1000,42),monthly=monthly)
        w=load_workbook(BytesIO(raw));chart=w['Monthly path by month']._charts[0]
        self.assertIn('$F$26:$F$37',chart.series[0].val.numRef.f)
        for series,field in zip(chart.series[:3],('month_p10','month_p50','month_p90')):
            np.testing.assert_allclose([float(p.v) for p in series.val.numRef.numCache.pt],monthly[2][2][field])
        w.close()

    def test_saved_oracle_ladder_links_to_the_existing_formula_table(self):
        values=load_workbook(ROOT/'samples/oracle_liquidity.xlsx',data_only=True)
        chart=values['Payment dates']._charts[0]
        self.assertEqual(chart.series[0].val.numRef.f,"'Payment dates'!$C$49:$C$54")
        cached=[float(p.v) for p in chart.series[0].val.numRef.numCache.pt]
        np.testing.assert_allclose(cached,[values['Payment dates'].cell(r,3).value for r in range(49,55)])
        self.assertEqual(len(values['Core charts']._charts),2)
        values.close()
