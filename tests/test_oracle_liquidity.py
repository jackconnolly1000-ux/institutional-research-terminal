"""Oracle cash-basis, funding exclusions, exact dates and payment reconciliations."""
import html
import json
import tempfile
import unittest
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

import oracle_liquidity as ol
import oracle_liquidity_view as view
from company_review import money

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests/fixtures/ORCL_2026Q1_liquidity.htm"
RAW = ROOT.parent / "data/raw/sec/ORCL"
M = 1_000_000


class OracleLiquidity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.saved = json.loads(ol.SAMPLE.read_text())
        cls.snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]

    def changed(self, change):
        root = ET.parse(FIXTURE).getroot()
        change(root)
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / "changed.htm"
            ET.ElementTree(root).write(p, encoding="utf-8", xml_declaration=True)
            return ol.build(p)

    def tag_change(self, key, attribute, value):
        tag = self.saved["inputs"][key]["tag"]
        def mutate(root):
            for fact in root.iter(f"{{{ol.fx.IX}}}nonFraction"):
                if fact.get("name") == tag:
                    if attribute == "text":
                        fact.text = value
                    else:
                        fact.set(attribute, value)
        return self.changed(mutate)

    def test_saved_analysis_matches_real_filing_facts(self):
        self.assertEqual(ol.build(FIXTURE), self.saved)
        inputs = self.saved["inputs"]
        expected = {"ocf":23103,"prepay":11363,"capex":28499,"equity":19909,"debt_repayments":4202,
                    "cp":0,"capex_financing":-830,"dividends":1565,"unpaid_capex":6247,
                    "closing_total":38934,"closing_cash":36369}
        self.assertEqual({k:inputs[k]["value"] for k in expected}, {k:v*M for k,v in expected.items()})
        self.assertTrue(all(f["accession"] == ol.of.QUARTER_ACCESSION and f["unit"] == "iso4217:USD"
                            and not f["dimensions"] for f in inputs.values()))
        self.assertEqual(inputs["prepay"]["period_start"], "2026-06-01")
        self.assertEqual(inputs["prepay"]["scale"],6)

    def test_cash_reconciles_on_the_restricted_cash_basis(self):
        c = self.saved["cash"]
        self.assertEqual(c["restricted_at_end"],2565*M)
        self.assertEqual(c["other_net"],-2666*M)
        i = {k:f["value"] for k,f in self.saved["inputs"].items()}
        self.assertEqual(i["opening_total"]+c["ocf_less_prepay"]+i["prepay"]+i["equity"]
                         -i["capex"]-i["debt_repayments"]+c["other_net"],i["closing_total"])
        self.assertEqual(i["closing_total"]-c["restricted_at_end"],i["closing_cash"])

    def test_historical_exclusions_remove_each_contribution_once(self):
        c=self.saved["cash"]
        self.assertEqual((c["ocf_less_prepay"],c["fcf"],c["fcf_less_prepay"]), (11740*M,-5396*M,-16759*M))
        self.assertEqual((c["change_without_equity"],c["change_without_equity_and_prepay"]),(-12264*M,-23627*M))
        self.assertAlmostEqual(c["prepay_share_ocf"],11363/23103)

    def test_missing_fact_wrong_unit_and_wrong_entity_are_not_zero(self):
        for attr, val in (("name","x:missing"),("unitRef","EUR")):
            with self.subTest(attr=attr):
                got=self.tag_change("prepay",attr,val)
                self.assertIsNone(got["cash"])
                self.assertEqual(got["status"],"Unavailable")
                self.assertIsNone(got["inputs"]["prepay"]["value"])
                self.assertEqual(got["schedules"]["finance"]["status"],"Checked")
        def mutate(root):
            for f in root.iter(f"{{{ol.fx.XBRLI}}}identifier"): f.text="1"
        got=self.changed(mutate)
        self.assertIsNone(got["cash"])
        self.assertIn("CIK",got["detail"])

    def test_wrong_period_or_negative_cash_use_is_refused(self):
        def mutate(root):
            for f in root.iter(f"{{{ol.fx.XBRLI}}}startDate"):
                if f.text==ol.START: f.text="2026-06-02"
        self.assertIsNone(self.changed(mutate)["cash"])
        got=self.tag_change("capex","sign","-")
        self.assertIsNone(got["cash"])
        self.assertIn("direction",got["detail"])

    def test_detail_must_tie_to_independent_cash_flow_controls(self):
        for key in ("equity","capex_financing","dividends","closing_total"):
            with self.subTest(key=key):
                got=self.tag_change(key,"text","1")
                self.assertIsNone(got["cash"])
                self.assertIn("reconcil",got["detail"])

    def test_payment_dates_and_liabilities_reconcile_without_borrowing(self):
        ss=self.saved["schedules"]
        self.assertTrue(all(s["status"]=="Checked" for s in ss.values()))
        self.assertEqual([ss[k]["rows"][0]["value"] for k in ol.SCHEDULES],[3219*M,620*M,5449*M])
        for s in ss.values():
            self.assertEqual((s["rows"][0]["due_start"],s["rows"][0]["due_end"]),("2026-09-01","2027-05-31"))
            self.assertEqual(s["rows"][1]["due_start"],"2027-06-01")
        self.assertEqual(self.saved["inputs"]["not_commenced"]["value"],288000*M)
        self.assertIsNone(self.saved["facility"]["available"])
        self.assertIsNone(self.saved["facility"]["headroom"])

    def test_bad_payment_total_blocks_only_its_schedule(self):
        for key in ("operating_0","operating_interest"):
            got=self.tag_change(key,"text","1")
            self.assertEqual(got["schedules"]["operating"]["status"],"Unavailable")
            self.assertEqual(got["schedules"]["finance"]["status"],"Checked")
            self.assertEqual(got["status"],"Checked")

    def test_snapshot_and_workbook_must_match_checked_sources(self):
        for index in (0,2): self.assertIsNone(ol.for_snapshot(self.snapshots[index]))
        bad=deepcopy(self.snapshots[1]);bad["metrics"]["cash"]["accession"]="another"
        self.assertIsNone(ol.for_snapshot(bad))
        bad=deepcopy(self.snapshots[1]);bad["metrics"]["cash"]["value"]+=M
        self.assertEqual(ol.for_snapshot(bad)["status"],"Unavailable")
        self.assertIsNotNone(ol.workbook_bytes(self.saved))
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/"bad.xlsx";p.write_bytes(b"not the checked workbook")
            with patch.object(ol,"WORKBOOK",p): self.assertIsNone(ol.workbook_bytes(self.saved))

    def test_view_explains_limits_and_keeps_other_companies_unaffected(self):
        shown=html.unescape(view.render(self.snapshots[1],money))
        for text in ("−$5.396B","−$16.759B","$11.363B","$19.909B","$2.565B","$36.369B",
                     "not forecasts", "recurrence", "fully utilized", "3.0x", "Unavailable", "6.02(c)"):
            self.assertIn(text,shown)
        self.assertNotIn("6.02(e)",shown)
        self.assertEqual(view.render(self.snapshots[0],money),"")

    def test_workbook_refuses_a_changed_annual_maturity_source(self):
        with tempfile.TemporaryDirectory() as folder:
            changed=Path(folder)/"annual.json"
            changed.write_bytes(ol.of.SAMPLE.read_bytes()+b"\n")
            with patch.object(ol.of,"SAMPLE",changed):
                self.assertIsNone(ol.workbook_bytes(self.saved))

    def test_workbook_refuses_a_changed_revolver_audit(self):
        with tempfile.TemporaryDirectory() as folder:
            changed=Path(folder)/"audit.json"
            changed.write_bytes(ol.oracle_revolver.SAMPLE.read_bytes()+b"\n")
            with patch.object(ol.oracle_revolver,"SAMPLE",changed):
                self.assertIsNone(ol.workbook_bytes(self.saved))

    def test_markdown_has_one_leading_style_and_no_currency_math(self):
        import company_review
        fragment = view.render(self.snapshots[1],money)
        self.assertNotIn("<style>",fragment)
        self.assertNotIn("$",fragment)
        shown = company_review.render(self.snapshots[1])
        self.assertEqual(shown.count("<style>"),1)
        self.assertNotIn("$",shown)
        from bs4 import BeautifulSoup
        parsed = BeautifulSoup(shown,"html.parser")
        self.assertFalse(parsed.select('section style'))
        self.assertTrue(any(h.get_text() == 'What payments are actually dated?' for h in parsed.find_all('h3')))
        self.assertTrue(any('scheduled to be repaid' in h.get_text() for h in parsed.find_all('h3')))

    def test_excel_preserves_formulas_values_sources_dates_and_missing_capacity(self):
        content=ol.workbook_bytes(self.saved)
        formulas=load_workbook(BytesIO(content),data_only=False)
        values=load_workbook(BytesIO(content),data_only=True)
        self.assertEqual(formulas.sheetnames,["Repayment review","Cash analysis","Payment dates","Data","About"])
        self.assertEqual(formulas["Cash analysis"]["B12"].value,"=B8-B10")
        self.assertEqual(values["Cash analysis"]["B12"].value,-16759*M)
        self.assertEqual(values["Cash analysis"]["B19"].value,0)
        self.assertEqual(values["Cash analysis"]["B26"].value,-23627*M)
        self.assertIsNone(values["Payment dates"]["B26"].value)
        self.assertIn("Unavailable",values["Payment dates"]["D26"].value)
        self.assertEqual(values["Payment dates"]["E6"].value,datetime(2026,9,1))
        self.assertEqual(values["Payment dates"]["F6"].value,datetime(2027,5,31))
        self.assertEqual(len(formulas["Cash analysis"]._charts),1)
        data=values["Data"]
        keys={data.cell(i,1).value:i for i in range(2,data.max_row+1)}
        row=keys["prepay"]
        self.assertEqual(data.cell(row,2).value,11363*M)
        self.assertEqual(data.cell(row,8).value,ol.of.QUARTER_ACCESSION)
        self.assertEqual(data.cell(row,9).value,ol.PREPAY)
        self.assertEqual(data.cell(row,13).value,ol.of.QUARTER_SOURCE)
        self.assertEqual(data.cell(keys["facility_drawn"],8).value,ol.of.ACCESSION)
        self.assertIsNone(values["Repayment review"]["E16"].value)
        self.assertEqual(values["Repayment review"]["F16"].value,"Unavailable")
        self.assertEqual(values["Repayment review"]["C16"].value,3219*M)
        self.assertEqual(values["Repayment review"]["B44"].value,7210*M)
        repayment=values["Repayment review"]
        self.assertEqual(repayment["B60"].value,10000*M)
        self.assertEqual(repayment["C60"].value,0)
        self.assertEqual(repayment["D60"].value,10000*M)
        self.assertEqual(formulas["Repayment review"]["D60"].value,"=IF(COUNT(B60:C60)=2,B60-C60,NA())")
        for row in (59,60,61):
            self.assertEqual(repayment[f"E{row}"].value,"Unavailable")
        for cell in ("C59","D59","B61","C61","D61"):
            self.assertIsNone(repayment[cell].value)
        self.assertEqual(repayment["B68"].value,12000*M)
        self.assertIn("not a $12B commitment",repayment["D68"].value)
        shown=html.unescape(view.revolver_html(self.snapshots[1],money))
        for phrase in ("May 31 undrawn commitment is $10B", "not a $12B commitment", "not assumed zero", "September 14, 2026"):
            self.assertIn(phrase,shown)

    @unittest.skipUnless((RAW/"orcl-20260831.htm").exists(),"Full saved Oracle filing not available")
    def test_full_saved_filing_matches_trimmed_fixture_and_credit_terms(self):
        self.assertEqual(ol.build(RAW/"orcl-20260831.htm"),self.saved)
        from bs4 import BeautifulSoup
        credit=RAW/"orcl-ex10_17.htm"
        if credit.exists():
            text=BeautifulSoup(credit.read_text(),"html.parser").get_text(" ",strip=True)
            for phrase in ("3.0 to 1.0","no event has occurred and is continuing","maximum of two times",
                           "total cash interest expense","fifth anniversary of the Effective Date"):
                self.assertIn(phrase,text)


if __name__=="__main__": unittest.main()
