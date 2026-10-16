"""Coverage keeps reported expense, lease context and contractual measures distinct."""
import html
import json
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

import oracle_interest as oi
import oracle_interest_view as view
import oracle_liquidity as ol
from company_review import money

ROOT = Path(__file__).resolve().parents[1]
ANNUAL = ROOT / "tests/fixtures/ORCL_2026_interest.htm"
QUARTER = ROOT / "tests/fixtures/ORCL_2027Q1_interest.htm"


class OracleInterest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.review = oi.build(ANNUAL, QUARTER)
        cls.snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]

    def test_hand_read_period_build_and_coverage(self):
        c = self.review["calculation"]
        self.assertEqual([c["periods"][p]["operating_income"] for p in ["fy", "current", "prior"]], [20606e6, 6728e6, 4277e6])
        self.assertEqual([c["periods"][p]["interest"] for p in ["fy", "current", "prior"]], [4599e6, 1428e6, 923e6])
        self.assertEqual((c["operating_income"], c["interest"]), (23057e6, 5104e6))
        self.assertAlmostEqual(c["value"], 4.517437304075235)
        self.assertEqual((self.review["period_start"], self.review["period_end"], self.review["tag"]), ("2025-09-01", "2026-08-31", None))

    def test_whole_filing_segment_facts_are_not_substituted(self):
        root = ET.parse(QUARTER).getroot()
        self.assertTrue(any(c.find('.//xbrldi:explicitMember', oi.fx.NS) is not None for c in root.iter(f"{{{oi.fx.XBRLI}}}context")))
        for f in self.review["inputs"].values():
            self.assertEqual(f["dimensions"], [])
            self.assertEqual(f["unit"], "iso4217:USD")
            self.assertEqual(str(f["entity_identifier"]).lstrip("0"), "1341439")
        f = self.review["inputs"]["current_operating_income"]
        self.assertEqual(f["context_id"], "C_f2f9ad85-08df-4e59-8948-bfd9405570de")

    def test_lease_interest_is_context_not_added_to_denominator(self):
        c = self.review["calculation"]
        self.assertEqual([c["periods"][p]["lease_interest"] for p in ["fy", "current", "prior"]], [279e6, 126e6, 48e6])
        self.assertEqual(c["lease_interest"], 357e6)
        self.assertNotEqual(c["interest"], (5104+357)*1e6)
        self.assertTrue(all(g["value"] is None and g["status"] == "Unavailable" for g in self.review["unavailable"].values()))

    def test_composition_and_capitalization_are_not_assumed(self):
        self.assertIn("do not state", self.review["basis"])
        self.assertIn("before any interest-income offset", self.review["basis"])
        self.assertIn("immaterial, not zero", self.review["composition"])
        self.assertIn("Borrowing-only interest coverage is Unavailable", self.review["limitations"])
        self.assertIn("contract-defined", self.review["limitations"])

    def test_wrong_period_entity_unit_tag_source_and_dimensions_are_refused(self):
        for key, value in [("period_start", "2026-06-02"), ("period_end", "2026-08-30"),
                           ("entity_identifier", "59478"), ("unit", "iso4217:EUR"),
                           ("tag", "us-gaap:InterestExpenseNonoperating"), ("accession", "other"),
                           ("source_sha256", "unchecked"), ("dimensions", [["a", "b"]])]:
            with self.subTest(key=key):
                i = deepcopy(self.review["inputs"]);i["current_interest"][key] = value
                c = oi.calculate(i)
                self.assertIsNone(c["interest"]);self.assertIsNone(c["value"])
                self.assertEqual(c["operating_income"], 23057e6)

    def test_missing_numerator_or_denominator_never_becomes_zero(self):
        for key in ["fy_operating_income", "current_interest"]:
            i = deepcopy(self.review["inputs"]);i[key]["value"] = None
            c = oi.calculate(i);self.assertIsNone(c["value"])
            self.assertEqual(c["interest"] if key.startswith("fy") else c["operating_income"], 5104e6 if key.startswith("fy") else 23057e6)

    def test_missing_optional_lease_context_does_not_block_reported_line(self):
        i = deepcopy(self.review["inputs"]);i["fy_lease_interest"]["value"] = None
        c = oi.calculate(i)
        self.assertIsNone(c["lease_interest"])
        self.assertAlmostEqual(c["value"], 23057/5104)

    def test_zero_and_negative_ttm_interest_do_not_divide(self):
        for values in [(0, 0, 0), (4599e6, 1428e6, 7000e6)]:
            i = deepcopy(self.review["inputs"])
            for role, v in zip(["fy", "current", "prior"], values):i[role + "_interest"]["value"] = v
            c = oi.calculate(i);self.assertIsNone(c["value"])
            self.assertIn("zero or negative", c["detail"])

    def test_negative_operating_income_remains_a_valid_loss(self):
        i = deepcopy(self.review["inputs"]);i["fy_operating_income"]["value"] = -20606e6
        self.assertLess(oi.calculate(i)["value"], 0)
        i["current_interest"]["value"] = -1
        self.assertIsNone(oi.calculate(i)["value"])

    def changed(self, mutate):
        root = ET.parse(QUARTER).getroot();mutate(root)
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / "changed.htm";ET.ElementTree(root).write(p)
            return oi.build(ANNUAL, p)

    def test_conflicting_duplicate_interest_is_refused(self):
        def mutate(root):
            f = next(f for f in root.iter(f"{{{oi.fx.IX}}}nonFraction") if f.get("name") == oi.TAGS["interest"])
            extra = deepcopy(f);extra.text = "1";root.append(extra)
        self.assertIsNone(self.changed(mutate)["calculation"]["value"])

    def test_segment_only_operating_income_does_not_replace_consolidated(self):
        def mutate(root):
            for f in root.iter(f"{{{oi.fx.IX}}}nonFraction"):
                if f.get("name") == oi.TAGS["operating_income"] and f.get("contextRef") == self.review["inputs"]["current_operating_income"]["context_id"]:
                    f.set("name", "x:removed")
        c = self.changed(mutate)["calculation"]
        self.assertIsNone(c["operating_income"]);self.assertIsNone(c["value"])
        self.assertEqual(c["interest"], 5104e6)

    def test_saved_review_matches_fixtures_and_full_sources(self):
        self.assertEqual(oi.load(), self.review)
        raw = ROOT.parent / "data/raw/sec/ORCL"
        if (raw / oi.SOURCES["annual"]["document"]).exists():
            self.assertEqual(oi.build(*(raw / oi.SOURCES[k]["document"] for k in ["annual", "quarter"])), self.review)

    def test_stale_review_and_wrong_snapshot_are_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / "review.json"
            with patch.object(oi, "SAMPLE", p):
                self.assertIsNone(oi.load())
                r = deepcopy(self.review);r["inputs"]["fy_interest"]["value"] += 1;r["calculation"] = oi.calculate(r["inputs"])
                p.write_text(json.dumps(r));self.assertIsNone(oi.load())
        self.assertIsNotNone(oi.for_snapshot(self.snapshots[1]))
        for s in [self.snapshots[0], self.snapshots[2]]:self.assertIsNone(oi.for_snapshot(s))
        s = deepcopy(self.snapshots[1]);s["metrics"]["cash"]["accession"] = "other"
        self.assertIsNone(oi.for_snapshot(s))

    def test_view_preserves_limits_source_links_and_other_companies(self):
        shown = html.unescape(view.render(self.snapshots[1], money))
        for text in ["4.52x", "$23.057B", "$5.104B", "September 1, 2025", "capitalized-interest basis",
                     "Borrowing-only", "not added", "3.0x", oi.SOURCES["annual"]["source_url"]]:
            self.assertIn(text, shown)
        self.assertEqual(view.render(self.snapshots[0], money), "")
        self.assertEqual(view.render(self.snapshots[2], money), "")
        self.assertNotIn("$", view.render(self.snapshots[1], money))

    def test_actual_workbook_has_formulas_and_matching_cached_values(self):
        from openpyxl import load_workbook
        b = ol.workbook_bytes(json.loads(ol.SAMPLE.read_text()));self.assertIsNotNone(b)
        v, f = load_workbook(BytesIO(b), data_only=True), load_workbook(BytesIO(b), data_only=False)
        self.assertEqual(v.sheetnames, ["Funding conclusion","Repayment review", "Cash analysis", "Payment dates", "Data", "About"])
        for cell, value in {"E45":23057e6,"E46":5104e6,"E47":357e6,"B49":23057/5104}.items():
            self.assertAlmostEqual(v["Cash analysis"][cell].value, value)
            self.assertTrue(f["Cash analysis"][cell].value.startswith("="))
        self.assertEqual(v["Cash analysis"]["B20"].value, 36369e6)
        self.assertEqual(v["Cash analysis"]["B19"].value, 0)

    def test_changed_coverage_evidence_blocks_stale_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / "interest.json";p.write_bytes(oi.SAMPLE.read_bytes() + b"\n")
            with patch.object(oi, "SAMPLE", p):self.assertIsNone(ol.workbook_bytes(json.loads(ol.SAMPLE.read_text())))


if __name__ == "__main__":
    unittest.main()
