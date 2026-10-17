"""Independent checked answers, period refusal and app/individual-export parity."""
import html
import json
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import company_review
import oracle_liquidity
import ratio_review as rr
import vistra_liquidity

ROOT = Path(__file__).resolve().parents[1]


class RatioReconciliation(unittest.TestCase):
    def setUp(self):
        self.snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]

    def test_checked_answers_from_independent_evidence(self):
        # Source amounts are USD millions, from the balance sheets and checked cash-flow/interest tables.
        answers = {"LLY": (64672, 47741, 54908, 8950, 28083, 9893, 33633, 1079),
                   "ORCL": (55630, 47514, 125337, 36369, 46940, 75660, 23057, 5104),
                   "VST": (10276, 10557, 19895, 435, 5121, 2866, 3563, 1117)}
        for s in self.snapshots:
            rows = {r["key"]: r for r in rr.build(s)["rows"]}
            ca, cl, debt, cash, ocf, capex, op, interest = answers[s["ticker"]]
            expected = {"current_ratio": ca / cl, "net_debt": (debt - cash) * 1e6,
                        "ocf_to_debt": ocf / debt, "fcf_to_debt": (ocf - capex) / debt,
                        "interest_coverage": op / interest}
            for key, value in expected.items():
                self.assertAlmostEqual(rows[key]["value"], value, msg=(s["ticker"], key))
            self.assertEqual(rows["current_ratio"]["period_start"], None)
            self.assertEqual(rows["ocf_to_debt"]["period_start"],
                             "2025-09-01" if s["ticker"] == "ORCL" else "2025-07-01")

    def test_wrong_balance_sheet_date_version_or_unit_is_refused(self):
        for key, bad in (("period_end", "2026-05-31"), ("accession", "other"), ("unit", "EUR")):
            s = deepcopy(self.snapshots[1]); s["metrics"]["current_ratio"]["inputs"]["current_assets"][key] = bad
            with self.assertRaises(ValueError):
                rr.build(s)

    def test_annual_only_and_mislabelled_derived_periods_are_refused(self):
        for metric, field, bad in (("operating_cash_flow", "period_start", "2026-01-01"),
                                   ("fcf_to_debt", "period_end", "2025-12-31"),
                                   ("free_cash_flow", "period_start", "2026-01-01"),
                                   ("current_ratio", "unit", "USD")):
            s = deepcopy(self.snapshots[0]); s["metrics"][metric][field] = bad
            with self.assertRaises(ValueError):
                rr.build(s)

    def test_overlapping_ytd_wrong_tag_or_source_cannot_enter_flow_ratios(self):
        for key, bad in (("coefficient", 1), ("tag", "other"), ("unit", "EUR"),
                         ("accession", "other"), ("period_start", "2025-04-01")):
            s = deepcopy(self.snapshots[0]); s["metrics"]["operating_cash_flow"]["inputs"][2][key] = bad
            with self.assertRaises(ValueError):
                rr.build(s)

    def test_stale_derived_result_and_missing_source_are_refused(self):
        for key in ("net_debt", "current_ratio", "ocf_to_debt", "fcf_to_debt"):
            s = deepcopy(self.snapshots[0]); s["metrics"][key]["value"] += 1
            with self.assertRaises(ValueError):
                rr.build(s)
        s = deepcopy(self.snapshots[0]); s["metrics"]["capital_spending"]["inputs"][0]["value"] = None
        with self.assertRaises(ValueError):
            rr.build(s)

    def test_interest_periods_must_match_and_missing_review_stays_unavailable(self):
        import oracle_interest
        interest = oracle_interest.load()
        with patch.object(oracle_interest, "load", return_value={**interest, "period_end": "2026-05-31"}):
            with self.assertRaises(ValueError):
                rr.build(self.snapshots[1])
        with patch.object(oracle_interest, "load", return_value=None):
            row = rr.build(self.snapshots[1])["rows"][-1]
            self.assertIsNone(row["value"])
            self.assertIn("Unavailable", row["basis"])

    def test_stale_nested_provenance_is_refused(self):
        for metric, field in (("net_debt", "cash"), ("ocf_to_debt", "flow"),
                              ("fcf_to_debt", "debt"), ("free_cash_flow", "capital_spending")):
            s = deepcopy(self.snapshots[0]); s["metrics"][metric]["inputs"][field]["value"] += 1
            with self.assertRaises(ValueError):
                rr.build(s)

    def test_app_evidence_shows_results_periods_formulas_and_specific_limits(self):
        for s in self.snapshots:
            text = html.unescape(company_review.evidence_html(s))
            for row in rr.build(s)["rows"][:-1]:
                self.assertIn(rr.display(row), text)
                self.assertIn(row["period_end"], text)
                self.assertIn(row["basis"], text)
            self.assertIn("total current liabilities", text)
            self.assertIn("not average debt", text)
        self.assertIn("customer prepayments", company_review.evidence_html(self.snapshots[1]))
        self.assertIn("LTSA", company_review.evidence_html(self.snapshots[2]))

    def test_all_individual_workbooks_match_values_formulas_units_dates_and_sources(self):
        for s in self.snapshots:
            stem = {"LLY": "lly", "ORCL": "oracle", "VST": "vistra"}[s["ticker"]]
            path = ROOT / f"samples/{stem}_liquidity.xlsx"
            v, f = load_workbook(path, data_only=True), load_workbook(path, data_only=False)
            for i, row in enumerate(rr.build(s)["rows"], 6):
                cell, formula = v["Debt and liquidity"].cell(i, 2), f["Debt and liquidity"].cell(i, 2)
                self.assertAlmostEqual(cell.value, row["value"], msg=(s["ticker"], row["key"]))
                self.assertTrue(formula.value.startswith("="))
                self.assertEqual(v["Debt and liquidity"].cell(i, 3).value.strftime("%Y-%m-%d"), row["period_end"])
                self.assertEqual(cell.number_format, '0.0%' if row["unit"] == "fraction" else
                                 '0.00"x"' if row["unit"] == "ratio" else '#,##0,,;(#,##0,,);"-"')
                self.assertIn(row["basis"], v["Debt and liquidity"].cell(i, 4).value)
            # Trace the ratio's current-asset/liability facts to actual dated, accession-specific input rows.
            values = list(v["Data"].values)
            for tag in ("us-gaap:AssetsCurrent", "us-gaap:LiabilitiesCurrent"):
                self.assertTrue(any(tag in row and s["metrics"]["cash"]["accession"] in row for row in values))
            self.assertEqual(json.loads(path.with_suffix('.manifest.json').read_text())["ratios_sha256"], rr.fingerprint())
            self.assertFalse(v["Debt and liquidity"].sheet_view.showGridLines)
            self.assertTrue(v["Debt and liquidity"]["B9"].border.right.style)
            v.close(); f.close()

    def test_ratio_rule_or_source_change_invalidates_each_download(self):
        with patch.object(rr, "fingerprint", return_value="changed"):
            self.assertIsNone(company_review.saved_workbook(self.snapshots, ticker="LLY"))
            self.assertIsNone(oracle_liquidity.workbook_bytes(json.loads(oracle_liquidity.SAMPLE.read_text())))
            self.assertIsNone(vistra_liquidity.workbook_bytes(json.loads(vistra_liquidity.SAMPLE.read_text())))


if __name__ == "__main__":
    unittest.main()
