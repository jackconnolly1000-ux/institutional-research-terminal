"""Real-filing cash controls, exact dimensions, dated resources and exported evidence."""
import html
import json
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
import filing_xbrl as fx
import vistra_liquidity as vl
import vistra_liquidity_view as view
from company_review import money, render

ROOT=Path(__file__).resolve().parents[1]
Q=ROOT/'tests/fixtures/VST_liquidity_quarter.htm'
K=ROOT/'tests/fixtures/VST_liquidity_annual.htm'
RAW=ROOT.parent/'data/raw/sec/VST'
M=1_000_000


class VistraLiquidity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.saved=json.loads(vl.SAMPLE.read_text())
        cls.snapshots=json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots']
        cls.snapshot=next(s for s in cls.snapshots if s['ticker']=='VST')

    def changed(self,key,attr,value,annual=False):
        root=ET.parse(K if annual else Q).getroot()
        fact=self.saved['inputs'][key]
        for e in root.iter('{'+fx.IX+'}nonFraction'):
            if e.get('name')==fact['tag'] and e.get('contextRef')==fact['context_id']:
                if attr=='text':e.text=value
                else:e.set(attr,value)
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'changed.htm';ET.ElementTree(root).write(p,encoding='utf-8',xml_declaration=True)
            return vl.build(Q if annual else p,p if annual else K)

    def test_real_fixtures_match_saved_review(self):
        self.assertEqual(vl.build(Q,K),self.saved)
        self.assertEqual(self.saved['status'],'Checked')
        self.assertEqual(self.saved['cash'],{'fcf':650*M,'distributions':959*M,'net_debt_cash':-162*M,
                          'other_investing':293*M,'restricted':46*M})
        self.assertEqual(self.saved['debt'],{'principal':19764*M,'baseline':19895*M})
        self.assertEqual(self.saved['annual']['principal'],17196*M)
        for f in self.saved['inputs'].values():
            self.assertIsNotNone(f['value'])
            self.assertEqual(f['unit'],'iso4217:USD')
            self.assertIn(f['accession'],[p['accession'] for p in vl.PERIODS.values()])

    @unittest.skipUnless((RAW/'vistra-20260630.htm').exists(),'Full SEC documents not saved locally')
    def test_full_saved_filings_match_trimmed_facts(self):
        self.assertEqual(vl.build(RAW/'vistra-20260630.htm',RAW/'vistra-20251231.htm'),self.saved)

    def test_cash_basis_and_debt_financing_are_distinct(self):
        v={k:f['value'] for k,f in self.saved['inputs'].items()}
        self.assertEqual(v['opening_total']+v['ocf']+v['investing']+v['financing'],v['closing_total'])
        self.assertEqual((v['closing_total'],v['cash'],v['restricted_current'],v['restricted_noncurrent']),
                         (481*M,435*M,41*M,5*M))
        self.assertEqual(v['debt_issued'],6422*M)
        self.assertEqual(v['nuveen_principal_paid'],19*M)
        self.assertEqual(self.saved['cash']['fcf']-self.saved['cash']['distributions'],-309*M)

    def test_missing_wrong_unit_and_conflicting_cash_facts_block_cash(self):
        for attr,value in [('name','x:Missing'),('unitRef','EUR'),('text','1')]:
            with self.subTest(attr=attr):
                result=self.changed('capex',attr,value)
                self.assertIsNone(result['cash']);self.assertEqual(result['status'],'Unavailable')
                self.assertIsNotNone(result['capacity'])

    def test_wrong_period_and_cash_use_sign_are_refused(self):
        result=self.changed('capex','contextRef','c-8')
        self.assertIsNone(result['cash'])
        self.assertIsNone(self.changed('capex','sign','-')['cash'])
        for k in ['debt_issued','debt_repaid','restricted_current','financing']:
            self.assertIsNone(self.changed(k,'text','1')['cash'])

    def test_exact_dimensions_required_default_reader_still_refuses(self):
        fact=self.saved['inputs']['revolver_available']
        # The undimensioned 5.860B total is a different reporting scope from the revolver.
        default=fx.read_instant(Q,vl.PERIODS['quarter']['accession'],fact['tag'],vl.END)
        self.assertEqual(default['value'],5860*M)
        self.assertEqual(default['dimensions'],[])
        got=fx.read_instant(Q,vl.PERIODS['quarter']['accession'],fact['tag'],vl.END,dimensions=vl.REVOLVER)
        self.assertEqual(got['value'],4408*M)
        self.assertEqual(got['context_id'],'c-660')
        self.assertIsNone(fx.read_instant(Q,'checked',fact['tag'],vl.END,
                            dimensions=[*vl.REVOLVER,('x:Axis','x:Member')])['value'])
        self.assertIsNone(fx.read_instant(Q,'checked',fact['tag'],vl.END,
                            dimensions=[vl.REVOLVER[0],(vl.INSTRUMENT,'vistra:WrongFacilityMember')])['value'])
        self.assertIsNone(self.changed('revolver_available','contextRef','c-8')['capacity'])

    def test_wrong_entity_is_unavailable(self):
        root=ET.parse(Q).getroot()
        for e in root.iter('{'+fx.XBRLI+'}identifier'):e.text='1'
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'wrong.htm';ET.ElementTree(root).write(p)
            r=vl.build(p,K)
        self.assertIsNone(r['cash']);self.assertIsNone(r['capacity']);self.assertIsNone(r['debt'])
        self.assertIn('Unavailable',r['detail'])

    def test_capacity_is_dated_net_of_lcs_not_current_drawability(self):
        r=self.saved;v={k:f['value'] for k,f in r['inputs'].items()}
        self.assertEqual(r['capacity']['reported_liquidity'],6295*M)
        self.assertEqual(v['revolver_limit']-v['revolver_drawn']-v['revolver_lc'],v['revolver_available'])
        self.assertEqual(v['revolver_lc']+v['loc_other']+v['loc_alternative'],3587*M)
        self.assertLess(v['commodity_available'],v['commodity_limit'])
        self.assertIsNone(r['capacity']['later_drawable'])
        self.assertTrue(all(item['later_available'] is None for item in r['resources']))
        self.assertEqual([v[k] for k in ['bcop_available','zero_available','txef_available']],[0,0,0])
        self.assertIsNone(self.changed('revolver_lc','text','1')['capacity'])
        self.assertIsNone(self.changed('commodity_available','text','1751')['capacity'])
        self.assertIsNotNone(self.changed('loc_total','text','1')['cash'])

    def test_project_debt_not_added_twice_and_combined_term_balance_read_once(self):
        r=self.saved;v={k:f['value'] for k,f in r['inputs'].items()}
        self.assertEqual(sum(v[k] for k in ['bridge','construction_total','zero_debt','txef_debt']),1633*M)
        self.assertEqual(v['pulaski_construction']+v['bcop_term'],v['construction_total'])
        self.assertEqual(v['shortterm'],0)
        self.assertEqual(r['inputs']['bcop_term']['context_id'],'c-686')
        self.assertIsNone(self.changed('bcop_term','text','1')['debt'])
        self.assertIsNone(self.changed('shortterm','name','x:Missing')['debt'])

    def test_selected_payment_dates_and_conditional_obligations(self):
        r=self.saved;p=r['payments']
        self.assertEqual([x['value']/M for x in p],[261,297,500,631,800,1300])
        self.assertEqual((p[1]['due_start'],p[1]['due_end']),('2026-12-01','2026-12-31'))
        self.assertIn('conversion',p[1]['detail'].lower())
        self.assertEqual(p[3]['due_end'],'2026-12-31')
        self.assertEqual(r['inputs']['nuveen_carrying']['value'],613*M)
        self.assertEqual(r['inputs']['cogentrix']['value'],2300*M)
        self.assertNotIn('helix_initial',[x['key'] for x in p])
        self.assertEqual(r['inputs']['helix_initial']['value'],500*M)

    def test_annual_schedule_is_historical_and_controlled(self):
        self.assertEqual([self.saved['inputs']['annual_'+str(i)]['value']/M for i in range(6)],
                         [1201,3435,786,2362,2853,6559])
        self.assertIsNone(self.changed('annual_0','text','1',annual=True)['annual'])
        self.assertIsNotNone(self.changed('annual_0','text','1',annual=True)['cash'])

    def test_snapshot_and_workbook_hash_gate(self):
        self.assertEqual(vl.for_snapshot(self.snapshot),self.saved)
        self.assertIsNone(vl.for_snapshot(self.snapshots[0]))
        changed=deepcopy(self.snapshot);changed['metrics']['cash']['period_end']='2026-07-01'
        self.assertIsNone(vl.for_snapshot(changed))
        changed=deepcopy(self.snapshot);changed['metrics']['cash']['value']=1
        self.assertEqual(vl.for_snapshot(changed)['status'],'Unavailable')
        self.assertIsNotNone(vl.workbook_bytes(self.saved))
        with patch.object(Path,'read_bytes',return_value=b'stale'):
            self.assertIsNone(vl.workbook_bytes(self.saved))

    def test_view_preserves_sections_and_carries_uncertainty(self):
        s=html.unescape(view.render(self.snapshot,money))
        for value in ['How does Vistra fund its repayments?','$650M','$959M','$6.295B','$1.092B',
                      'Pulaski','Current drawable amounts are Unavailable','Oct 15','2026-10-05',
                      '5.50x','35%','Nuveen','Helix','Cogentrix','otherwise non-recourse']:
            self.assertIn(value,s)
        self.assertNotIn('<style>',s)
        self.assertNotIn('$',view.render(self.snapshot,money))
        total=render(self.snapshot)
        self.assertEqual(total.count('<style>'),1)
        self.assertIn('What sits outside baseline borrowing?',total)
        self.assertIn('What remained after capital spending?',total)
        self.assertEqual(view.render(self.snapshots[0],money),'')

    def test_later_events_count_issuance_once_and_correct_filing_date(self):
        r=self.saved
        issued=[e for e in r['events'] if e['kind']=='New debt issued after balance sheet']
        self.assertEqual(len(issued),1)
        october=next(e for e in r['events'] if e['accession'].endswith('038468'))
        self.assertEqual((october['event_date'],october['filed']),('2026-09-30','2026-10-05'))
        self.assertIn('35%',r['terms'][0]['detail']);self.assertIn('5.50x',r['terms'][0]['detail'])
        self.assertIn('undrawn',r['terms'][0]['detail'])

    def test_workbook_values_dates_formulas_chart_and_blank_capacity(self):
        b=vl.WORKBOOK.read_bytes();w=load_workbook(BytesIO(b),data_only=True)
        formulas=load_workbook(BytesIO(b),data_only=False)
        self.assertEqual(w.sheetnames,['Funding conclusion','Borrowing and leases','Debt and liquidity','Cash review','Payment dates','Financing capacity','Data','About'])
        self.assertEqual([w['Cash review'][k].value for k in ['B8','B13','B14','B23','B37','B47']],
                         [650*M,959*M,-309*M,-162*M,-341*M,435*M])
        for c in ['B32','B35','B39','B44','B49']:self.assertEqual(w['Cash review'][c].value,0)
        self.assertEqual(w['Financing capacity']['B16'].value,6295*M)
        self.assertIsNone(w['Financing capacity']['B9'].value)
        self.assertEqual(w['Financing capacity']['C9'].value,'Unavailable')
        self.assertEqual(w['Payment dates']['B9'].value.strftime('%Y-%m-%d'),'2026-12-31')
        self.assertEqual(formulas['Cash review']['B14'].value,'=B8-B13')
        self.assertEqual(len(formulas['Payment dates']._charts),1)
        for s in formulas:
            self.assertFalse(s.merged_cells.ranges)
            self.assertFalse(s.sheet_view.showGridLines)


if __name__=='__main__':unittest.main()
