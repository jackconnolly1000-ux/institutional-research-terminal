"""Every source URL in the three company workbooks is a native hyperlink to exactly that URL."""
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

import company_review
import oracle_liquidity as ol
import vistra_liquidity as vl

ROOT = Path(__file__).resolve().parents[1]
EXPECTED = {"lly_liquidity": 134, "oracle_liquidity": 140, "vistra_liquidity": 175}


class WorkbookHyperlinks(unittest.TestCase):
    def test_every_source_url_is_a_link_to_itself(self):
        for stem, expected in EXPECTED.items():
            wb = load_workbook(ROOT / f"samples/{stem}.xlsx")
            urls = [c for ws in wb for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and (c.value.startswith("https://") or c.value.startswith("Source: https://"))]
            self.assertEqual(len(urls), expected, stem)
            for cell in urls:
                self.assertIsNotNone(cell.hyperlink, f"{stem} {cell.parent.title}!{cell.coordinate}")
                self.assertEqual(cell.hyperlink.target, cell.value.removeprefix("Source: "))

    def test_guides_no_longer_say_to_copy_urls(self):
        for stem in EXPECTED:
            wb = load_workbook(ROOT / f"samples/{stem}.xlsx")
            text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
            self.assertNotIn("Copy a full URL", text)
            self.assertNotIn("copyable source text", text)
            self.assertIn("Select a source URL to open it in a browser.", text)

    def test_linked_workbooks_are_still_served(self):
        snaps = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]
        by = {s["ticker"]: s for s in snaps}
        self.assertIsNotNone(company_review.saved_workbook(snaps, ticker="LLY"))
        self.assertIsNotNone(ol.workbook_bytes(ol.for_snapshot(by["ORCL"])))
        self.assertIsNotNone(vl.workbook_bytes(vl.for_snapshot(by["VST"])))


if __name__ == "__main__":
    unittest.main()
