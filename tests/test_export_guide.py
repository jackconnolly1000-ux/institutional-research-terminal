"""Export navigation and interpretation must agree with the actual company outputs."""
import json
from pathlib import Path
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import company_history
import company_review
import export_guide
import funding_conclusions
import oracle_liquidity
import vistra_liquidity

ROOT=Path(__file__).resolve().parents[1]


class IndividualExportParity(unittest.TestCase):
    def setUp(self):
        self.snapshots=json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots']

    def test_contents_cover_actual_tabs_and_preserve_conclusions(self):
        for s in self.snapshots:
            t=s['ticker'];stem={'LLY':'lly','ORCL':'oracle','VST':'vistra'}[t]
            w=load_workbook(ROOT/f'samples/{stem}_liquidity.xlsx',data_only=True)
            self.assertEqual([name for name,_ in export_guide.contents(t)],w.sheetnames)
            self.assertEqual([w['Funding conclusion'].cell(i,1).value for i in range(21,21+len(w.sheetnames))],w.sheetnames)
            self.assertEqual(w['Funding conclusion']['B6'].value,funding_conclusions.load(t)['conclusion'])
            self.assertIn('Unavailable',w['Funding conclusion']['B17'].value)
            self.assertEqual(len(w['Core charts']._charts),2)
            self.assertIn('Data',w['Core charts']['D41'].value)
            self.assertIn('Annual column:',w['History']['D10'].value)
            if t in ('LLY','ORCL'):self.assertIn('Latest column only:',w['History']['D7'].value)
            self.assertEqual(json.loads((ROOT/f'samples/{stem}_liquidity.manifest.json').read_text())['guide_sha256'],export_guide.fingerprint())
            w.close()

    def test_older_history_basis_cannot_inherit_ttm_capitalization_amount(self):
        for s in self.snapshots:
            model=company_history.build(s);basis=model['annual_interest_basis']
            if s['ticker']=='LLY':self.assertIn('Annual interest is reported net',basis);self.assertIn('unstated',basis)
            if s['ticker']=='ORCL':self.assertIn('Unavailable',basis)
            if s['ticker']=='VST':
                self.assertNotIn('151',basis);self.assertNotIn('1,117',basis)
                self.assertIn('before the capitalized-interest deduction',basis)
                self.assertTrue(all('docs/' not in r['annual_detail'] for r in model['rows']))

    def test_stale_export_layout_refuses_all_three_downloads(self):
        with patch.object(export_guide,'fingerprint',return_value={'changed':True}):
            self.assertIsNone(company_review.saved_workbook(self.snapshots,ticker='LLY'))
            self.assertIsNone(oracle_liquidity.workbook_bytes(oracle_liquidity.for_snapshot(self.snapshots[1])))
            self.assertIsNone(vistra_liquidity.workbook_bytes(vistra_liquidity.for_snapshot(self.snapshots[2])))

    def test_optional_combined_scope_is_explicit(self):
        self.assertIn('original core metric summary',export_guide.COMBINED_CAPTION)
        self.assertIn("company's primary download",export_guide.COMBINED_CAPTION)
        self.assertIn('does not update the checked narrative',export_guide.CAPTION)
