"""Reconcile the actual Oracle schedule; refuse source gaps and misleading date combinations."""
import html
import json
import sys
import tempfile
import unittest
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import company_review as view  # noqa: E402
import oracle_financing as of  # noqa: E402
import reviewed_financing as rf  # noqa: E402

FIXTURE = ROOT / "tests/fixtures/ORCL_2026_maturities.htm"
QUARTER = ROOT / "tests/fixtures/ORCL_2026Q1_financing.htm"
RAW = ROOT.parent / "data/raw/sec/ORCL"
M = 1_000_000


class OracleFinancing(unittest.TestCase):
    def setUp(self):
        self.saved = json.loads(of.SAMPLE.read_text())
        self.snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]
        self.oracle = self.snapshots[1]

    def test_saved_schedule_rebuilds_and_matches_printed_note(self):
        got = of.build(FIXTURE, QUARTER)
        self.assertEqual(got, self.saved)
        self.assertEqual([f["value"] // M for f in got["buckets"]], [7210, 10145, 5500, 7250, 9750, 90250])
        self.assertEqual((got["total"], got["first_five_years"]), (130105*M, 39855*M))
        self.assertAlmostEqual(got["first_five_share"], 39855/130105)
        self.assertEqual((got["as_of"], got["anchor"], got["current_schedule"]),
                         ("2026-05-31", "2026-08-31", None))
        self.assertEqual(got["buckets"][0]["due_start"], "2026-06-01")
        self.assertIsNone(got["buckets"][-1]["due_end"])
        self.assertTrue(all(f["scale"] == 6 and f["unit"] == "iso4217:USD" for f in got["buckets"]))

    def test_principal_and_carrying_bases_reconcile(self):
        c = self.saved["controls"]
        self.assertEqual((c["discount_and_costs"]["value"], c["carrying_debt"]["value"]), (564*M, 129541*M))
        self.assertEqual(self.saved["total"] - c["discount_and_costs"]["value"], c["carrying_debt"]["value"])
        self.assertNotEqual(self.saved["total"], self.oracle["metrics"]["gross_debt"]["value"])
        self.assertEqual([f["value"] for f in self.saved["quarter_flows"]], [19909*M, 4202*M])

    def tamper(self, transform, quarter=False):
        tree = ET.parse(QUARTER if quarter else FIXTURE)
        transform(tree.getroot())
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "changed.htm"
            tree.write(path, encoding="utf-8", xml_declaration=True)
            return of.build(FIXTURE, path) if quarter else of.build(path, QUARTER)

    def test_missing_or_inconsistent_annual_facts_block_the_total(self):
        bucket = "us-gaap:" + of.BUCKETS[0][3]
        cases = [
            (bucket, "name", "x:WrongTag", "no us-gaap:"),
            (bucket, "text", "7,211", "do not equal"),
            ("us-gaap:" + of.CONTROLS["discount_and_costs"], "text", "565", "does not equal"),
            (bucket, "sign", "-", "negative"),
            (bucket, "unitRef", "unknown", "units other than"),
        ]
        for tag, attr, value, why in cases:
            with self.subTest(tag=tag, attr=attr):
                def change(root):
                    for f in root.iter(f"{{{of.fx.IX}}}nonFraction"):
                        if f.get("name") == tag:
                            if attr == "text":
                                f.text = value
                            else:
                                f.set(attr, value)
                got = self.tamper(change)
                self.assertEqual(got["status"], "Unavailable")
                self.assertIsNone(got["total"])
                self.assertIn(why, got["detail"])

    def test_wrong_entity_or_date_is_refused(self):
        for part, value, why in (("identifier", "0000000001", "CIK does not match"),
                                  ("instant", "2026-05-30", "no undimensioned")):
            with self.subTest(part=part):
                def change(root):
                    for e in root.iter(f"{{{of.fx.XBRLI}}}{part}"):
                        if part == "identifier" or e.text == of.AS_OF:
                            e.text = value
                got = self.tamper(change)
                self.assertIsNone(got["total"])
                self.assertIn(why, got["detail"])

    def test_absent_document_is_unavailable_not_zero(self):
        got = of.build(FIXTURE.with_name("missing.htm"))
        self.assertEqual(got["status"], "Unavailable")
        self.assertIsNone(got["total"])

    def test_quarter_failure_does_not_erase_valid_historical_schedule(self):
        def change(root):
            for e in root.iter(f"{{{of.fx.XBRLI}}}identifier"):
                e.text = "1"
        got = self.tamper(change, quarter=True)
        self.assertEqual(got["total"], 130105*M)
        self.assertTrue(all(f["value"] is None and "CIK" in f["detail"] for f in got["quarter_flows"]))

    def test_schedule_attaches_only_to_the_checked_company_and_anchor(self):
        for ticker in (0, 2):
            self.assertIsNone(of.for_snapshot(self.snapshots[ticker]))
        for key, bad in (("accession", "another-version"), ("period_end", "2026-11-30")):
            changed = deepcopy(self.oracle)
            changed["metrics"]["cash"][key] = bad
            self.assertIsNone(of.for_snapshot(changed))

    def test_oracle_reviews_match_filing_metadata_and_leave_unknowns_unchecked(self):
        rows = [{**f, "accession": acc, "review": "Not checked"} for acc, f in of.FILINGS.items()]
        rows.append({"form": "8-K", "filed": "2026-10-15", "accession": "later", "review": "Not checked"})
        inv = {"since": of.ANCHOR, "through": "2026-10-15", "status": "incomplete", "filings": rows}
        got = rf.annotate(inv, "ORCL")
        self.assertEqual([r["review"] for r in got["filings"]], ["Checked 2026-10-14"]*2 + ["Not checked"])
        self.assertEqual(got["status"], "incomplete")
        self.assertEqual(inv["filings"][0]["review"], "Not checked")
        for key, bad in (("form", "8-K/A"), ("filed", "2026-09-11"), ("period", "2026-09-09")):
            changed = deepcopy(inv)
            changed["filings"][0][key] = bad
            self.assertEqual(rf.annotate(changed, "ORCL")["filings"][0]["review"], "Not checked")
        inv["since"] = "2026-05-31"
        self.assertTrue(all(r["review"] == "Not checked" for r in rf.annotate(inv, "ORCL")["filings"]))

    def test_view_distinguishes_principal_dates_and_publication_from_transaction(self):
        shown = html.unescape(view.maturity_html(self.oracle))
        for text in ("Historical schedule at May 31, 2026", "$10.145B", "$90.250B thereafter", "30.6%",
                     "Fiscal 2027 below includes June–August", "not a replacement", "$19.909B", "$4.202B",
                     "already reflected", "no financing event", "index through October 15, 2026", "not added"):
            self.assertIn(text, shown)
        self.assertEqual(view.maturity_html(self.snapshots[0]), "")

    def test_excel_keeps_numeric_inputs_formulas_dates_and_chart(self):
        wb = load_workbook(BytesIO(view.workbook(self.snapshots)))
        schedule, raw, events = (wb[n] for n in ("Oracle maturities", "Oracle debt inputs", "Oracle financing"))
        self.assertEqual([raw.cell(i, 2).value for i in range(2, 8)], [7210*M, 10145*M, 5500*M, 7250*M, 9750*M, 90250*M])
        self.assertEqual(schedule["D8"].value, "=SUM(D2:D7)")
        self.assertEqual(schedule["D12"].value, "=D8-D11")
        self.assertEqual(schedule["D14"].value, "=D12-D13")
        self.assertEqual(schedule["D16"].value, "=D15/D8")
        self.assertEqual(schedule["E17"].value, "Unavailable")
        self.assertIsNone(schedule["D17"].value)
        self.assertEqual(raw["C2"].value, datetime(2026, 5, 31))
        self.assertEqual((events["A5"].value, events["H5"].value), (datetime(2026, 9, 12), datetime(2026, 9, 14)))
        self.assertEqual(raw["H2"].value, of.ACCESSION)
        self.assertEqual(events["E2"].value, "='Oracle debt inputs'!B11")
        self.assertIsNone(events["E4"].value)  # earnings filing is not another equity transaction
        self.assertEqual(len(schedule._charts), 1)
        self.assertEqual(schedule._charts[0].series[0].val.numRef.f, "'Oracle maturities'!$D$2:$D$6")
        self.assertEqual(wb["Debt and liquidity"]["C12"].value, "='Liquidity inputs'!D14+'Liquidity inputs'!D15")
        for name in ("Oracle maturities", "Oracle debt inputs", "Oracle financing"):
            self.assertEqual(wb[name]["A1"].font.name, "Arial")
            self.assertFalse(wb[name].sheet_view.showGridLines)

    @unittest.skipUnless((RAW / "orcl-20260531.htm").exists(), "full saved Oracle filings not available")
    def test_full_saved_sources_match_trimmed_fixtures_and_eight_k_findings(self):
        full = of.build(RAW / "orcl-20260531.htm", RAW / "orcl-20260831.htm")
        self.assertEqual(full["total"], self.saved["total"])
        for group in ("buckets", "quarter_flows"):
            self.assertEqual([(f["tag"], f["value"], f["context_id"]) for f in full[group]],
                             [(f["tag"], f["value"], f["context_id"]) for f in self.saved[group]])
        from bs4 import BeautifulSoup
        checks = {
            "8k-000119312526387905-orcl-20260910.htm": ["September 10, 2026", "$1,625", "$0.50"],
            "8k-000119312526387905-orcl-ex99_1.htm": ["During Q1 FY 2027", "$20 billion", "before commissions"],
            "8k-000119312526389753-d20034d8k.htm": ["September 12, 2026", "No Oracle stock was sold"],
            "8k-000119312526389753-d20034dex991.htm": ["September 12, 2026", "No Oracle stock was sold"],
        }
        for filename, phrases in checks.items():
            text = " ".join(BeautifulSoup((RAW / filename).read_text(), "html.parser").stripped_strings)
            for phrase in phrases:
                self.assertIn(phrase, text.replace("\xa0", " "), filename)


if __name__ == "__main__":
    unittest.main()
