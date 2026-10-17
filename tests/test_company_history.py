import json
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
import company_history as h
import company_review

ROOT = Path(__file__).resolve().parents[1]


class History(unittest.TestCase):
    def setUp(self):
        self.snapshots = json.loads((ROOT/'samples/liquidity_pilots.json').read_text())['snapshots']

    def test_two_dates_and_independent_annual_answers(self):
        expected = {'LLY': (7268e6,42503e6,35235e6,55629/35228,16813e6,7841e6,8972e6,16813/42503,8972/42503,26302/895),
                    'ORCL': (31289e6,129541e6,98252e6,46567/41764,31977e6,55663e6,-23686e6,31977/129541,-23686/129541,20606/4599),
                    'VST': (785e6,None,None,9179/11814,4070e6,2752e6,1318e6,None,None,1906/1107)}
        for s in self.snapshots:
            model=h.build(s)
            self.assertNotEqual(model['annual_end'], model['period_end'])
            self.assertIn('overlap',model['note'])
            for r,v in zip(model['rows'],expected[s['ticker']]):
                if v is None:self.assertIsNone(r['annual_value']);self.assertIn('not checked',r['annual_detail'])
                else:self.assertAlmostEqual(r['annual_value'],v)
                self.assertEqual(r['value'],next(x for x in __import__('ratio_review').build(s)['rows'] if x['key']==r['key'])['value'])

    def test_changed_facts_are_refused(self):
        with patch.object(h,'CHECKED_SHA256','changed'):
            with self.assertRaises(ValueError):h.build(self.snapshots[0])

    def test_app_has_two_dated_sources_and_explicit_missing_history(self):
        for s in self.snapshots:
            text=h.render(s,company_review.money)
            self.assertIn('What changed since the fiscal year end?',text)
            self.assertIn('Annual source filing',text)
            self.assertIn('Fiscal year',text)
            self.assertIn('Later TTM',text)
            self.assertIn('overlap',text)
        self.assertIn('Unavailable',h.render(self.snapshots[2],company_review.money))

    def test_workbooks_reconcile_every_history_cell_and_keep_unknowns_text(self):
        for s in self.snapshots:
            stem={'LLY':'lly','ORCL':'oracle','VST':'vistra'}[s['ticker']]
            p=ROOT/f'samples/{stem}_liquidity.xlsx';v=load_workbook(p,data_only=True);f=load_workbook(p,data_only=False)
            for i,r in enumerate(h.build(s)['rows'],6):
                if r['annual_value'] is None:self.assertEqual(v['History'].cell(i,2).value,'Unavailable');self.assertEqual(f['History'].cell(i,2).data_type,'s')
                else:self.assertAlmostEqual(v['History'].cell(i,2).value,r['annual_value']);self.assertTrue(f['History'].cell(i,2).value.startswith('='))
                self.assertAlmostEqual(v['History'].cell(i,3).value,r['value'])
            self.assertEqual(json.loads(p.with_suffix('.manifest.json').read_text())['history_sha256'],h.fingerprint())
            v.close();f.close()
