"""Diagnostic: does fetch_firm_data's risk-factor comparison use the right filings and the right text?

Standalone copy of the filing-selection and text-extraction logic in app.py fetch_firm_data
(lines 180-259 as of commit c8bfd6b). It can't import app.py, which runs Streamlit code at import,
so the copied lines are marked COPIED and must be kept in sync by hand if fetch_firm_data changes.
Never calls the Anthropic API.

Usage (from app/):
    .venv/bin/python tests/check_risk_extraction.py              # parts A and B
    .venv/bin/python tests/check_risk_extraction.py --excerpts   # also part C (needs an .xlsx in data/exports)
    .venv/bin/python tests/check_risk_extraction.py --excerpts --tickers AAPL   # limit to some tickers
"""
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup

TICKERS = ["AAPL", "FAST", "STRL"]
PROJECT = Path(__file__).resolve().parents[2]          # ~/Stock Project
RAW_DIR = PROJECT / "data" / "raw" / "sec"             # downloaded filings (cached by accession number)
EXPORT_DIR = PROJECT / "data" / "exports"
MIN_GAP = 0.15                                          # seconds between SEC requests: ~6-7/s, under SEC's 10/s

# --- COPIED from fetch_firm_data (app.py lines 219-220) ---
risk_regex = r'(?:item\s+1a\b[\.\:\s\-\—\xa0]*(?:risk\s+factors)?)(.*?)(?:item\s+(?:1b|2)\b[\.\:\s\-\—\xa0]|$)'
MDA_REGEX_10K = r'(?:item\s+7\b[\.\:\s\-\—\xa0]*(?:management[\'\’]?s\s+discussion)?)(.*?)(?:item\s+(?:7a|8)\b[\.\:\s\-\—\xa0]|$)'
MDA_REGEX_10Q = r'(?:item\s+2\b[\.\:\s\-\—\xa0]*(?:management[\'\’]?s\s+discussion)?)(.*?)(?:item\s+(?:3|4)\b[\.\:\s\-\—\xa0]|$)'
RISK_FALLBACK = (5000, 35000)
MDA_FALLBACK = (40000, 80000)
MIN_MATCH_LEN = 4000


# --- SEC access: same headers and retry rule as app.py sec_get_request, plus a throttle ---
_last_request = [0.0]
request_count = [0]

def sec_get(url: str, headers: dict):
    for attempt in range(3):
        wait = MIN_GAP - (time.time() - _last_request[0])
        if wait > 0:
            time.sleep(wait)
        _last_request[0] = time.time()
        request_count[0] += 1
        try:
            resp = requests.get(url, headers=headers, timeout=12)
            if resp.status_code == 200: return resp
            elif resp.status_code == 429: time.sleep(1.5 * (attempt + 1))
        except Exception: time.sleep(1.0)
    return None

def get_document(ticker, cik, acc, doc, headers):
    """Filing documents never change once filed, so cache them on disk by accession number."""
    path = RAW_DIR / ticker / f"{acc}_{doc}"
    if path.exists():
        return path.read_bytes(), "cache"
    resp = sec_get(f"https://www.sec.gov/Archives/edgar/data/{cik}/{acc}/{doc}", headers)
    if not resp:
        return None, "FAILED"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(resp.content)
    return resp.content, "downloaded"

def clean(html: bytes) -> str:
    # COPIED: app.py line 227
    return re.sub(r'\s+', ' ', BeautifulSoup(html, 'html.parser').get_text(separator=' ', strip=True))


# --- Extraction with a full match report ---
def ended_by(terminator: str) -> str:
    t = terminator.strip().lower()
    if not t:
        return "end of document"
    m = re.match(r'item\s+(\w+)', t)
    return f"Item {m.group(1).upper()}" if m else repr(terminator)

def extract(text: str, regex: str, fallback: tuple):
    """Same selection rule as the app: longest group(1) whose stripped length > 4000, else the fallback slice."""
    matches = []
    for m in re.compile(regex, re.I | re.S).finditer(text):
        body = m.group(1)
        matches.append({
            "start": m.start(),
            "heading": text[m.start():m.start(1)],
            "len": len(body.strip()),
            "passed": len(body.strip()) > MIN_MATCH_LEN,
            "ended_by": ended_by(text[m.end(1):m.end()]),
            "head": body.strip()[:80],
            "tail": body.strip()[-200:],
            "body": body,
        })
    passing = [x for x in matches if x["passed"]]
    if passing:
        chosen = max(passing, key=lambda x: len(x["body"]))  # max() keeps the first on ties, like the app
        return matches, matches.index(chosen), chosen["body"]
    return matches, None, text[fallback[0]:fallback[1]]

def report_extraction(label, text, regex, fallback, cut, cut_before_strip=False, heading_word="risk"):
    matches, chosen_idx, chosen = extract(text, regex, fallback)
    # The app cuts current text after strip ([:20000]) but prior text before strip ([:18000] then .strip())
    sent = chosen[:cut].strip() if cut_before_strip else chosen.strip()[:cut]
    print(f"\n  {label}: {len(matches)} regex match(es); cleaned filing is {len(text):,} chars")
    for i, x in enumerate(matches):
        flag = "CHOSEN" if i == chosen_idx else ("passed" if x["passed"] else "too short")
        print(f"    #{i+1} at char {x['start']:,}: {x['len']:,} chars [{flag}], ended by {x['ended_by']}")
        print(f"       heading: {x['heading'][:60]!r}")
        print(f"       starts : {x['head']!r}")
        print(f"       ends   : ...{x['tail']!r}")
    source = f"match #{chosen_idx+1}" if chosen_idx is not None else f"FALLBACK SLICE [{fallback[0]}:{fallback[1]}]"
    print(f"    -> used {source}: {len(chosen.strip()):,} chars, {len(sent):,} survive the {cut:,} cut")
    print(f"    -> first 300 chars: {chosen.strip()[:300]!r}")
    heading_ok = chosen_idx is not None and heading_word in matches[chosen_idx]["heading"].lower()
    return {
        "source": source, "len": len(chosen.strip()), "sent": len(sent), "sent_text": sent,
        "fallback": chosen_idx is None,
        "ended_by": matches[chosen_idx]["ended_by"] if chosen_idx is not None else "-",
        "heading_ok": heading_ok,
    }


# --- Part C helpers ---
def normalize(s: str) -> str:
    s = s.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")
    s = s.replace("—", "-").replace("–", "-").replace("\xa0", " ").replace("…", "...")
    return re.sub(r'\s+', ' ', s).strip()

def longest_run(excerpt: str, hay: str) -> float:
    words = excerpt.split()
    for k in range(len(words), 0, -1):
        if any(" ".join(words[i:i + k]) in hay for i in range(len(words) - k + 1)):
            return k / len(words)
    return 0.0

def check_excerpts(sent_by_ticker, full_by_ticker):
    import openpyxl
    exports = sorted(EXPORT_DIR.glob("*.xlsx"), key=lambda p: p.stat().st_mtime)
    if not exports:
        print("\nPART C skipped: no .xlsx in data/exports"); return
    path = exports[-1]
    print(f"\n{'=' * 90}\nPART C: excerpts in {path.name} (saved {datetime.fromtimestamp(path.stat().st_mtime):%Y-%m-%d %H:%M})")
    wb = openpyxl.load_workbook(path, read_only=True)
    found_sent = found_full = total = 0
    misses = []
    for sheet in ("CapEx & Catalysts", "Supply Chain & Risks"):
        if sheet not in wb.sheetnames: continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        col_t, col_q = rows[0].index("Ticker"), rows[0].index("Filing Excerpt")
        for r in rows[1:]:
            t, q = r[col_t], r[col_q]
            if not q or t not in sent_by_ticker: continue
            total += 1
            pieces = [p.strip() for p in normalize(str(q)).split("...") if p.strip()]
            if all(p in sent_by_ticker[t] for p in pieces): found_sent += 1
            elif all(p in full_by_ticker[t] for p in pieces): found_full += 1
            else: misses.append((t, sheet, min(longest_run(p, sent_by_ticker[t]) for p in pieces), q))
    print(f"  {total} excerpts: {found_sent} word for word in text sent to Claude, "
          f"{found_full} only elsewhere in the full filing (not from the input), {len(misses)} not found")
    print("  Longest matching word run for misses. Rough guide, not a rule: ~60-90% suggests paraphrase, lower suggests invention.")
    for t, sheet, pct, q in misses:
        print(f"    {t} [{sheet}] {pct:.0%}: {str(q)[:160]!r}")


# --- Main ---
def run_ticker(ticker, want_excerpts):
    # COPIED: app.py lines 180-217 (filing selection)
    headers = {'User-Agent': f'InstitutionalResearch analyst_{ticker.lower()}@research-terminal.com'}
    data = sec_get("https://www.sec.gov/files/company_tickers.json", headers).json()
    cik_padded = next((str(v['cik_str']).zfill(10) for v in data.values() if v['ticker'].upper() == ticker.upper()), None)
    filings = sec_get(f"https://data.sec.gov/submissions/CIK{cik_padded}.json", headers).json()['filings']['recent']
    form_list = filings['form']
    cik = str(int(cik_padded))
    idx_curr = next((i for i, f in enumerate(form_list) if f in ['10-K', '10-Q']), -1)
    idx_prior = next((i for i in range(idx_curr + 1, len(form_list)) if form_list[i] in ['10-K', '10-Q']), -1)
    form_type = form_list[idx_curr]

    def describe(i):
        return form_list[i], filings['filingDate'][i], filings['accessionNumber'][i]

    print(f"\n{'=' * 90}\n{ticker}  (CIK {cik})")
    c_form, c_date, c_acc = describe(idx_curr)
    html_c, how_c = get_document(ticker, cik, c_acc.replace('-', ''), filings['primaryDocument'][idx_curr], headers)
    print(f"  CURRENT : {c_form} filed {c_date}, accession {c_acc} ({how_c})")
    if idx_prior != -1:
        p_form, p_date, p_acc = describe(idx_prior)
        html_p, how_p = get_document(ticker, cik, p_acc.replace('-', ''), filings['primaryDocument'][idx_prior], headers)
    else:
        p_form = p_date = p_acc = "-"; html_p, how_p = None, "no prior filing"
    comparison = ("TRUE 10-K vs 10-K" if c_form == p_form == "10-K"
                  else f"{c_form} vs {p_form}" + (" (quarter over quarter)" if c_form == p_form == "10-Q" else " (mixed)"))
    print(f"  PRIOR   : {p_form} filed {p_date}, accession {p_acc} ({how_p})")
    print(f"  COMPARISON: {comparison}")

    text_c = clean(html_c)
    cur_risk = report_extraction("CURRENT RISK FACTORS", text_c, risk_regex, RISK_FALLBACK, 20000)
    cur_mda = report_extraction("CURRENT MD&A", text_c, MDA_REGEX_10K if form_type == '10-K' else MDA_REGEX_10Q, MDA_FALLBACK, 20000, heading_word="management")
    prior_risk = None
    if html_p:
        prior_risk = report_extraction("PRIOR RISK FACTORS", clean(html_p), risk_regex, RISK_FALLBACK, 18000, cut_before_strip=True)

    # Part B extra: the two most recent 10-Ks, regardless of what the app picks
    tenk = [i for i, f in enumerate(form_list) if f == '10-K'][:2]
    tenk_results = []
    print(f"\n  --- What a 10-K vs 10-K comparison would get ---")
    for i in tenk:
        k_form, k_date, k_acc = describe(i)
        html_k, how_k = get_document(ticker, cik, k_acc.replace('-', ''), filings['primaryDocument'][i], headers)
        print(f"\n  10-K filed {k_date}, accession {k_acc} ({how_k})")
        if html_k:
            tenk_results.append((k_date, report_extraction("10-K RISK FACTORS", clean(html_k), risk_regex, RISK_FALLBACK, 20000)))

    sent_text = ""
    if want_excerpts:
        # COPIED: app.py lines 244-257 (8-K text sent to Claude: first four 8-Ks over 100 chars, 4,500 chars each)
        k8 = 0
        for i, f in enumerate(form_list):
            if f in ['8-K', '8-K/A'] and k8 < 4:
                html_8k, _ = get_document(ticker, cik, filings['accessionNumber'][i].replace('-', ''), filings['primaryDocument'][i], headers)
                if html_8k:
                    t8 = clean(html_8k)
                    if len(t8) > 100:
                        sent_text += " " + t8[:4500]; k8 += 1
        sent_text = normalize(" ".join([cur_risk["sent_text"], cur_mda["sent_text"], prior_risk["sent_text"] if prior_risk else "", sent_text]))

    return {"ticker": ticker, "current": f"{c_form} {c_date}", "prior": f"{p_form} {p_date}", "comparison": comparison,
            "cur_risk": cur_risk, "cur_mda": cur_mda, "prior_risk": prior_risk, "tenk": tenk_results,
            "sent_norm": sent_text, "full_norm": normalize(text_c) if want_excerpts else ""}


def fmt(r):
    if r is None: return "not downloaded"
    if r["fallback"]: return f"FALLBACK, {r['len']:,} -> {r['sent']:,}"
    return f"match, {'heading' if r['heading_ok'] else 'NO heading'}, ends {r['ended_by']}, {r['len']:,} -> {r['sent']:,}"

def main():
    want_excerpts = "--excerpts" in sys.argv
    tickers = sys.argv[sys.argv.index("--tickers") + 1].upper().split(",") if "--tickers" in sys.argv else TICKERS
    print(f"Risk extraction check, run {datetime.now():%Y-%m-%d %H:%M}, tickers {tickers}. No Anthropic calls.")
    results = [run_ticker(t, want_excerpts) for t in tickers]

    print(f"\n{'=' * 90}\nSUMMARY  (len -> chars that survive the cut)")
    for r in results:
        print(f"\n{r['ticker']}: current {r['current']}, prior {r['prior']} -> {r['comparison']}")
        print(f"  current risk : {fmt(r['cur_risk'])}")
        print(f"  current MD&A : {fmt(r['cur_mda'])}")
        print(f"  prior risk   : {fmt(r['prior_risk'])}")
        for date, k in r["tenk"]:
            print(f"  10-K {date}: {fmt(k)}")
        real = len(r["tenk"]) == 2 and all(not k["fallback"] and k["heading_ok"] and k["ended_by"] == "Item 1B" for _, k in r["tenk"])
        print(f"  10-K vs 10-K gets real Risk Factors from both (heuristic: matched heading, ended at Item 1B): {'YES' if real else 'NO'}")
    print(f"\nSEC requests made: {request_count[0]}")

    if want_excerpts:
        check_excerpts({r["ticker"]: r["sent_norm"] for r in results}, {r["ticker"]: r["full_norm"] for r in results})


if __name__ == "__main__":
    main()
