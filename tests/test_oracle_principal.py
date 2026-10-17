"""Oracle's August 31, 2026 principal calendar: instrument facts, stated terms and reconciliation rules.

Expected values below are typed from the printed 10-K debt table (Note 6, pp. 83–84) and the 10-Q,
independently of the XBRL reader. Raw-filing checks skip on a machine without data/raw/sec/ORCL.
"""
import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import sec_data
import oracle_principal as op

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests/fixtures"
RAW = ROOT.parent / "data/raw/sec/ORCL"
M = 1_000_000


def build():
    return op.build(FIX / "ORCL_2026_principal.htm", FIX / "ORCL_2027Q1_principal.htm")


class Calendar(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.r = build()

    def row(self, label):
        return next(i for i in self.r["instruments"] if i["label"] == label)

    def test_saved_sample_is_the_fixture_build(self):
        saved = json.loads(op.SAMPLE.read_text(encoding="utf-8"))
        saved.pop("source_sha256")
        self.assertEqual(saved, self.r)
        self.assertIsNotNone(op.load())

    def test_instruments_match_the_printed_table(self):
        self.assertEqual(self.r["status"], "Checked")
        self.assertEqual(len(self.r["instruments"]), 58)   # 54 fixed, 2 floating, term loan, commercial paper
        printed = {"$3,000M, 2.65%, due July 2026": (3_000, "2026-07-15"),
                   "$2,250M, 2.80%, due April 2027": (2_250, "2027-04-01"),
                   "$5,000M, 6.70%, due February 2056": (5_000, "2056-02-04"),
                   "$500M, Compounded SOFR plus 0.76%, due August 2028": (500, "2028-08-03"),
                   "Term loan, $5,630M borrowed, SOFR plus 1.35%, due August 2027": (5_137, "2027-08-16")}
        for label, (millions, maturity) in printed.items():
            self.assertEqual(self.row(label)["principal_may_31"], millions * M, label)
            self.assertEqual(self.row(label)["maturity"], maturity, label)
        cp = self.row("Commercial paper notes")
        self.assertEqual(cp["principal_may_31"], 1_468 * M)
        self.assertIsNone(cp["maturity"])
        self.assertEqual(sum(i["principal_may_31"] for i in self.r["instruments"]), 130_105 * M)

    def test_every_check_agrees_and_none_is_forced(self):
        statuses = {c["label"]: c["status"] for c in self.r["checks"]}
        self.assertEqual(len(statuses), 15)
        self.assertTrue(all(s in ("Agrees", "Consistent", "Consistent, not adopted") for s in statuses.values()))
        buckets = [c for c in self.r["checks"] if "maturity bucket" in c["label"]]
        self.assertEqual([c["actual"] for c in buckets],
                         [7_210_625_000, 10_144_750_000, 5_500 * M, 7_250 * M, 9_750 * M, 90_250 * M])

    def test_only_the_scheduled_note_matures_and_the_grouped_line_is_not_subtracted(self):
        self.assertEqual([m["label"] for m in self.r["matured"]], ["$3,000M, 2.65%, due July 2026"])
        rec = self.r["reconciliation"]
        self.assertEqual(rec["repayments_quarter"], 4_202 * M)
        self.assertEqual(rec["scheduled_principal_quarter"], 3_070_375_000)   # $3,000M note + June 30 installment
        # The $1,131.6M residual is compared with MD&A's ~$1.1B, never taken out of commercial paper
        self.assertIsNone(rec["commercial_paper"])
        self.assertIsNone(rec["principal_total_aug"])
        self.assertEqual(rec["commercial_paper_range"], [287 * M, 387 * M])
        self.assertEqual(rec["implied_discount_range"], [517 * M, 617 * M])

    def test_august_calendar_and_next_twelve_months(self):
        self.assertEqual([(b["label"], b["principal"]) for b in self.r["calendar"]],
                         [("Remainder FY2027", 2_672_250_000), ("FY2028", 10_144_750_000), ("FY2029", 5_500 * M),
                          ("FY2030", 7_250 * M), ("FY2031", 9_750 * M), ("Thereafter", 90_250 * M)])
        near = self.r["next_twelve_months"]
        self.assertEqual(near["principal"], 7_317 * M)
        self.assertEqual([x["date"] for x in near["rows"]],
                         ["2026-09-30", "2026-12-31", "2027-03-31", "2027-04-01", "2027-06-30", "2027-08-16"])
        self.assertEqual(self.r["reconciliation"]["known_principal_excluding_cp"],
                         sum(b["principal"] for b in self.r["calendar"]))
        self.assertEqual(self.r["reconciliation"]["known_principal_excluding_cp"], 125_567 * M)

    def test_term_loan_schedule_follows_the_stated_formula(self):
        s = op.term_loan_schedule(5_630 * M)
        self.assertEqual([i["amount"] for i in s["installments"]], [70_375_000] * 8 + [140_750_000] * 4)
        self.assertEqual(s["final"], {"date": "2027-08-16", "amount": 4_504 * M, "basis": s["final"]["basis"]})
        self.assertEqual((s["balance_may_31_2025"], s["balance_may_31_2026"], s["balance_aug_31_2026"]),
                         (5_418_875_000, 5_137_375_000, 5_067 * M))
        self.assertEqual(sum(i["amount"] for i in s["installments"]) + s["final"]["amount"], 5_630 * M)

    def test_principal_by_period_matches_the_calendar(self):
        import oracle_liquidity as ol
        parts = [op.principal_between(self.r, start, end) for _, start, end in ol.DUE]
        self.assertEqual(parts, [2_672_250_000, 10_144_750_000, 5_500 * M, 7_250 * M, 9_750 * M, 1_250 * M, 89_000 * M])
        self.assertEqual(sum(parts), self.r["reconciliation"]["known_principal_excluding_cp"])

    def test_fiscal_year_boundaries(self):
        self.assertEqual(op.bucket_of("2026-05-31"), "Fiscal 2026")
        self.assertEqual(op.bucket_of("2026-06-01"), "Fiscal 2027")
        self.assertEqual(op.bucket_of("2031-05-31"), "Fiscal 2031")
        self.assertEqual(op.bucket_of("2031-06-01"), "Thereafter")

    def test_a_changed_bucket_is_reported_not_reconciled(self):
        with mock.patch.dict(op.BUCKETS_MAY, {"Fiscal 2029": 5_600}):
            r = build()
        self.assertEqual(r["status"], "Differs")
        bad = next(c for c in r["checks"] if c["label"].startswith("Fiscal 2029"))
        self.assertEqual((bad["status"], bad["difference"]), ("Differs", -100 * M))

    def test_a_missing_fact_makes_the_calendar_unavailable(self):
        with tempfile.TemporaryDirectory() as d:
            q = Path(d) / "q.htm"
            q.write_text((FIX / "ORCL_2027Q1_principal.htm").read_text(encoding="utf-8")
                         .replace('name="us-gaap:RepaymentsOfDebt"', 'name="us-gaap:Renamed"'), encoding="utf-8")
            r = op.build(FIX / "ORCL_2026_principal.htm", q)
        self.assertEqual(r["status"], "Unavailable")
        self.assertIn("RepaymentsOfDebt", r["detail"])

    def test_load_refuses_a_calendar_that_does_not_agree(self):
        with tempfile.TemporaryDirectory() as d:
            bad = copy.deepcopy(self.r)
            bad["status"] = "Differs"
            (Path(d) / op.SAMPLE.name).write_text(json.dumps(bad), encoding="utf-8")
            self.assertIsNone(op.load(d))
            self.assertIsNone(op.load(Path(d) / "missing"))


@unittest.skipUnless((RAW / "orcl-20260531.htm").is_file(), "saved Oracle filings not on this machine")
class RawFilings(unittest.TestCase):
    def test_full_filings_rebuild_the_saved_calendar(self):
        saved = json.loads(op.SAMPLE.read_text(encoding="utf-8"))
        for name, digest in saved.pop("source_sha256").items():
            self.assertEqual(hashlib.sha256((RAW / name).read_bytes()).hexdigest(), digest)
        self.assertEqual(op.build(RAW / "orcl-20260531.htm", RAW / "orcl-20260831.htm"), saved)

    def test_cited_text_appears_in_the_filings(self):
        quarter = " ".join(sec_data.html_to_text((RAW / "orcl-20260831.htm").read_bytes()).split())
        annual = " ".join(sec_data.html_to_text((RAW / "orcl-20260531.htm").read_bytes()).split())
        for item in op.QUARTER_TEXT.values():
            self.assertTrue(item["phrase"] in quarter, item["phrase"])
        # The term-loan clauses sit in separate bullets across a page break, so each is checked on its own
        for clause in ("reduced by any prepayments multiplied by 1.25% on September 30, 2024 and quarterly thereafter until June 30, 2026",
                       "reduced by any prepayments multiplied by 2.50% on September 30, 2026 and quarterly thereafter until June 30, 2027",
                       "will become fully due and payable on August 16, 2027",
                       "may be extended at our sole option by up to 2 years"):
            self.assertTrue(clause in annual, clause)


class Workbook(unittest.TestCase):
    """The Payment dates block in the served Oracle workbook carries the same calendar as the app."""

    @classmethod
    def setUpClass(cls):
        from openpyxl import load_workbook
        path = ROOT / "samples/oracle_liquidity.xlsx"
        cls.values = load_workbook(path, data_only=True)["Payment dates"]
        cls.formulas = load_workbook(path)["Payment dates"]
        cls.r = op.load()

    def find(self, label):
        return next(r for r in range(35, self.values.max_row + 1) if self.values.cell(r, 1).value == label)

    def test_calendar_rows_match_the_model(self):
        self.assertEqual(self.values["A35"].value, "August 31, 2026 borrowing principal by instrument")
        for bucket in self.r["calendar"]:
            row = self.find(bucket["label"])
            self.assertEqual(self.values.cell(row, 3).value, bucket["principal"])
            self.assertTrue(str(self.formulas.cell(row, 3).value).startswith("=SUMIFS("))
        self.assertEqual(self.values.cell(self.find("Next twelve months, excluding commercial paper"), 3).value, 7_317 * M)
        self.assertEqual(self.values.cell(self.find("Total, excluding commercial paper"), 3).value, 125_567 * M)
        self.assertEqual(self.values.cell(self.find("Difference"), 3).value, 0)

    def test_unknowns_stay_unavailable_and_are_not_numbers(self):
        cp = self.find("Commercial paper")
        self.assertIsNone(self.values.cell(cp, 3).value)
        self.assertIn("Unavailable", self.values.cell(cp, 4).value)
        table = next(r for r in range(35, self.values.max_row + 1) if self.values.cell(r, 1).value == "Commercial paper notes")
        self.assertIsNone(self.values.cell(table, 5).value)
        low = self.find("Commercial paper implied, low")
        self.assertEqual((self.values.cell(low, 3).value, self.values.cell(low + 1, 3).value), (287 * M, 387 * M))

    def test_every_instrument_row_keeps_its_source(self):
        rows = [r for r in range(35, self.values.max_row + 1) if (self.values.cell(r, 6).value or "") in
                ("Senior notes", "Term loan", "Commercial paper")]
        self.assertEqual(len(rows), 58)
        for r in rows:
            self.assertIn("us-gaap:DebtInstrumentCarryingAmount context C_", self.values.cell(r, 4).value)
            self.assertIn(op.ANNUAL["accession"], self.values.cell(r, 4).value)

    def test_download_is_refused_when_the_calendar_changes(self):
        import oracle_liquidity as ol
        snapshot = next(s for s in json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]
                        if s["ticker"] == "ORCL")
        self.assertIsNotNone(ol.workbook_bytes(ol.for_snapshot(snapshot)))
        with tempfile.TemporaryDirectory() as d:
            changed = Path(d) / op.SAMPLE.name
            changed.write_bytes(op.SAMPLE.read_bytes() + b" ")
            with mock.patch.object(op, "SAMPLE", changed):
                self.assertIsNone(ol.workbook_bytes(ol.for_snapshot(snapshot)))


if __name__ == "__main__":
    unittest.main()
