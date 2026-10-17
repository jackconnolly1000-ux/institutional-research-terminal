"""The Oracle stress scenario in the app: Oracle only, labelled, editable, and recorded in its own workbook."""
import os
import unittest
from io import BytesIO
from unittest import mock

from openpyxl import load_workbook
from streamlit.runtime.memory_media_file_storage import MemoryMediaFileStorage
from streamlit.testing.v1 import AppTest

import liquidity_simulation as ls
import simulation_view as sv

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SERVED = {}
REAL_STORE = MemoryMediaFileStorage.load_and_get_id


def record(storage, data, mimetype, kind, filename=None):
    file_id = REAL_STORE(storage, data, mimetype, kind, filename)
    SERVED[file_id] = data
    return file_id


class ScenarioPanel(unittest.TestCase):
    def setUp(self):
        p = mock.patch.object(MemoryMediaFileStorage, "load_and_get_id", record)
        p.start()
        self.addCleanup(p.stop)
        self.app = AppTest.from_file(os.path.join(ROOT, "app.py"), default_timeout=90)
        self.app.run()

    def pick(self, ticker):
        self.app.segmented_control[0].set_value(ticker).run()
        self.assertFalse(self.app.exception, [e.value for e in self.app.exception])

    def view(self):
        return next((c for c in self.app.segmented_control if c.key == "review_view"), None)

    def scenario(self):
        self.pick("ORCL")
        self.view().set_value(sv.VIEWS[1]).run()
        self.assertFalse(self.app.exception, [e.value for e in self.app.exception])

    def text(self):
        return "\n".join([m.value for m in self.app.markdown] + [c.value for c in self.app.caption])

    def test_only_oracle_has_the_scenario_view(self):
        self.pick("ORCL")
        self.assertEqual(self.view().options, list(sv.VIEWS))
        self.assertEqual(self.view().value, "Checked review")           # the checked review opens first
        self.assertEqual(len(self.app.dataframe), 0)                       # no scenario figures in the review
        self.assertIn("View → Stress scenario (illustrative)", self.text())
        for ticker in ("LLY", "VST"):
            self.pick(ticker)
            self.assertIsNone(self.view())
            self.assertEqual(len(self.app.dataframe), 0)

    def test_scenario_view_is_separate_and_framed(self):
        self.scenario()
        text = self.text()
        self.assertIn(ls.LABEL, text)
        self.assertIn("not from checked filing facts", text)
        self.assertIn("irt-card tone-caution", text)                       # the caution banner
        self.assertNotIn("Funding conclusion for ORCL", text)              # the checked review is not mixed in
        self.assertNotIn("Download ORCL review and evidence",
                         [b.proto.label for b in self.app.get("download_button")])
        self.assertEqual(len(self.app.dataframe), 4)   # results, month ends, correlation, uses outside the model
        self.assertEqual(len(self.app.dataframe[0].value), 4)
        self.assertEqual(list(self.app.dataframe[2].value["Correlation setting"]), list(ls.COMPARE_CORRELATIONS))

    def test_scenario_edits_survive_a_visit_to_the_review(self):
        self.scenario()
        self.app.slider(key="sim_rho").set_value(0.5).run()
        self.view().set_value(sv.VIEWS[0]).run()
        self.assertIn("Funding conclusion for ORCL", self.text())
        self.view().set_value(sv.VIEWS[1]).run()
        self.assertEqual(self.app.slider(key="sim_rho").value, 0.5)

    def test_view_choice_survives_switching_company(self):
        self.scenario()
        self.pick("LLY")
        self.pick("ORCL")
        self.assertEqual(self.view().value, sv.VIEWS[1])

    def test_an_invalid_range_runs_nothing(self):
        self.scenario()
        next(n for n in self.app.number_input if n.key == "sim_ocf_low").set_value(60.0).run()
        self.assertTrue(any("low ≤ most likely ≤ high" in w.value for w in self.app.warning))
        self.assertEqual(len(self.app.dataframe), 0)

    def test_the_download_records_inputs_seed_and_results(self):
        self.scenario()
        button = next(b for b in self.app.get("download_button") if b.proto.label == "Download this scenario (Excel)")
        wb = load_workbook(BytesIO(SERVED[button.proto.url.rsplit("/", 1)[-1].split(".")[0]]))
        self.assertEqual(wb.sheetnames, ["Scenario", "Results", "Correlation", "Monthly path", "Monthly path by month",
                                         "Outside-use effects", "Uses not modelled"])
        self.assertEqual([wb["Correlation"].cell(r, 2).value for r in (2, 3, 4)], list(ls.COMPARE_CORRELATIONS))
        self.assertEqual(wb["Scenario"]["A2"].value, ls.LABEL)
        self.assertIn("Seed 2026, 10,000 trials", wb["Scenario"]["A3"].value)
        self.assertEqual(wb["Results"].max_row, 5)
        expected = ls.simulate(ls.oracle_inputs(), ls.oracle_decisions()[0], 10_000, 2026)
        self.assertAlmostEqual(wb["Results"]["E2"].value, expected["mean"], delta=1)
        self.assertAlmostEqual(wb["Results"]["J2"].value, expected["share_below_zero"])

    def test_the_correlation_setting_reaches_every_result_and_the_download(self):
        self.scenario()
        self.app.slider(key="sim_rho").set_value(0.5).run()
        self.assertFalse(self.app.exception)
        button = next(b for b in self.app.get("download_button") if b.proto.label == "Download this scenario (Excel)")
        wb = load_workbook(BytesIO(SERVED[button.proto.url.rsplit("/", 1)[-1].split(".")[0]]))
        self.assertEqual(wb["Scenario"]["C12"].value, 0.5)
        from dataclasses import replace
        tied = replace(ls.oracle_inputs(), correlation=0.5)
        expected = ls.simulate(tied, ls.oracle_decisions()[0], 10_000, 2026)
        self.assertAlmostEqual(wb["Results"]["E2"].value, expected["mean"], delta=1)
        self.assertAlmostEqual(wb["Results"]["F2"].value, expected["sd"], delta=1)

    def test_panel_separates_facts_from_assumptions_and_states_its_limits(self):
        self.scenario()
        text = "\n".join([m.value for m in self.app.markdown] + [c.value for c in self.app.caption])
        self.assertIn("**Filing facts:**", text)
        self.assertIn("**Assumptions:**", text)
        self.assertIn("each payment date within a month", text)
        self.assertIn("not probabilities of default", text)
        self.assertNotIn("probability", " ".join(self.app.dataframe[0].value.columns).lower())
        uses = self.app.dataframe[3].value
        self.assertEqual(list(uses["Scenario"]), ["Base model", "Finance-lease payments (principal proxy)", "Short-term capex financing (cumulative proceeds proxy)",
                                                 "Purchase and other obligations", ls.ALL_THREE])
        self.assertIn("Sensitivity amount ($B)", uses.columns)
        self.assertIn("not established balances or payments due", text)
        self.assertIn("financing cash flow", text)                     # why finance-lease principal is outside OCF

    def test_download_records_every_use_left_out_and_its_effect(self):
        self.scenario()
        button = next(b for b in self.app.get("download_button") if b.proto.label == "Download this scenario (Excel)")
        wb = load_workbook(BytesIO(SERVED[button.proto.url.rsplit("/", 1)[-1].split(".")[0]]))
        effects = list(wb["Outside-use effects"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(effects), 4 * 5)                          # every decision, base + three + all
        self.assertEqual({r[1] for r in effects if r[2] == 0}, {"Base model"})
        notes = list(wb["Uses not modelled"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(notes), 3 + len(ls.EXCLUDED) + len(ls.UNQUANTIFIED))
        kinds = [wb["Scenario"].cell(r, 2).value for r in range(6, 13)]
        self.assertEqual(kinds[:3], ["Filing fact", "Filing fact", "Assumption"])
        everything = " ".join(str(c.value) for s in wb for row in s.iter_rows() for c in row if c.value)
        self.assertNotIn("Shortfall probability", everything)
        self.assertNotIn("finance lease payments (inside OCF", everything)   # the corrected claim
        self.assertNotIn("Added use", everything)
        self.assertIn(ls.SENSITIVITY_NOTE, everything)
        self.assertEqual(wb["Outside-use effects"]["C1"].value, "Sensitivity amount (USD, displayed in millions)")


    def test_month_end_table_sits_beside_the_end_of_window_result(self):
        self.scenario()
        table = self.app.dataframe[1].value
        self.assertEqual(list(table["Cash at"]), ["Aug 31, 2027", "Lowest month end"])
        end = ls.simulate(ls.oracle_inputs(), ls.oracle_decisions()[0], 10_000, 2026)
        path = ls.monthly_path(ls.oracle_inputs(), ls.oracle_decisions()[0], 10_000, 2026)
        self.assertAlmostEqual(table[sv.SHARE][0], end["share_below_zero"] * 100)
        self.assertAlmostEqual(table[sv.SHARE][1], path["share_below_zero_any_month"] * 100)
        self.assertAlmostEqual(table["Mean ($B)"][1], round(path["lowest_mean"] / ls.B, 2))
        captions = " ".join(c.value for c in self.app.caption)
        self.assertIn("Aug 2027 in 100.0% of trials", captions)
        self.assertIn(r"Apr 2027 \$2.250B", captions)
        self.assertIn("September 2026", captions)                      # where commercial paper is assumed due
        self.assertNotIn("probability", " ".join(table.columns).lower())

    def test_download_records_the_month_end_path(self):
        self.scenario()
        button = next(b for b in self.app.get("download_button") if b.proto.label == "Download this scenario (Excel)")
        wb = load_workbook(BytesIO(SERVED[button.proto.url.rsplit("/", 1)[-1].split(".")[0]]))
        summary = list(wb["Monthly path"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(summary), 4)
        path = ls.monthly_path(ls.oracle_inputs(), ls.oracle_decisions()[0], 10_000, 2026)
        self.assertAlmostEqual(summary[0][3], path["share_below_zero_any_month"])
        self.assertAlmostEqual(summary[0][5], path["lowest_mean"], delta=1)
        self.assertEqual(wb["Monthly path"]["D2"].number_format, "0.0%")
        self.assertEqual(wb["Monthly path"]["F2"].number_format, '#,##0,,;(#,##0,,);"-"')
        by = wb["Monthly path by month"]
        self.assertEqual(by.max_row, 1 + 4 * 12)
        self.assertIn("(USD, displayed in millions)", by["D1"].value)
        april = next(r for r in by.iter_rows(min_row=2, max_row=13, values_only=True) if r[1] == "Apr 2027")
        self.assertEqual(april[3], 2_250_000_000)
        self.assertIsInstance(april[4], float)
        self.assertEqual(by["C2"].number_format, "yyyy-mm-dd")
        self.assertEqual((by["E2"].font.name, by["E2"].font.sz), ("Arial", 10.0))
        self.assertEqual(wb["Scenario"]["A17"].value, ls.MONTHLY_BASIS)


class PrincipalCalendarWording(unittest.TestCase):
    """Audit (Oct 17): no Oracle text may still say August 31 principal is unknown now that the calendar is checked."""

    def test_review_and_combined_workbook_point_to_the_calendar(self):
        import zipfile
        import company_review
        with open(os.path.join(ROOT, "oracle_liquidity_view.py"), encoding="utf-8") as f:
            source = f.read()
        self.assertNotIn("Unavailable: remaining debt principal at August 31", source)
        self.assertIn("is dated in the principal calendar above", source)
        self.assertNotIn("does not establish remaining August 31 payments", company_review.LIMITS["ORCL"])
        about = zipfile.ZipFile(os.path.join(ROOT, "samples", "liquidity_pilots.xlsx")).read("xl/worksheets/sheet3.xml")
        self.assertNotIn(b"August schedule and revolver availability are Unavailable", about)
        self.assertIn(b"August 31 note and term-loan principal is dated", about)


class ScenarioWorkbook(unittest.TestCase):
    def test_workbook_is_arial_10_with_numbers_as_numbers(self):
        i, decisions = ls.oracle_inputs(), ls.oracle_decisions()
        results = [(ls.simulate(i, d, 1_000, 1), d) for d in decisions]
        wb = load_workbook(BytesIO(sv.workbook(i, results, 1_000, 1)))
        cell = wb["Results"]["E2"]
        self.assertIsInstance(cell.value, float)
        self.assertEqual((cell.font.name, cell.font.sz), ("Arial", 10.0))
        self.assertEqual(wb["Scenario"]["A14"].value, ls.TIMING_LIMIT)
        self.assertEqual(wb["Scenario"]["A15"].value, sv.SHARE_NOTE)
        self.assertEqual(wb["Scenario"]["C12"].number_format, "+0.0;-0.0;0.0")   # a ratio, not $ millions
        self.assertNotIn("Correlation", wb.sheetnames)                            # only when a comparison is passed


if __name__ == "__main__":
    unittest.main()
