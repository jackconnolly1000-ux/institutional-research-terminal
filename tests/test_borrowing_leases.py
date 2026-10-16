"""Presentation keeps current borrowing, dated leases and future commitments distinct."""
import html
import json
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import borrowing_leases as bl
import company_review as view
import oracle_liquidity as ol
import vistra_liquidity as vl

ROOT = Path(__file__).resolve().parents[1]


class BorrowingLeasePresentation(unittest.TestCase):
    def setUp(self):
        self.snapshots = json.loads((ROOT / "samples/liquidity_pilots.json").read_text())["snapshots"]

    def test_checked_components_and_lease_dates(self):
        expected = {"LLY": (54908e6, [1362e6, None], ["2025-12-31", "2026-06-30"]),
                    "ORCL": (125337e6, [34621e6, 9185e6], ["2026-08-31"] * 2),
                    "VST": (19895e6, [105e6, 222e6], ["2025-12-31"] * 2)}
        for s in self.snapshots:
            m = bl.build(s)
            total, leases, dates = expected[s["ticker"]]
            self.assertEqual(m["baseline"], total)
            self.assertEqual(sum(f["value"] for f in m["borrowing"]), total)
            self.assertEqual([f["value"] for f in m["leases"]], leases)
            self.assertEqual([f["period_end"] for f in m["leases"]], dates)
            self.assertEqual(m["commitment"]["value"], 288e9 if s["ticker"] == "ORCL" else None)
            self.assertEqual(m["latest_leases_available"], s["ticker"] == "ORCL")

    def test_unavailable_is_not_zero_and_old_dates_remain_visible(self):
        for s in (self.snapshots[0], self.snapshots[2]):
            text = html.unescape(view.render(s))
            self.assertIn("Lease liabilities at 2026-06-30: Unavailable", text)
            self.assertIn("As of 2025-12-31", text)
            self.assertIn("No checked amount for leases not yet commenced", text)
            self.assertNotIn("$0.000B", text)
        self.assertEqual(bl.build(self.snapshots[2])["borrowing"][2]["value"], 0)

    def test_no_new_instrument_breakdown_or_lease_addition(self):
        m = bl.build(self.snapshots[1])
        self.assertEqual([f["label"] for f in m["borrowing"]], ["Current borrowing", "Noncurrent borrowing"])
        text = html.unescape(view.render(self.snapshots[1]))
        self.assertIn("does not split these borrowing lines by instrument", text)
        self.assertIn("Neither is added to baseline gross debt", text)

    def test_mismatched_anchor_or_borrowing_pair_is_refused(self):
        with self.assertRaises(KeyError):
            bl.build(self.snapshots[0], {})
        for key in ("accession", "period_end", "unit"):
            s = deepcopy(self.snapshots[0]); s["metrics"]["gross_debt"]["inputs"][0][key] = "wrong"
            with self.assertRaises(ValueError):
                bl.build(s)
        s = deepcopy(self.snapshots[0]); s["metrics"]["cash"]["accession"] = "wrong"
        with self.assertRaises(ValueError):
            bl.build(s)

    def test_lease_total_or_unit_conflict_is_refused(self):
        for key, value in (("value", 1), ("unit", "EUR"), ("dimensions", [["axis", "member"]]),
                           ("period", "2026-05-31")):
            evidence = json.loads(bl.SAMPLE.read_text())
            evidence["companies"]["ORCL"]["facts"]["OperatingLeaseLiability"][key] = value
            with self.assertRaises(ValueError):
                bl.build(self.snapshots[1], evidence)
        evidence = json.loads(bl.SAMPLE.read_text())
        del evidence["companies"]["ORCL"]["facts"]["FinanceLeaseLiabilityCurrent"]
        with self.assertRaises(ValueError):
            bl.build(self.snapshots[1], evidence)

    def test_individual_workbooks_match_the_same_checked_model(self):
        for s in self.snapshots:
            ticker = s["ticker"]
            stem = {"LLY": "lly", "ORCL": "oracle", "VST": "vistra"}[ticker]
            path = ROOT / f"samples/{stem}_liquidity.xlsx"
            v, f = load_workbook(path, data_only=True), load_workbook(path, data_only=False)
            sheet, formulas = v["Borrowing and leases"], f["Borrowing and leases"]
            records = {r[0].value: r for r in sheet if r[0].value}
            m = bl.build(s)
            self.assertEqual(records["Baseline gross debt"][1].value, m["baseline"])
            for lease in m["leases"]:
                label = lease["label"] + (" — total" if lease["value"] is not None else "")
                row = records[label]
                self.assertEqual(row[1].value, lease["value"] if lease["value"] is not None else "Unavailable")
                self.assertEqual(row[2].value.strftime("%Y-%m-%d"), lease["period_end"])
                if lease["value"] is not None:
                    self.assertTrue(formulas.cell(row[0].row, 2).value.startswith("=B"))
            self.assertFalse(sheet.merged_cells)
            self.assertEqual(records["Leases not yet commenced"][1].value, m["commitment"]["value"] or "Unavailable")
            self.assertEqual(json.loads(path.with_suffix(".manifest.json").read_text())["borrowing_leases_sha256"], bl.fingerprint())
            v.close(); f.close()

    def test_new_evidence_changes_invalidate_all_individual_downloads(self):
        with patch.object(bl, "fingerprint", return_value="changed"):
            self.assertIsNone(view.saved_workbook(self.snapshots, ticker="LLY"))
            self.assertIsNone(ol.workbook_bytes(json.loads(ol.SAMPLE.read_text())))
            self.assertIsNone(vl.workbook_bytes(json.loads(vl.SAMPLE.read_text())))
