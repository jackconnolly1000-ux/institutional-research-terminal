"""Release 2, steps 2–3: the Oracle stress scenario in the app, with editable assumptions and an Excel record.

The model is liquidity_simulation.py. This view shows the math model and the assumption basis before any result
(as in class), runs the four decisions side by side, shows a stability check, and offers the scenario as its own
workbook. It is kept out of the checked evidence workbooks on purpose: assumptions are not filing facts.
"""
from __future__ import annotations

from dataclasses import replace
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import liquidity_simulation as ls

B = 1_000_000_000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _billions(st, label, value, key, help=None):
    return st.number_input(label, value=round(value / B, 3), step=0.5, format="%.3f", key=key, help=help) * B


def assumptions(st, inputs, decisions):
    """Editable inputs, defaulting to the checked basis. Returns (inputs, decisions, trials, seed) or None."""
    st.markdown("**Assumptions** · USD billions · defaults come from the checked filings")
    c1, c2, c3 = st.columns(3)
    t = inputs.ocf_ex_prepayments
    low = _billions(c1, "OCF excl. prepayments, low", t.low, "sim_ocf_low")
    mode = _billions(c2, "Most likely", t.mode, "sim_ocf_mode")
    high = _billions(c3, "High", t.high, "sim_ocf_high")
    st.caption(t.basis)
    c4, c5 = st.columns(2)
    prepay = _billions(c4, "Prepayments, high (low and most likely are $0)", inputs.prepayments.high, "sim_prepay")
    dividends = _billions(c5, "Dividends over the year (a decision)", decisions[0].dividends, "sim_div",
                          help="Default: Jun–Aug 2026 dividends of $1.565B, times four.")
    c6, c7 = st.columns(2)
    capex_a = _billions(c6, "Capex, decision A", decisions[0].capex, "sim_capex_a", help="Default: TTM cash capex.")
    capex_b = _billions(c7, "Capex, decision B", decisions[2].capex, "sim_capex_b", help="Default: fiscal 2026 cash capex.")
    c8, c9 = st.columns(2)
    trials = c8.selectbox("Trials", [10_000, 50_000], key="sim_trials")
    seed = int(c9.number_input("Seed", value=ls.DEFAULT_SEED, step=1, key="sim_seed"))
    try:
        edited = replace(inputs, ocf_ex_prepayments=ls.Triangular(low, mode, high, t.basis),
                         prepayments=ls.Triangular(0, 0, max(prepay, 0), inputs.prepayments.basis))
    except ValueError:
        st.warning("The OCF range needs low ≤ most likely ≤ high. Nothing was run.")
        return None
    choices = [ls.Decision(f"{name}{', revolver assumed usable' if rev else ''}", capex, dividends, rev)
               for name, capex in (("Decision A", capex_a), ("Decision B", capex_b)) for rev in (False, True)]
    return edited, choices, trials, seed


def results_table(st, results):
    import pandas as pd
    rows = [{"Decision": r["decision"], "Capex ($B)": round(d.capex / B, 2),
             "Mean ending cash ($B)": round(r["mean"] / B, 2), "10th percentile ($B)": round(r["p10"] / B, 2),
             "Shortfall probability": r["probability_shortfall"] * 100,
             "Mean shortfall when short ($B)": round(r["mean_shortfall"] / B, 2)}
            for r, d in results]
    st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch",
                 column_config={"Shortfall probability": st.column_config.NumberColumn(format="%.1f%%")})


def histogram(st, result):
    import plotly.graph_objects as go
    import theme as T
    from charts import CHART_CONFIG
    fig = go.Figure(go.Histogram(x=result["ending_cash"] / B, nbinsx=60, marker_color=T.FLOW_OUT,
                                 hovertemplate="%{x:.1f}B: %{y} trials<extra></extra>"))
    fig.add_vline(x=0, line_color=T.TEXT_STRONG, line_width=1)
    fig.add_annotation(x=0, y=1, yref="paper", text="zero", showarrow=False, yanchor="bottom",
                       font=dict(color=T.TEXT_MUTED, size=11))
    fig.update_layout(height=280, margin=dict(l=56, r=12, t=24, b=40), paper_bgcolor="rgba(0,0,0,0)",
                      plot_bgcolor="rgba(0,0,0,0)", bargap=0.04, showlegend=False,
                      font=dict(color=T.TEXT, size=11),
                      xaxis=dict(title="Ending cash, USD billions", gridcolor=T.BORDER, zerolinecolor=T.BORDER),
                      yaxis=dict(title="Trials", gridcolor=T.BORDER))
    st.plotly_chart(fig, width="stretch", theme=None, config=CHART_CONFIG)


def workbook(inputs, results, trials, seed) -> bytes:
    """The scenario on screen as its own record: inputs with sources, seed, trials and results. Arial 10."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Scenario"
    stamp = datetime.now(ZoneInfo("America/New_York")).strftime("%Y-%m-%d %H:%M ET")
    ws.append(["Oracle liquidity stress scenario"])
    ws.append([ls.LABEL])
    ws.append([f"Horizon {ls.HORIZON[0]} to {ls.HORIZON[1]}. Created {stamp}. Seed {seed}, {trials:,} trials; "
               "rerunning with the same inputs and seed reproduces these results."])
    ws.append([])
    ws.append(["Input", "Low (USD)", "Most likely (USD)", "High (USD)", "Basis / source"])
    t, p, c = inputs.ocf_ex_prepayments, inputs.prepayments, inputs.commercial_paper_due
    rows = [("Starting cash", inputs.starting_cash, inputs.starting_cash, inputs.starting_cash, inputs.sources["starting_cash"]),
            ("Dated debt principal due", inputs.principal_due, inputs.principal_due, inputs.principal_due, inputs.sources["principal_due"]),
            ("OCF excluding prepayments", t.low, t.mode, t.high, t.basis),
            ("Specified customer prepayments", p.low, p.mode, p.high, p.basis),
            ("Commercial paper due", c.low, c.mode, c.high, c.basis),
            ("Revolver commitment (hypothetical)", inputs.revolver_commitment, inputs.revolver_commitment,
             inputs.revolver_commitment, inputs.sources["revolver_commitment"])]
    for row in rows:
        ws.append(list(row))
    ws.append([])
    ws.append(["Not modelled: operating and finance lease payments (inside OCF; principal split not disclosed), "
               "purchase obligations (overlap capex), interest (inside OCF), new borrowing or equity, and the $288B "
               "of leases not yet commenced (no payment dates)."])
    res = wb.create_sheet("Results")
    res.append(["Decision", "Capex (USD)", "Dividends (USD)", "Revolver assumed usable", "Mean ending cash (USD)",
                "Std dev (USD)", "10th percentile (USD)", "Median (USD)", "90th percentile (USD)",
                "Shortfall probability", "Mean shortfall when short (USD)"])
    for r, d in results:
        res.append([r["decision"], d.capex, d.dividends, "Yes" if d.revolver else "No", r["mean"], r["sd"], r["p10"],
                    r["p50"], r["p90"], r["probability_shortfall"], r["mean_shortfall"]])
    money_fmt = '#,##0,,;(#,##0,,);"-"'
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10, bold=cell.row in ((1, 5) if sheet.title == "Scenario" else (1,)))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float | int) and not isinstance(cell.value, bool):
                    cell.number_format = "0.0%" if (sheet.title == "Results" and cell.column == 10) else money_fmt
        sheet.freeze_panes = "A2" if sheet.title == "Results" else None
        for col in "ABCDEFGHIJK":
            sheet.column_dimensions[col].width = 22
    wb["Scenario"].column_dimensions["E"].width = 70
    wb["Scenario"].column_dimensions["A"].width = 34
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def render(st, snapshot) -> None:
    """Oracle only: the other reviews have no dated principal calendar, so their obligations side would be invented."""
    if snapshot.get("ticker") != "ORCL":
        return
    inputs = ls.oracle_inputs()
    if inputs is None:
        st.caption("Stress scenario Unavailable: the checked Oracle evidence it reads did not load.")
        return
    with st.expander("Stress scenario (illustrative) · Sep 2026–Aug 2027"):
        st.markdown(f"**{ls.LABEL}**")
        st.markdown(
            "**Question:** if operating cash flow varied within the range below, how often would August 31 cash plus that "
            "cash flow fall short of the cash uses the filings date?\n\n"
            "**Model, per trial:** ending cash = $36.369B cash + OCF excluding prepayments + prepayments − capex − "
            "dividends − $7.317B dated debt principal − commercial paper due (+ a revolver draw up to $10B, only where "
            "the decision assumes it is usable; drawable credit is Unavailable in the filings). OCF, prepayments and "
            "commercial paper are uncertain (triangular); capex and dividends are decisions.".replace("$", "\\$"))
        chosen = assumptions(st, inputs, ls.oracle_decisions())
        if chosen is None:
            return
        edited, decisions, trials, seed = chosen
        results = [(ls.simulate(edited, d, trials, seed), d) for d in decisions]
        st.markdown("**Results by decision**")
        results_table(st, results)
        pick = st.radio("Show the distribution for", [d.name for d in decisions], horizontal=True, key="sim_pick")
        histogram(st, next(r for r, d in results if d.name == pick))
        check = ls.simulate(edited, decisions[0], trials * 5, seed + 1)
        drift = abs(check["probability_shortfall"] - results[0][0]["probability_shortfall"])
        st.caption(f"Stability check: rerunning decision A with {trials * 5:,} trials and a new seed changes the "
                   f"shortfall probability by {drift:.1%}. The logic, boundary, distribution and stability checks "
                   "from class also run as automated tests.")
        st.caption("Not modelled: lease payments and interest (already inside OCF), purchase obligations (overlap "
                   "capex), new borrowing or equity, and the $288B of leases not yet commenced. The funding "
                   "conclusion above is unchanged: future funding coverage remains Unavailable.".replace("$", "\\$"))
        st.download_button("Download this scenario (Excel)", workbook(edited, results, trials, seed),
                           file_name=f"Oracle_Stress_Scenario_seed{seed}.xlsx", mime=XLSX, key="sim_download")
