"""Financial questions first; filing mechanics remain available in the evidence disclosure.

The notes are selected, manually checked disclosures, not an exhaustive obligations list. Each belongs to the
pilot's specified 10-Q and remains separate from borrowing and ratios (docs/pilot_evidence_*.md).
"""
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

from ui import esc
from oracle_liquidity_view import CSS as ORACLE_LIQUIDITY_CSS
from vistra_liquidity_view import CSS as VISTRA_LIQUIDITY_CSS

NOTES = {
    "LLY": [
        {"name": "Pending AtaiBeckley acquisition", "amount": "Approximately $2.8B",
         "date": "2026-08-05", "kind": "Conditional closing payment, disclosed after quarter end", "location": "10-Q MD&A, p. 31",
         "detail": "Potential payment at closing, subject to closing conditions. Not recognized as borrowing."},
        {"name": "Venture fund commitments", "amount": "Approximately $806M",
         "date": "2026-06-30", "kind": "Unfunded commitment", "location": "10-Q MD&A, p. 31; Note 8",
         "detail": "Expected to be paid over up to ten years. Not recognized as borrowing."},
    ],
    "ORCL": [
        {"name": "Leases not yet commenced", "amount": "$288B",
         "date": "2026-08-31", "kind": "Future commitment", "location": "10-Q Note 6, p. 13",
         "detail": "Substantially all data centers; expected to commence in fiscal 2027–2029 with 15–19 year terms. Not on the balance sheet or in the lease maturities table."},
        {"name": "Operating lease liabilities", "amount": "$34.621B",
         "date": "2026-08-31", "kind": "Recognized liability", "location": "10-Q Note 6, p. 12",
         "detail": "Recognized leases, separate from borrowing and from leases not yet commenced."},
        {"name": "Finance lease liabilities", "amount": "$9.185B",
         "date": "2026-08-31", "kind": "Recognized liability", "location": "10-Q Note 6, p. 12",
         "detail": "Shown separately under this review's borrowing definition; excluded from baseline gross debt."},
        {"name": "Purchase and other obligations", "amount": "$34.150B",
         "date": "2026-08-31", "kind": "Unrecorded commitment", "location": "10-Q Note 6, p. 13",
         "detail": "Primarily cloud infrastructure components and data center power. Overlap with other commitments is not established; no combined total is presented."},
    ],
    "VST": [
        {"name": "Nuveen final payment", "amount": "$631M due Dec 31, 2026",
         "date": "2026-06-30", "kind": "Recognized purchase obligation", "location": "10-Q Note 12, p. 28",
         "detail": "Principal and interest on the minority-interest purchase. The current balance-sheet obligation is $613M. Both describe the same obligation; neither is added to baseline borrowing."},
        {"name": "Helix initial commitment", "amount": "$500M",
         "date": "2026-06-30", "kind": "Callable investment commitment", "location": "10-Q MD&A, p. 69; Note 7, p. 15",
         "detail": "Remains callable by the fund's general partner, who determines timing and amount. Amount funded to date is not separately disclosed."},
        {"name": "Helix additional commitment", "amount": "Up to $500M",
         "date": "2026-06-30", "kind": "Conditional investment commitment", "location": "10-Q MD&A, p. 69; Note 7, p. 15",
         "detail": "The specified milestones had not occurred at June 30. Vistra may also elect to fund it regardless of those milestones."},
    ],
}

LIMITS = {
    "LLY": "June lease liabilities, updated principal maturities, drawable credit and detailed payment dates remain Unavailable. The selected acquisitions and manufacturing review is not a complete review of contingencies, tax, collaboration or investment obligations.",
    "ORCL": "Commercial paper is not separately disclosed in this 10-Q. The capex-financing balance and the later status of the 10-K's $3.3B lessor guarantee remain unresolved. The May 31 maturity schedule does not establish remaining August 31 payments or current revolver availability.",
    "VST": "June 30 lease liabilities are Unavailable. Derivatives, litigation and the remaining retirement-obligation detail need further review. Credit-support overlaps are unresolved; the selected notes above are not a complete obligations list.",
}

CSS = """<style>
.lr-review {max-width:1180px; margin:8px auto 0; color:var(--text); font-variant-numeric:tabular-nums;}
.lr-review * {box-sizing:border-box;}
.lr-review h2 {font-size:2rem;line-height:1.2;letter-spacing:-.025em;padding:0;margin:6px 0 12px;color:var(--text-strong);}
.lr-review h3 {font-size:1.25rem;line-height:1.35;margin:0 0 12px;padding:0;color:var(--text-strong);}
.lr-review p {font-size:.9rem;line-height:1.55;margin:8px 0;}
.lr-review .lr-meta {font-size:.78rem;color:var(--text-muted);line-height:1.5;}
.lr-lead {border-bottom:1px solid var(--border);padding:0 0 18px;}
.lr-lead .lr-answer {font-size:1.12rem;max-width:850px;color:var(--text-strong);margin-top:12px;}
.lr-summary {display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:16px;padding:18px 0;border-bottom:1px solid var(--border);}
.lr-stat-label {font-size:.8rem;color:var(--text-muted);}
.lr-stat-value {font-size:1.5rem;font-weight:650;line-height:1.5;color:var(--text-strong);}
.lr-stat-detail {font-size:.75rem;color:var(--text-muted);}
.lr-grid {display:grid;grid-template-columns:1.15fr 1fr;gap:30px;margin:24px 0;}
.lr-section {min-width:0;}
.lr-chart {margin:18px 0;}
.lr-bar-row {display:grid;grid-template-columns:115px minmax(40px,1fr) 85px;gap:12px;align-items:center;margin:12px 0;font-size:.8rem;}
.lr-bar-value {text-align:right;color:var(--text-strong);white-space:nowrap;}
.lr-track {display:grid;grid-template-columns:1fr 1fr;height:19px;}
.lr-half {display:flex;align-items:center;min-width:0;}
.lr-half:first-child {justify-content:flex-end;border-right:1px solid var(--border-strong);}
.lr-bar {display:block;height:12px;background:var(--text-muted);min-width:1px;}
.lr-row {display:flex;justify-content:space-between;gap:16px;border-bottom:1px solid var(--border);padding:10px 0;font-size:.85rem;}
.lr-row span {max-width:75%;}
.lr-row strong {white-space:nowrap;font-weight:600;}
.lr-ratios {display:grid;grid-template-columns:1fr 1fr;gap:18px;border-top:1px solid var(--border);padding-top:14px;margin-top:20px;}
.lr-ratios strong {display:block;font-size:1.2rem;color:var(--text-strong);}
.lr-obligations {border-top:1px solid var(--border);padding:22px 0 8px;}
.lr-obligation {display:grid;grid-template-columns:minmax(160px,1fr) minmax(0,2fr);gap:24px;border-bottom:1px solid var(--border);padding:14px 0;}
.lr-obligation p {margin:0 0 5px;}
.lr-obligation strong {color:var(--text-strong);}
.lr-limits {padding:15px 0 4px;}
.lr-maturity {border-top:1px solid var(--border);padding:22px 0;}
.lr-maturity-row {display:grid;grid-template-columns:85px minmax(20px,1fr) 90px;gap:12px;align-items:center;margin:15px 0;font-size:.82rem;}
.lr-maturity-track {height:14px;background:var(--surface-raised);}
.lr-maturity-bar {display:block;height:14px;background:var(--text-muted);}
.lr-event {display:grid;grid-template-columns:145px minmax(0,1fr);gap:20px;border-top:1px solid var(--border);padding:14px 0;}
.lr-event p {margin:0 0 7px;}
.lr-review a {color:var(--link);}
.lr-history-scroll {overflow-x:auto;}
.lr-history table {width:100%;border-collapse:collapse;font-size:.85rem;}
.lr-history th,.lr-history td {padding:10px 8px;border-bottom:1px solid var(--border);text-align:right;}
.lr-history th:first-child,.lr-history td:first-child {text-align:left;}
.lr-review + .irt-evidence {max-width:1180px;margin:12px auto;}
.lr-funding {padding:22px 0;border-bottom:1px solid var(--border);}
.lr-funding-row {display:grid;grid-template-columns:165px minmax(0,1fr);gap:20px;padding:12px 0;border-top:1px solid var(--border);}
.lr-funding-row h4 {font-size:.9rem;margin:8px 0;font-weight:600;}
.lr-funding-row p {margin-top:6px;}
.lr-funding .lr-meta {overflow-wrap:anywhere;}
.lr-funding details {margin-top:14px;font-size:.85rem;}
@media(max-width:760px) {.lr-funding-row {grid-template-columns:1fr;gap:0;}}
@media(max-width:760px) {.lr-grid {grid-template-columns:1fr;gap:28px;} .lr-summary {grid-template-columns:repeat(2,minmax(0,1fr));} .lr-obligation,.lr-event {grid-template-columns:1fr;gap:4px;} .lr-review h2 {font-size:1.7rem;} .lr-review h3 {font-size:1.15rem;} .lr-bar-row {grid-template-columns:96px minmax(30px,1fr) 75px;gap:7px;} .lr-stat-value {font-size:1.3rem;}}
</style>"""

# Nested style blocks are parsed as raw text by the Markdown renderer and can hide headings.
CSS = CSS.replace("</style>", ORACLE_LIQUIDITY_CSS + VISTRA_LIQUIDITY_CSS + "</style>")


def money(value, digits=2):
    if value is None:
        return "Unavailable"
    magnitude = Decimal(str(abs(value)))
    if 0 < magnitude < 1_000_000_000:
        millions = (magnitude / 1_000_000).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        return f"{'−' if value < 0 else ''}${millions:,.0f}M"
    billions = (magnitude / 1_000_000_000).quantize(Decimal(10) ** -digits, rounding=ROUND_HALF_UP)
    return f"{'−' if value < 0 else ''}${billions:,.{digits}f}B"


def pct(metric):
    return "Unavailable" if metric["value"] is None else f"{metric['value']:.1%}"


def headline(snapshot):
    m = snapshot["metrics"]
    fcf = m["free_cash_flow"]["value"]
    if fcf is None:
        return "TTM cash generation could not be established from the checked inputs."
    if fcf < 0:
        return f"TTM reported capital spending exceeded operating cash flow by {money(-fcf)}."
    if fcf == 0:
        return "TTM operating cash flow equalled reported capital spending."
    return f"TTM operating cash flow exceeded reported capital spending by {money(fcf)}."


def maturity_html(snapshot):
    """A historical principal schedule with subsequent disclosures kept on their actual dates."""
    import oracle_financing as of
    review = of.for_snapshot(snapshot)
    if review is None:
        return ""
    if review["status"] != "Checked":
        return ('<section class="lr-maturity"><h3>Debt maturity review</h3>'
                f'<p>Unavailable: {esc(review["detail"])}</p></section>')
    output = '<section class="lr-maturity"><h3>When was borrowing scheduled to be repaid?</h3>'
    output += (f'<p class="lr-meta">Historical schedule at May 31, 2026 · principal in USD billions · '
               f'<a href="{esc(review["source_url"])}" target="_blank" rel="noopener noreferrer">10-K Note 6, p. 84</a>'
               ' · fiscal years end May 31</p>')
    output += ('<p>The saved August 31 10-Q does not disclose an updated debt maturity schedule. '
               'Fiscal 2027 below includes June–August, which had already elapsed by that balance-sheet date.</p>')
    output += '<div class="lr-grid"><div>'
    maximum = max(f["value"] for f in review["buckets"][:5]) or 1
    for fact in review["buckets"][:5]:
        output += (f'<div class="lr-maturity-row"><span>{esc(fact["label"].replace("Fiscal", "FY"))}</span>'
                   f'<div class="lr-maturity-track" aria-hidden="true"><span class="lr-maturity-bar" '
                   f'style="width:{fact["value"] / maximum * 100:.4f}%"></span></div>'
                   f'<strong class="lr-bar-value">{esc(money(fact["value"], 3))}</strong></div>')
    output += (f'<p><strong>{esc(money(review["buckets"][5]["value"], 3))} thereafter</strong>'
               ' · fiscal 2032 onward, combined. This is not a single-year payment.</p></div><div>')
    share = f'{review["first_five_share"]:.1%}' if review["first_five_share"] is not None else "Unavailable"
    output += (f'<p>{esc(money(review["first_five_years"], 3))}, or {share} '
               f'of the {esc(money(review["total"], 3))} principal total, was scheduled through fiscal 2031 '
               'at May 31.</p>')
    c = review["controls"]
    output += (f'<p class="lr-meta">The schedule excludes interest and leases. Principal of {esc(money(review["total"], 3))} less '
               f'{esc(money(c["discount_and_costs"]["value"], 3))} of discounts and issuance costs equals '
               f'{esc(money(c["carrying_debt"]["value"], 3))} carrying debt at May 31. '
               'This reconciles to the annual balance sheet, not the later August balance sheet.</p>')
    cash = snapshot["metrics"]["cash"]["value"]
    current = snapshot["metrics"]["gross_debt"]["inputs"][0]["value"]
    output += (f'<p>At August 31, the 10-Q reports {esc(money(cash))} cash and '
               f'{esc(money(current, 3))} current borrowing. Current borrowing is a carrying amount, '
               'not a replacement for the annual principal schedule.</p>')
    output += (f'<p class="lr-meta"><a href="{esc(of.QUARTER_SOURCE)}" target="_blank" rel="noopener noreferrer">'
               '10-Q balance sheet, p. 1</a>. These figures and historical cash flows do not establish '
               'a funding surplus or shortfall for future payment dates.</p></div></div>')
    output += '<h3>What financing was reported afterward?</h3>'
    output += '<div class="lr-event"><div class="lr-meta">Jun 1–Aug 31, 2026</div><div>'
    for fact in review["quarter_flows"]:
        output += f'<p><strong>{esc(fact["label"])}: {esc(money(fact["value"], 3))}</strong></p>'
        if fact["value"] is None:
            output += f'<p class="lr-meta">Unavailable: {esc(fact["detail"])}</p>'
    output += ('<p>The cash-flow statement groups senior notes, term loans and other repayments. '
               'The MD&A separately describes about $3.1B scheduled debt repayment and $1.1B commercial-paper '
               'repayment, while the cash-flow commercial-paper row shows a dash. This unresolved '
               'presentation difference prevents assigning repayments to the old maturity buckets.</p>')
    output += (f'<p class="lr-meta"><a href="{esc(of.QUARTER_SOURCE)}" target="_blank" rel="noopener noreferrer">'
               '10-Q, pp. 5 and 33–35, filed Sep 11, 2026</a>. Quarter flows already affect the August balance sheet.</p></div></div>')
    for accession, event in of.FILINGS.items():
        output += (f'<article class="lr-event"><div class="lr-meta">Event {esc(event["event_date"])}<br>'
                   f'Filed {esc(event["filed"])}</div><div><p><strong>{esc(event["kind"])}</strong></p>'
                   f'<p>{esc(event["finding"])}</p><p class="lr-meta">'
                   f'<a href="{esc(event["source_url"])}" target="_blank" rel="noopener noreferrer">'
                   f'8-K · {esc(accession)}</a> · '
                   f'<a href="{esc(event["related"][0][1])}" target="_blank" rel="noopener noreferrer">Exhibit 99.1</a>'
                   '</p></div></article>')
    output += ('<p class="lr-meta">These two saved 8-Ks and their exhibits were checked. Neither establishes '
               'new borrowing after August 31. ')
    import oracle_revolver
    audit = oracle_revolver.for_snapshot(snapshot)
    output += ('The saved SEC financing-candidate index through October 15, 2026 lists only these two later filings. '
               if audit and audit["status"] == "Checked" else 'The saved financing-index check is Unavailable. ')
    output += ('Later or mismatched filings remain Not checked. Events are not added '
               'to the dated balance-sheet amounts.</p>')
    output += '<details><summary>Revolver terms and limits of this review</summary>'
    for term in of.TERMS:
        output += (f'<p><strong>{esc(term["label"])}</strong> · May 31, 2026<br>{esc(term["detail"])} '
                   f'<a href="{esc(term["source_url"])}" target="_blank" rel="noopener noreferrer">{esc(term["location"])}</a></p>')
    output += '<p class="lr-meta">The $10B revolver matures March 6, 2031, subject to the agreement. Capacity is not a payment due in that year.</p></details></section>'
    return output.replace("$", "&#36;")


def render(snapshot):
    """Deterministic, source-linked explanations; no ratings or implied repayment assumptions."""
    m = snapshot["metrics"]
    end = date.fromisoformat(m["cash"]["period_end"]).strftime("%b %d, %Y")
    source = snapshot["source_url"]
    output = CSS + '<section class="lr-review">'
    output += f'<header class="lr-lead"><div class="lr-meta">{esc(snapshot["ticker"])} · Debt and liquidity review</div>'
    output += f'<h2>{esc(snapshot["company_name"])}</h2><div class="lr-meta">Balance sheet: {esc(end)} · '
    output += f'<a href="{esc(source)}" target="_blank" rel="noopener noreferrer">Source 10-Q</a> · Balance-sheet facts checked {esc(snapshot["checked_as_of"])}</div>'
    output += f'<p class="lr-answer">{esc(headline(snapshot))}</p>'
    output += '<p class="lr-meta">Cash flows cover the trailing twelve months to this balance-sheet date. Borrowing is at carrying amount and excludes leases. Figures are rounded; evidence retains exact inputs.</p></header>'
    output += '<div class="lr-summary">'
    for label, value, detail in (
        ("Cash and equivalents", money(m["cash"]["value"]), "Excludes restricted cash"),
        ("Baseline gross debt", money(m["gross_debt"]["value"]), "Borrowing, excluding leases"),
        ("Net debt", money(m["net_debt"]["value"]), "Gross debt less cash"),
        ("Current ratio", f'{m["current_ratio"]["value"]:.2f}x' if m["current_ratio"]["value"] is not None else "Unavailable", "Current assets / current liabilities"),
    ):
        output += f'<div><div class="lr-stat-label">{esc(label)}</div><div class="lr-stat-value">{esc(value)}</div><div class="lr-stat-detail">{esc(detail)}</div></div>'
    from funding_conclusions import render as funding_html
    output += '</div>' + funding_html(snapshot)
    output += '<div class="lr-grid"><section class="lr-section"><h3>What remained after capital spending?</h3>'
    output += f'<div class="lr-meta">TTM ended {esc(end)} · USD billions · 10-K and 10-Q cash-flow statements</div>'
    import core_charts
    output += core_charts.render_flow(snapshot, money)
    output += f'<p class="lr-meta">{esc(m["free_cash_flow"]["notes"][0])}</p>'
    output += '<div class="lr-ratios">'
    for key, label in (("ocf_to_debt", "TTM operating cash flow / debt"), ("fcf_to_debt", "TTM free cash flow / debt")):
        output += f'<div><span class="lr-stat-label">{esc(label)}</span><strong>{esc(pct(m[key]))}</strong></div>'
    output += '</div><p class="lr-meta">Historical flows relative to period-end borrowing; these percentages are not repayment forecasts. FCF excludes acquisitions, dividends and other cash uses.</p></section>'
    output += '<section class="lr-section">'
    import borrowing_leases
    try:
        borrowing_model = borrowing_leases.build(snapshot)
        output += borrowing_leases.render_borrowing(borrowing_model, money)
    except (OSError, ValueError, KeyError, TypeError):
        borrowing_model = None
        output += '<h3>Borrowing and lease presentation unavailable</h3><p>The checked evidence is missing or does not match this review.</p>'
    assets = m["current_ratio"]["inputs"]["current_assets"]
    liabilities = m["current_ratio"]["inputs"]["current_liabilities"]
    if assets["value"] is not None and liabilities["value"] is not None:
        gap = assets["value"] - liabilities["value"]
        direction = "exceeded" if gap >= 0 else "were below"
        output += (f'<p>Current assets of {esc(money(assets["value"]))} {direction} current liabilities of '
                   f'{esc(money(liabilities["value"]))} by {esc(money(abs(gap)))}.</p>')
    output += '<p class="lr-meta">Current assets include items other than cash. This comparison does not establish whether specific payments can be met; a maturity schedule and financing terms are still needed.</p></section></div>'
    if borrowing_model is not None:
        output += borrowing_leases.render_leases(borrowing_model, money)
    from oracle_liquidity_view import render as oracle_liquidity_html
    output += oracle_liquidity_html(snapshot, money)
    from vistra_liquidity_view import render as vistra_liquidity_html
    output += vistra_liquidity_html(snapshot, money)
    from lly_liquidity_view import render as lilly_liquidity_html
    output += lilly_liquidity_html(snapshot, money)
    from lly_interest_view import render as lilly_interest_html
    output += lilly_interest_html(snapshot, money)
    from oracle_interest_view import render as oracle_interest_html
    output += oracle_interest_html(snapshot, money)
    from vistra_interest_view import render as vistra_interest_html
    output += vistra_interest_html(snapshot, money)
    import company_history
    output += company_history.render(snapshot, money)
    output += core_charts.render_balances(snapshot, money)
    output += maturity_html(snapshot)
    output += '<section class="lr-obligations"><h3>What sits outside baseline borrowing?</h3><p class="lr-meta">Selected disclosures, each shown separately. Commitments are not all payable now and are not added to debt or to each other.</p>'
    for item in NOTES[snapshot["ticker"]]:
        # The legacy workbook still consumes NOTES; the app shows these in its dedicated lease section.
        if snapshot["ticker"] == "ORCL" and item["name"] in {
                "Leases not yet commenced", "Operating lease liabilities", "Finance lease liabilities"}:
            continue
        output += (f'<article class="lr-obligation"><div><p><strong>{esc(item["name"])}</strong></p>'
                   f'<p>{esc(item["amount"])}</p><div class="lr-meta">As of {esc(item["date"])} · {esc(item["kind"])}</div></div>'
                   f'<div><p>{esc(item["detail"])}</p><a class="lr-meta" href="{esc(source)}" target="_blank" rel="noopener noreferrer">{esc(item["location"])}</a></div></article>')
    output += '</section><section class="lr-limits"><h3>What remains unresolved?</h3>'
    remaining = {"VST": "Historical interest-component coverage is shown above. Receivables/Nuveen allocation, Other composition and an all-in financing-cost denominator remain Unavailable.",
                 "ORCL": "Historical reported-line interest coverage is shown above; borrowing-only coverage, capitalization basis and covenant headroom remain Unavailable.",
                 "LLY": "Historical interest coverage uses derived operating income; its capitalization and other limits are stated above."}[snapshot["ticker"]]
    output += f'<p class="lr-meta">{esc(LIMITS[snapshot["ticker"]])} {remaining}</p></section></section>'
    return output.replace("$", "&#36;")


def evidence_html(snapshot: dict) -> str:
    """Read-only display; every input keeps its filing and selection information."""
    import ui
    from liquidity_pilot import leaf_inputs
    import ratio_review
    reconciled = {r["key"]: r for r in ratio_review.build(snapshot)["rows"]}
    metrics = snapshot["metrics"]
    output = '<details class="irt-evidence"><summary>Calculations and filing evidence</summary>'
    for key, metric in metrics.items():
        row = reconciled[key]
        output += f'<p><strong>{ui.esc(metric["label"])}</strong>'
        period = ((row["period_start"] + " to ") if row["period_start"] else "As of ") + row["period_end"]
        output += f'<br>{ui.esc(ratio_review.display(row))} · {ui.esc(period)}'
        if metric.get("formula"):
            output += f' · {ui.esc(metric["formula"])}'
        if metric["value"] is None:
            output += f'<br>Unavailable: {ui.esc(metric["detail"])}'
        for note in metric.get("notes", []):
            output += f'<br>{ui.esc(note)}'
        output += f'<br>{ui.esc(row["basis"])}'
        output += '</p><ul class="irt-ol" style="overflow-wrap:anywhere">'
        # Ratios and net debt reference the metrics already documented above.
        facts = [] if metric.get("unit") == "fraction" or metric["label"] in {"Net debt", "Free cash flow"} else leaf_inputs(metric)
        for fact in facts:
            value = f"{fact['value']:,.0f} USD" if fact["value"] is not None else "Unavailable"
            context = (f" · Context {fact['context_id']}, scale {fact['scale']}, decimals {fact['decimals']}"
                       if fact.get("context_id") else "")
            period = ((fact["period_start"] + " to ") if fact.get("period_start") else "") + fact["period_end"]
            output += (f'<li>{ui.esc(fact["label"])}: {ui.esc(value)} · {ui.esc(fact["tag"])}<br>'
                       f'Period {ui.esc(period)} · {ui.esc(fact.get("form") or "Form unavailable")} '
                       f'filed {ui.esc(fact.get("filed") or "Unavailable")} · '
                       f'{ui.esc(fact.get("accession") or "Accession unavailable")}<br>'
                       f'{ui.esc(fact.get("selection_rule") or fact.get("detail", ""))}{ui.esc(context)}'
                       f'<br><a href="{ui.esc(fact.get("source_url") or snapshot["source_url"])}" target="_blank" rel="noopener noreferrer">Source filing</a></li>')
        output += '</ul>'
    import company_history
    import borrowing_leases
    history = company_history.build(snapshot)
    lease_model = borrowing_leases.build(snapshot)
    lease_inputs = []
    for item in lease_model["leases"]:
        lease_inputs.extend(item.get("inputs", []))
        if item.get("reported_total"):
            lease_inputs.append(item["reported_total"])
    if lease_model["commitment"]["value"] is not None:
        lease_inputs.append(lease_model["commitment"])
    groups = [("Annual history source inputs", history["annual_facts"] + history["annual_flows"]),
              ("Lease source inputs", lease_inputs)]
    for heading, inputs in groups:
        output += f'<h4>{ui.esc(heading)}</h4><ul class="irt-ol" style="overflow-wrap:anywhere">'
        for fact in inputs:
            period = ((fact["period_start"] + " to ") if fact.get("period_start") else "As of ") + fact["period_end"]
            output += (f'<li>{ui.esc(fact.get("label") or fact["tag"])}: {fact["value"]:,.0f} USD · {ui.esc(fact["tag"])}<br>'
                       f'{ui.esc(period)} · {ui.esc(fact["form"])} filed {ui.esc(fact["filed"])} · {ui.esc(fact["accession"])}<br>'
                       f'{ui.esc(fact["selection_rule"])} · Context {ui.esc(fact.get("context_id") or "Unavailable")} · '
                       f'scale {ui.esc(fact.get("scale"))} · decimals {ui.esc(fact.get("decimals"))}<br>'
                       f'<a href="{ui.esc(fact["source_url"])}" target="_blank" rel="noopener noreferrer">Source filing</a></li>')
        output += '</ul>'
    debt = metrics["gross_debt"]
    # The reconciled borrowing explanation already includes filing-specific caveats;
    # metric notes were displayed with the calculation above.
    caveats = []
    if debt.get("not_added"):
        caveats.append("Excluded from baseline debt: " + debt["not_added"])
    output += ''.join(f'<p class="irt-note">{ui.esc(n)}</p>' for n in caveats)
    return output + '</details>'


def add_oracle_sheets(wb, snapshot):
    """Keep principal, carrying value and financing events on their own reporting bases."""
    import oracle_financing as of
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font
    from datetime import datetime
    review = of.for_snapshot(snapshot)
    if review is None:
        return
    schedule = wb.create_sheet("Oracle maturities")
    schedule.append(["Principal schedule", "Due from", "Due through", "USD (displayed in millions)",
                     "Reporting basis", "Notes"])
    if review["status"] != "Checked":
        schedule.append(["Unavailable", None, None, None, "", review["detail"]])
        return
    raw = wb.create_sheet("Oracle debt inputs")
    raw.append(["Input", "USD (displayed in millions)", "Reporting date", "Due from", "Due through", "Form", "Filed",
                "Accession", "Tag", "Context ID", "Scale", "Decimals", "Source", "Filing location", "Notes"])
    stamp = lambda text: datetime.strptime(text, "%Y-%m-%d") if text else None
    inputs = [*review["buckets"], *review["controls"].values()]
    labels = [f["label"] for f in review["buckets"]] + ["Reported principal total", "Discounts and issuance costs", "Reported carrying debt"]
    for fact, label in zip(inputs, labels):
        raw.append([label, fact["value"], stamp(review["as_of"]), stamp(fact.get("due_start")),
                    stamp(fact.get("due_end")), review["form"], stamp(review["filed"]), review["accession"],
                    fact["tag"], fact["context_id"], fact["scale"], fact["decimals"], review["source_url"],
                    review["location"], "Historical principal schedule; not rolled forward to August 31."])
        raw.cell(raw.max_row, 2).font = Font(name="Arial", size=10, color="0000FF")
    for i, fact in enumerate(review["buckets"], 2):
        schedule.append([fact["label"], stamp(fact["due_start"]), stamp(fact["due_end"]),
                         f"='Oracle debt inputs'!B{i}", "As of May 31, 2026",
                         "Fiscal 2032 onward, combined" if i == 7 else
                         "Full fiscal year; includes elapsed months at the August balance sheet" if i == 2 else ""])
        schedule.cell(i, 4).font = Font(name="Arial", size=10, color="008000")
    for label, formula, note in (
        ("Calculated principal total", "=SUM(D2:D7)", "Excludes interest and leases"),
        ("Reported principal total", "='Oracle debt inputs'!B8", ""),
        ("Principal reconciliation difference", "=D8-D9", "Zero means the six buckets equal the reported total"),
        ("Less: discounts and issuance costs", "='Oracle debt inputs'!B9", ""),
        ("Calculated carrying debt", "=D8-D11", "Same May 31 date and filing"),
        ("Reported carrying debt", "='Oracle debt inputs'!B10", ""),
        ("Carrying reconciliation difference", "=D12-D13", "Zero means principal less costs equals carrying debt"),
        ("Principal through fiscal 2031", "=SUM(D2:D6)", ""),
        ("Share through fiscal 2031", "=D15/D8" if review["total"] > 0 else None, "Fraction displayed as %"),
    ):
        schedule.append([label, None, None, formula, "As of May 31, 2026", note])
        if formula and "Oracle debt inputs" in formula:
            schedule.cell(schedule.max_row, 4).font = Font(name="Arial", size=10, color="008000")
    schedule["D16"].number_format = '0.0%;(0.0%);"-"'
    schedule.append(["August schedule", None, None, None, "Unavailable", review["current_schedule_detail"]])
    schedule.append(["Chart scope", None, None, None, "Annual buckets only", "Thereafter is combined in row 7 and is not plotted as one year."])
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Principal due by fiscal year at May 31, 2026"
    chart.add_data(Reference(schedule, min_col=4, min_row=2, max_row=6), titles_from_data=False)
    chart.set_categories(Reference(schedule, min_col=1, min_row=2, max_row=6))
    chart.legend = None
    # openpyxl calls the numeric axis y_axis even for a horizontal bar chart.
    chart.y_axis.title = "USD billions"
    chart.y_axis.scaling.min = 0
    chart.y_axis.numFmt = "0.0,,,"
    chart.x_axis.scaling.orientation = "maxMin"
    chart.series[0].graphicalProperties.solidFill = "737985"
    chart.series[0].graphicalProperties.line.solidFill = "737985"
    chart.height, chart.width = 9, 19
    schedule.add_chart(chart, "H2")
    timeline = wb.create_sheet("Oracle financing")
    timeline.append(["Event date", "Period start", "Period end", "Finding", "USD (displayed in millions)", "Amount basis",
                     "Form", "Filed", "Accession", "Tag", "Context ID", "Source", "Exhibit", "Notes"])
    for fact in review["quarter_flows"]:
        raw.append([fact["label"], fact["value"], stamp(fact["period_end"]), stamp(fact["period_start"]),
                    stamp(fact["period_end"]), fact["form"], stamp(fact["filed"]), fact["accession"],
                    fact["tag"], fact.get("context_id"), fact.get("scale"), fact.get("decimals"),
                    fact["source_url"], fact["location"], fact.get("detail", "")])
        raw.cell(raw.max_row, 2).font = Font(name="Arial", size=10, color="0000FF")
        value = f"='Oracle debt inputs'!B{raw.max_row}" if fact["value"] is not None else None
        timeline.append([None, stamp(fact["period_start"]), stamp(fact["period_end"]), fact["label"], value,
                         "Quarter cash flow, already in August balance sheet", fact["form"], stamp(fact["filed"]),
                         fact["accession"], fact["tag"], fact.get("context_id"), fact["source_url"], "",
                         "Debt repayments are not assigned to May maturity buckets; commercial-paper presentation differs in MD&A."
                         if fact["label"] == "Debt repayments" else "Equity, not debt; net of issuance costs."])
        timeline.cell(timeline.max_row, 5).font = Font(name="Arial", size=10, color="008000")
    for accession, event in of.FILINGS.items():
        timeline.append([stamp(event["event_date"]), None, None, event["kind"], None,
                         "Document review; no additional amount summed", event["form"], stamp(event["filed"]),
                         accession, None, None, event["source_url"], event["related"][0][1], event["finding"]])
    for term in of.TERMS:
        timeline.append([None, None, stamp(term["as_of"]), term["label"], term["capacity"],
                         "Historical facility capacity, not cash or debt" if term["capacity"] is not None else "Contract definition",
                         "10-K", stamp(review["filed"]), review["accession"], None, None, term["source_url"], None,
                         term["detail"] + " " + ("Matures " + term["maturity"] if term["maturity"] else term["location"])])
        if term["capacity"] is not None:
            timeline.cell(timeline.max_row, 5).font = Font(name="Arial", size=10, color="0000FF")
    for sheet, columns in ((schedule, (2, 3)), (raw, (3, 4, 5, 7)), (timeline, (1, 2, 3, 8))):
        for col in columns:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, col).number_format = "mmm d, yyyy"
    usd = '#,##0,,;(#,##0,,);"-"'
    for sheet, col in ((schedule, 4), (raw, 2), (timeline, 5)):
        for row in range(2, sheet.max_row + 1):
            if (sheet, row) != (schedule, 16):
                sheet.cell(row, col).number_format = usd
    for cell in (schedule["D10"], schedule["D14"]):
        cell.number_format = '#,##0.00,,;(#,##0.00,,);"-"'
    for sheet in (schedule, raw, timeline):
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, sheet.max_column).column_letter}{sheet.max_row}"


def workbook(snapshots: list[dict]) -> bytes:
    """A row per reported input, with Excel formulas referencing those rows for derived results."""
    from io import BytesIO
    from liquidity_pilot import PILOTS, leaf_inputs
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Debt and liquidity"
    ws.append(["Ticker", "Metric", "Value", "Unit", "Period end", "Formula", "Status", "Notes"])
    raw = wb.create_sheet("Liquidity inputs")
    raw.append(["Ticker", "Metric", "Reported input", "Value (USD)", "Period end", "Form", "Filed",
                "Accession", "Tag", "Selection rule", "Source document", "Notes",
                "Context ID", "Scale", "Decimals", "Period start", "TTM role", "Coefficient"])
    usd_format = '#,##0,,;(#,##0,,);"-"'
    for snapshot in snapshots:
        metric_rows = {}
        for key, metric in snapshot["metrics"].items():
            metric_rows[key] = ws.max_row + 1
            refs = []
            # Net debt references the two already exported metrics; no duplicate input rows.
            derived_refs = {"net_debt": ("gross_debt", "cash", "-"),
                            "free_cash_flow": ("operating_cash_flow", "capital_spending", "-"),
                            "ocf_to_debt": ("operating_cash_flow", "gross_debt", "/"),
                            "fcf_to_debt": ("free_cash_flow", "gross_debt", "/")}
            if key not in derived_refs:
                for fact in leaf_inputs(metric):
                    raw.append([snapshot["ticker"], metric["label"], fact["label"], fact["value"],
                                fact["period_end"], fact.get("form"), fact.get("filed"), fact.get("accession"),
                                fact.get("tag"), fact.get("selection_rule"), fact.get("source_url") or snapshot["source_url"],
                                "; ".join(fact.get("notes", [])) or fact.get("detail", ""),
                                fact.get("context_id"), fact.get("scale"), fact.get("decimals"),
                                fact.get("period_start"), fact.get("role"), fact.get("coefficient")])
                    raw.cell(raw.max_row, 4).number_format = usd_format
                    raw.cell(raw.max_row, 4).font = Font(name="Arial", size=10, color="0000FF")
                    refs.append((fact.get("coefficient", 1), f"'Liquidity inputs'!D{raw.max_row}"))
            value = None
            if metric["value"] is not None:
                if key in derived_refs:
                    a, b, op = derived_refs[key]
                    value = f"=C{metric_rows[a]}{op}C{metric_rows[b]}"
                elif key == "current_ratio":
                    value = "=" + "/".join(ref for _, ref in refs)
                else:
                    value = "=" + refs[0][1] + "".join(("-" if sign < 0 else "+") + ref for sign, ref in refs[1:])
            notes = [*metric.get("notes", []), *metric.get("caveats", [])]
            if metric.get("not_added"):
                notes.append("Excluded: " + metric["not_added"])
            ws.append([snapshot["ticker"], metric["label"], value,
                       "USD (displayed in millions)" if metric["unit"] == "USD" else
                       "fraction (displayed as %)" if metric["unit"] == "fraction" else "ratio",
                       metric["period_end"], metric.get("formula") or metric.get("tag"),
                       ("Reported" if key == "cash" else "Calculated") if metric["value"] is not None else "Unavailable",
                       "; ".join(notes) or metric.get("detail", "")])
            ws.cell(ws.max_row, 3).number_format = (usd_format if metric["unit"] == "USD" else
                                                   '0.0%;(0.0%);"-"' if metric["unit"] == "fraction" else '0.00"x"')
            if key == "cash":
                ws.cell(ws.max_row, 3).font = Font(name="Arial", size=10, color="008000")
    about = wb.create_sheet("About")
    about.append(["Item", "Description"])
    about.append(["Snapshot", "Saved pilot evidence checked October 14, 2026. Not live data or investment advice."])
    about.append(["Units", "Dollar inputs are stored in USD and displayed in millions. Ratios are multiples."])
    about.append(["Calculation", "Excel recalculates formulas when opened. Each formula traces to reported inputs."])
    about.append(["Colors", "Blue numbers are reported inputs. Green formulas link a reported value from another sheet. Black formulas calculate results."])
    about.append(["Scope", "Baseline borrowing excludes leases. FCF subtracts the reported capital-spending line from OCF; acquisitions are excluded. Historical flow/debt ratios are not repayment forecasts. This is not a complete obligations analysis."])
    obligations = wb.create_sheet("Other obligations")
    obligations.append(["Ticker", "Disclosure", "Amount as disclosed", "Reporting date", "Classification", "Filing location", "Accession", "Source", "Notes"])
    for snapshot in snapshots:
        for item in NOTES[snapshot["ticker"]]:
            obligations.append([snapshot["ticker"], item["name"], item["amount"], item["date"], item["kind"],
                                item["location"], PILOTS[snapshot["ticker"]]["accession"], snapshot["source_url"], item["detail"]])
        add_oracle_sheets(wb, snapshot)
    if "Oracle maturities" in wb.sheetnames:
        about.append(["Oracle maturity review", "Historical principal schedule at May 31, 2026. Annual buckets exclude interest and leases. August schedule and revolver availability are Unavailable. Inputs and subsequent document checks are in Oracle debt inputs and Oracle financing; no events are rolled into the balance sheet."])
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        for row in sheet:
            for cell in row:
                color = cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else "000000"
                cell.font = Font(name="Arial", size=10, color=color, bold=cell.row == 1)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 23
    ws.column_dimensions["H"].width = 65
    raw.column_dimensions["K"].width = 55
    raw.column_dimensions["L"].width = 55
    about.column_dimensions["B"].width = 90
    if "Oracle maturities" in wb.sheetnames:
        for name in ("Oracle maturities", "Oracle debt inputs", "Oracle financing"):
            sheet = wb[name]
            sheet.column_dimensions["A"].width = 38 if name != "Oracle financing" else 19
            sheet.row_dimensions[1].height = 32
            for row in range(2, sheet.max_row + 1):
                sheet.row_dimensions[row].height = (34 if row < 17 else 65) if name == "Oracle maturities" else 88
        for row in range(2, wb["Oracle maturities"].max_row + 1):
            wb["Oracle maturities"].cell(row, 5).alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        wb["Oracle maturities"].column_dimensions["F"].width = 55
        wb["Oracle debt inputs"].column_dimensions["M"].width = 55
        wb["Oracle debt inputs"].column_dimensions["O"].width = 60
        wb["Oracle financing"].column_dimensions["D"].width = 55
        wb["Oracle financing"].column_dimensions["L"].width = 55
        wb["Oracle financing"].column_dimensions["M"].width = 55
        wb["Oracle financing"].column_dimensions["N"].width = 80
    out = BytesIO()
    wb.save(out)
    return out.getvalue()



def saved_workbook(snapshots, folder=None, ticker=None):
    """Serve the styled export only with the exact source snapshots used to build it."""
    import hashlib
    import json
    from pathlib import Path
    folder = Path(folder) if folder is not None else Path(__file__).parent / "samples"
    try:
        if ticker not in {None, "LLY"}:
            return None
        stem = "lly_liquidity" if ticker == "LLY" else "liquidity_pilots"
        manifest = json.loads((folder / f"{stem}.manifest.json").read_text())
        saved = json.loads((folder / "liquidity_pilots.json").read_text())
        if snapshots != saved["snapshots"]:
            return None
        if ticker == "LLY":
            import export_guide
            if manifest.get("guide_sha256") != export_guide.fingerprint():
                return None
            import core_charts
            if manifest.get("charts_sha256") != core_charts.fingerprint():
                return None
            import company_history
            if manifest.get("history_sha256") != company_history.fingerprint():
                return None
            import ratio_review
            if manifest.get("ratios_sha256") != ratio_review.fingerprint():
                return None
            import borrowing_leases
            if manifest.get("borrowing_leases_sha256") != borrowing_leases.fingerprint():
                return None
            import funding_conclusions as funding
            if manifest.get("funding_sha256") != funding.fingerprint() or funding.load(ticker, folder) is None:
                return None
        expected = {"liquidity_pilots.json", "lly_liquidity.json", "lly_financing.json", "lly_interest.json"} if ticker else {"liquidity_pilots.json", "oracle_financing.json"}
        for name, digest in manifest["sources"].items():
            if name not in expected:
                return None
            if hashlib.sha256((folder / name).read_bytes()).hexdigest() != digest:
                return None
        if set(manifest["sources"]) != expected:
            return None
        content = (folder / f"{stem}.xlsx").read_bytes()
        return content if hashlib.sha256(content).hexdigest() == manifest["workbook_sha256"] else None
    except (OSError, ValueError, KeyError, TypeError):
        return None


SELECTOR_NAMES = {"LLY": "Eli Lilly", "ORCL": "Oracle", "VST": "Vistra"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def render_download(st, snapshot, snapshots) -> None:
    """Each company's checked workbook, served only while its manifest still matches the saved sources."""
    ticker = snapshot["ticker"]
    import export_guide
    st.caption(export_guide.CAPTION)
    if ticker == "LLY":
        content = saved_workbook(snapshots, ticker="LLY")
        if content:
            st.download_button("Download LLY review and evidence", content,
                file_name="Lilly_Liquidity_2026-06-30.xlsx", mime=XLSX)
        else:
            st.caption("Lilly review workbook Unavailable: its saved sources must match the workbook.")
    if ticker == "VST":
        import vistra_liquidity as vl
        review = vl.for_snapshot(snapshot)
        content = vl.workbook_bytes(review) if review else None
        if content:
            st.download_button("Download VST review and evidence", content,
                file_name="Vistra_Liquidity_2026-06-30.xlsx", mime=XLSX)
        else:
            st.caption("Vistra repayment workbook Unavailable: its saved review must match the workbook.")
    if ticker == "ORCL":
        import oracle_liquidity as ol
        review = ol.for_snapshot(snapshot)
        content = ol.workbook_bytes(review) if review else None
        if content:
            st.download_button("Download ORCL review and evidence", content,
                file_name="Oracle_Liquidity_2026-08-31.xlsx", mime=XLSX)
        else:
            st.caption("Oracle cash-analysis workbook Unavailable: its saved source review must match the workbook.")


def render_app() -> None:
    """The app's default page: pick a checked company and its saved review opens directly.

    Only the selected review and its workbook are built, so the page makes no SEC, Yahoo or Claude request and
    needs no secrets. The selection survives a visit to general filing research (persist_state="session")."""
    import json
    import streamlit as st
    from ui import notice
    from liquidity_pilot import SAMPLE, financing_html
    try:
        saved = json.loads(SAMPLE.read_text(encoding="utf-8"))
        snapshots = {s["ticker"]: s for s in saved["snapshots"]}
    except (OSError, ValueError, KeyError, TypeError):
        snapshots = {}
    if not snapshots:
        st.markdown(notice("Saved reviews unavailable", "The saved debt and liquidity reviews could not be loaded. "
                           "No live figures are substituted.", "caution"), unsafe_allow_html=True)
        return
    # A selection from an older session (or a company no longer saved) falls back to the first checked review.
    if st.session_state.get("review_company") not in snapshots:
        st.session_state.review_company = next(iter(snapshots))
    ticker = st.segmented_control(
        "Company", list(snapshots), required=True, key="review_company", persist_state="session",
        format_func=lambda t: f"{SELECTOR_NAMES.get(t, t)} ({t})")
    snapshot = snapshots[ticker]
    st.caption("Saved filing reviews · source dates shown in each review · no live or paid request")
    try:
        page = render(snapshot) + evidence_html(snapshot) + financing_html(snapshot)
    except (OSError, KeyError, TypeError, ValueError):
        # A saved file edited out of shape shows a stated gap, not a stack trace or partial figures.
        import traceback
        traceback.print_exc()
        st.markdown(notice(f"{ticker} review unavailable", "This saved review is incomplete or out of date, so it is "
                           "not shown. The other saved reviews are unaffected.", "caution"), unsafe_allow_html=True)
        return
    render_download(st, snapshot, saved["snapshots"])
    st.markdown(page, unsafe_allow_html=True)
    with st.expander("Optional combined export"):
        import export_guide
        st.caption(export_guide.COMBINED_CAPTION)
        content = saved_workbook(saved["snapshots"])
        if content:
            st.download_button("Download all three company reviews (optional)", content,
                file_name="Liquidity_Pilot_2026-10-14.xlsx", mime=XLSX)
        else:
            st.caption("Combined workbook Unavailable: its saved sources must match the workbook.")
